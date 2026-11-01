# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1", "pyyaml>=6"]
# ///
"""Проверка docx на оформление по ГОСТ (ГОСТ 7.32 / требования КР к научным работам).

    uv run check_format.py work.docx [--json] [--fix]

Поля, шрифт, кегль, интервал, отступ, выравнивание, номера страниц, заголовки,
подписи рисунков и таблиц, оглавление, список литературы, оставленные метки.
--fix — исправить безопасное (поля, шрифт, кегль, интервал, отступ, выравнивание
основного текста, номера страниц) в копию «<имя>_исправлено.docx»; заголовки,
подписи и список литературы не трогаются.
Нормы — ГОСТ; если выше файла есть university/settings.yaml, шрифт, кегль, интервал,
отступ и поля берутся оттуда (раздел document).
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import asdict, dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

sys.path.insert(0, str(Path(__file__).parent))
from common import DataError, fail, utf8_stdout  # noqa: E402
from docxkit import MARK_RE, _field, _set_fonts  # noqa: E402
from settings import DEFAULTS, load_settings  # noqa: E402

FONT = "Times New Roman"
SIZE = 14.0
SPACING = 1.5
INDENT_CM = 1.25
MARGINS = {"left": (30, 30), "right": (10, 15), "top": (20, 20), "bottom": (20, 20)}
TOL_MM = 1.0
REFS_RE = re.compile(r"^\s*(список\s+(использованн\w+\s+)?(литературы|источников)|библиографический список|"
                     r"литература|пайдаланылган адабияттар\w*.*|адабияттар\w*|references|bibliography)\s*$", re.I)
TOC_RE = re.compile(r"^\s*(содержание|оглавление|мазмуну|contents)\s*$", re.I)
CAPTION_RE = re.compile(r"^\s*(рисунок|рис\.|сүрөт|figure|fig\.|таблица|table)\s*([\dА-ЯA-Z.]*)\s*(.)?", re.I)
NON_BODY_STYLES = ("heading", "заголовок", "title", "toc", "caption", "название", "оглавление", "subtitle")


@dataclass
class Problem:
    level: str      # critical | warning
    area: str
    text: str
    fixable: bool = False


# ---------- действующие значения с учётом стилей ----------

def _chain(style):
    while style is not None:
        yield style
        style = style.base_style


def _defaults(doc):
    return doc.styles.element.find(qn("w:docDefaults"))


def eff_para(doc, p, attr):
    v = getattr(p.paragraph_format, attr)
    if v is not None:
        return v
    for st in _chain(p.style):
        v = getattr(st.paragraph_format, attr)
        if v is not None:
            return v
    if attr == "line_spacing":
        d = _defaults(doc)
        sp = d.find(qn("w:pPrDefault") + "/" + qn("w:pPr") + "/" + qn("w:spacing")) if d is not None else None
        if sp is not None and sp.get(qn("w:line")) and sp.get(qn("w:lineRule"), "auto") == "auto":
            return int(sp.get(qn("w:line"))) / 240
        return 1.0
    return None


def _rfonts_name(el) -> str | None:
    if el is None:
        return None
    f = el.find(qn("w:rFonts"))
    if f is None:
        return None
    if f.get(qn("w:ascii")):
        return f.get(qn("w:ascii"))
    if f.get(qn("w:asciiTheme")):
        return "шрифт темы (обычно Calibri)"
    return None


def eff_font(doc, p, run) -> str:
    if run.font.name:
        return run.font.name
    rpr = run._r.rPr
    th = _rfonts_name(rpr)
    if th:
        return th
    for st in _chain(p.style):
        n = _rfonts_name(st.element.rPr)
        if n:
            return n
    d = _defaults(doc)
    n = _rfonts_name(d.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))) if d is not None else None
    return n or "Calibri"


def eff_size(doc, p, run) -> float:
    if run.font.size:
        return run.font.size.pt
    for st in _chain(p.style):
        if st.font.size:
            return st.font.size.pt
    d = _defaults(doc)
    sz = d.find(qn("w:rPrDefault") + "/" + qn("w:rPr") + "/" + qn("w:sz")) if d is not None else None
    return int(sz.get(qn("w:val"))) / 2 if sz is not None else 10.0


def is_heading(p) -> bool:
    name = (p.style.name or "").lower()
    if any(name.startswith(x) for x in ("heading", "заголовок")):
        return True
    ppr = p._p.pPr
    return ppr is not None and ppr.find(qn("w:outlineLvl")) is not None


def is_body(doc, p) -> bool:
    name = (p.style.name or "").lower()
    if any(x in name for x in NON_BODY_STYLES) or is_heading(p):
        return False
    text = p.text.strip()
    if len(text) < 60 or CAPTION_RE.match(text) or "TOC \\o" in p._p.xml:
        return False
    al = eff_para(doc, p, "alignment")
    return al not in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT)


def _mm(length) -> float:
    return round(length / 36000, 1) if length is not None else 0.0


# ---------- проверка ----------

def check(path: str | Path) -> tuple[list[Problem], dict]:
    doc = Document(str(path))
    probs: list[Problem] = []
    add = lambda lvl, area, text, fix=False: probs.append(Problem(lvl, area, text, fix))  # noqa: E731

    for i, s in enumerate(doc.sections, 1):
        where = f" (раздел {i})" if len(doc.sections) > 1 else ""
        if s.page_width and (abs(_mm(s.page_width) - 210) > 2 or abs(_mm(s.page_height) - 297) > 2):
            add("critical", "страница", f"формат {_mm(s.page_width):.0f}×{_mm(s.page_height):.0f} мм, нужен A4 210×297{where}", True)
        for side, (lo, hi) in MARGINS.items():
            v = _mm(getattr(s, f"{side}_margin"))
            if v < lo - TOL_MM or v > hi + TOL_MM:
                need = f"{lo}" if lo == hi else f"{lo}–{hi}"
                ru = {"left": "левое", "right": "правое", "top": "верхнее", "bottom": "нижнее"}[side]
                add("critical", "поля", f"{ru} поле {v:g} мм, нужно {need} мм{where}", True)

    has_page = False
    for s in doc.sections:
        for part in (s.footer, s.header, s.first_page_footer, s.even_page_footer):
            try:
                xml = part._element.xml if part._element is not None else ""
            except Exception:
                xml = ""
            if re.search(r"instrText[^>]*>\s*PAGE\b", xml) or 'w:fldSimple w:instr=" PAGE' in xml \
                    or re.search(r'w:instr="\s*PAGE', xml):
                has_page = True
    if not has_page:
        add("critical", "номера страниц", "нет номеров страниц (нужны внизу по центру)", True)

    normal = doc.styles["Normal"]
    nf = _rfonts_name(normal.element.rPr) or _rfonts_name(
        _defaults(doc).find(qn("w:rPrDefault") + "/" + qn("w:rPr")))
    if nf != FONT:
        add("warning", "стиль «Обычный»", f"шрифт стиля {nf or 'не задан'}, нужен {FONT}", True)

    body = [p for p in doc.paragraphs if is_body(doc, p)]
    stats = {"paragraphs": len(doc.paragraphs), "body": len(body), "tables": len(doc.tables),
             "images": len(doc.inline_shapes)}
    bad = {"font": [], "size": [], "spacing": [], "indent": [], "align": []}
    for p in body:
        runs = [r for r in p.runs if r.text.strip()]
        fonts = {eff_font(doc, p, r) for r in runs}
        sizes = {eff_size(doc, p, r) for r in runs}
        if fonts - {FONT}:
            bad["font"].append((p, ", ".join(sorted(fonts - {FONT}))))
        if sizes - {SIZE}:
            bad["size"].append((p, ", ".join(f"{x:g}" for x in sorted(sizes - {SIZE}))))
        ls = eff_para(doc, p, "line_spacing")
        if not isinstance(ls, float) or abs(ls - SPACING) > 0.05:
            bad["spacing"].append((p, f"{ls:g}" if isinstance(ls, float) else "точный"))
        ind = eff_para(doc, p, "first_line_indent")
        if ind is None or abs(ind / 360000 - INDENT_CM) > 0.1:
            bad["indent"].append((p, f"{(ind or 0) / 360000:.2f} см"))
        if eff_para(doc, p, "alignment") != WD_ALIGN_PARAGRAPH.JUSTIFY:
            bad["align"].append((p, ""))

    def report(key, area, what, lvl_share=0.1):
        items = bad[key]
        if not items:
            return
        share = len(items) / max(len(body), 1)
        lvl = "critical" if share > lvl_share else "warning"
        sample = items[0][0].text.strip()[:60]
        vals = sorted({v for _, v in items if v})
        add(lvl, area, f"{what}: {len(items)} из {len(body)} абзацев"
            + (f" ({', '.join(vals[:3])})" if vals else "") + f"; например: «{sample}…»", True)

    report("font", "шрифт", f"не {FONT}")
    report("size", "кегль", f"кегль не {SIZE:g} пт")
    report("spacing", "интервал", f"интервал не {SPACING:g}")
    report("indent", "абзацный отступ", f"отступ первой строки не {INDENT_CM} см", 0.5)
    report("align", "выравнивание", "не по ширине", 0.5)

    heads = [p for p in doc.paragraphs if is_heading(p) and p.text.strip()]
    stats["headings"] = len(heads)
    if not heads and len(body) > 30:
        add("warning", "заголовки", "заголовки набраны без стилей «Заголовок 1/2» — оглавление не соберётся")
    for h in heads:
        if h.text.strip().endswith("."):
            add("warning", "заголовки", f"точка в конце заголовка: «{h.text.strip()[:60]}»")

    texts = [p.text.strip() for p in doc.paragraphs]
    has_toc_field = "TOC \\o" in doc.element.body.xml or any(TOC_RE.match(x) for x in texts)
    if len(body) > 40 and not has_toc_field:
        add("warning", "оглавление", "нет содержания / оглавления")

    tab_caps = [x for x in texts if re.match(r"^(таблица|table)\s+[\dА-ЯA-Z]", x, re.I)]
    fig_caps = [x for x in texts if re.match(r"^(рисунок|рис\.|сүрөт|figure|fig\.)\s*[\dА-ЯA-Z]", x, re.I)]
    data_tables = [t_ for t_ in doc.tables if len(t_.rows) > 1]
    if len(tab_caps) < len(data_tables):
        add("warning", "подписи", f"таблиц {len(data_tables)}, подписей «Таблица N – …» {len(tab_caps)}")
    if len(fig_caps) < stats["images"]:
        add("warning", "подписи", f"рисунков {stats['images']}, подписей «Рисунок N – …» {len(fig_caps)}")
    for c in tab_caps + fig_caps:
        if c.lower().startswith(("рис.", "fig.")) or not re.match(r"^\S+\s+[\dА-ЯA-Z.]+\s+[–—-]\s+\S", c):
            add("warning", "подписи", f"подпись не по форме «Рисунок 1 – Название»: «{c[:60]}»")

    ref_idx = next((i for i, x in enumerate(texts) if REFS_RE.match(x)), None)
    if ref_idx is None:
        if len(body) > 20:
            add("warning", "литература", "не найден раздел «Список литературы» / «Список использованных источников»")
    else:
        entries = []
        for p in doc.paragraphs[ref_idx + 1:]:
            if is_heading(p) or re.match(r"^(приложение|тиркеме|appendix)\b", p.text.strip(), re.I):
                break
            if p.text.strip():
                entries.append(p.text.strip())
        stats["references"] = len(entries)
        if not entries:
            add("critical", "литература", "список литературы пуст")
        apa_like = sum(1 for e in entries if re.search(r"\(\d{4}[a-z]?\)\.", e))
        ieee_like = sum(1 for e in entries if re.match(r"^\[\d+\]", e))
        if entries and apa_like < len(entries) / 2 and ieee_like < len(entries) / 2:
            no_dash = [e for e in entries if " – " not in e and "//" not in e]
            if no_dash:
                add("warning", "литература", f"{len(no_dash)} записей не по ГОСТ Р 7.0.100 (нет «. – » / «//»): "
                    f"«{no_dash[0][:70]}»")
        no_date = [e for e in entries if re.search(r"https?://", e) and not re.search(
            r"дата обращения|accessed|кирүү күнү", e, re.I) and "doi" not in e.lower()]
        if no_date and apa_like < len(entries) / 2:
            add("warning", "литература", f"{len(no_date)} электронных источников без «(дата обращения: …)»")

    marks = []
    for x in texts + [c.text for t_ in doc.tables for r in t_.rows for c in r.cells]:
        marks += [m.group(0) for m in MARK_RE.finditer(x)]
    if marks:
        uniq = sorted({m.split(":")[0].strip("[]") for m in marks})
        add("critical", "метки", f"в тексте осталось меток: {len(marks)} ({', '.join(uniq)}) — чистовик сдавать нельзя")
    return probs, stats


def use_settings(start: Path) -> None:
    """Нормы проверки из university/settings.yaml → document; нет файла — ГОСТ (значения выше)."""
    global FONT, SIZE, SPACING, INDENT_CM, MARGINS
    d = load_settings(start)["document"]
    FONT, SIZE, SPACING, INDENT_CM = d["font"], float(d["size"]), float(d["line_spacing"]), float(d["indent_cm"])
    m, dm = d["margins_mm"], DEFAULTS["document"]["margins_mm"]
    MARGINS = {side: MARGINS[side] if m[side] == dm[side] else (m[side], m[side]) for side in MARGINS}


# ---------- безопасное исправление ----------

def fix(path: str | Path) -> Path:
    doc = Document(str(path))
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin, s.right_margin = Mm(MARGINS["left"][1]), Mm(MARGINS["right"][1])
        s.top_margin, s.bottom_margin = Mm(MARGINS["top"][1]), Mm(MARGINS["bottom"][1])
    if not any(re.search(r"instrText[^>]*>\s*PAGE\b", s.footer._element.xml) for s in doc.sections):
        for s in doc.sections:
            p = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _field(p.add_run(), "PAGE", "1")
    normal = doc.styles["Normal"]
    _set_fonts(normal, FONT)
    normal.font.size = Pt(SIZE)
    normal.paragraph_format.line_spacing = SPACING
    for p in doc.paragraphs:
        if not is_body(doc, p):
            continue
        pf = p.paragraph_format
        pf.line_spacing = SPACING
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if pf.left_indent in (None, 0):
            pf.first_line_indent = Cm(INDENT_CM)
        for r in p.runs:
            r.font.name = FONT
            rpr = r._r.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None:
                for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                    fonts.attrib.pop(qn(a), None)
                fonts.set(qn("w:eastAsia"), FONT)
                fonts.set(qn("w:cs"), FONT)
            r.font.size = Pt(SIZE)
    path = Path(path)
    out = path.with_name(f"{path.stem}_исправлено{path.suffix}")
    doc.save(str(out))
    return out


def print_report(probs: list[Problem], stats: dict, path: Path) -> None:
    crit = [p for p in probs if p.level == "critical"]
    warn = [p for p in probs if p.level == "warning"]
    print(f"Файл: {path}")
    print(f"Абзацев основного текста: {stats['body']}, таблиц: {stats['tables']}, рисунков: {stats['images']}, "
          f"заголовков: {stats.get('headings', 0)}, записей в списке литературы: {stats.get('references', '—')}")
    if not probs:
        print("Замечаний нет: оформление соответствует ГОСТ.")
        return
    for title, items in (("КРИТИЧНО", crit), ("ВНИМАНИЕ", warn)):
        if items:
            print(f"\n{title} ({len(items)}):")
            for i, p in enumerate(items, 1):
                print(f"  {i}. [{p.area}] {p.text}" + ("  (исправимо --fix)" if p.fixable else ""))


def main() -> None:
    utf8_stdout()
    ap = argparse.ArgumentParser(description="Проверка docx на оформление по ГОСТ")
    ap.add_argument("file")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--fix", action="store_true", help="исправить безопасное в копию")
    args = ap.parse_args()
    path = Path(args.file).expanduser().resolve()
    if not path.exists() or path.suffix.lower() != ".docx":
        fail(f"нужен существующий файл .docx: {path}")
    try:
        use_settings(path.parent)
    except DataError as e:
        fail(str(e))
    try:
        probs, stats = check(path)
    except Exception as e:  # битый docx
        fail(f"не удалось открыть {path.name}: {e}")
    if args.fix:
        out = fix(path)
        probs, stats = check(out)
        print(f"Исправленная копия: {out}\n")
        path = out
    if args.json:
        print(json.dumps({"file": str(path), "stats": stats, "problems": [asdict(p) for p in probs]},
                         ensure_ascii=False, indent=2))
    else:
        print_report(probs, stats, path)
    sys.exit(1 if any(p.level == "critical" for p in probs) else 0)


if __name__ == "__main__":
    main()
