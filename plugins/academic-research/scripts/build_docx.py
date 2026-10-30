# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1", "pyyaml>=6"]
# ///
"""Сборка научной работы в docx по ГОСТ из work.yaml + chapters/*.md + sources.yaml.

    uv run build_docx.py "<папка работы>/work.yaml" [--part avtoreferat] [--strict]

Титул по типу работы, оглавление (поле Word), нумерация глав и разделов,
подписи рисунков и таблиц, ссылки [@key] → [1] / (Автор, год), список
литературы через format_refs. Непроверенные источники в список не попадают,
ссылка на них печатается как [НЕ ПРОВЕРЕНО]. Итог: out/<type>.docx.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from common import (DISSERTATIONS, SHORT_TYPES, DataError, fail, load_sources,  # noqa: E402
                    load_work, out_dir, university_name, utf8_stdout)
from docxkit import (Ctx, add_inline, heading, new_document, page_break, para,  # noqa: E402
                     parse_cite, render_markdown, scan_citations, toc)
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm  # noqa: E402
from format_refs import cite, format_entry, sort_entries  # noqa: E402
from i18n import t  # noqa: E402


# ---------- титульные листы ----------

def _blank(doc, n: int) -> None:
    for _ in range(n):
        para(doc)


def _org_block(doc, w: dict, lang: str) -> None:
    para(doc, t("ministry", lang), align="center", spacing=1.0)
    para(doc, university_name(w, lang), bold=True, align="center", spacing=1.0)
    if w.get("faculty"):
        para(doc, w["faculty"], align="center", spacing=1.0)
    if w.get("department"):
        para(doc, f'{t("department", lang)} «{w["department"]}»', align="center", spacing=1.0)


def _person(p: dict) -> str:
    return ", ".join(x for x in (p.get("degree"), p.get("name")) if x)


def title_student(doc, w: dict, lang: str) -> None:
    """Реферат, доклад, курсовая, дипломная, магистерская."""
    typ = w["type"]
    _org_block(doc, w, lang)
    _blank(doc, 1)
    if typ in ("diplom", "magistr"):
        for line in (t("admit", lang), t("head", lang), t("blank", lang), "«___» ____________ 20___ г."):
            para(doc, line, align="right", spacing=1.0)
        _blank(doc, 2)
    else:
        _blank(doc, 4)
    para(doc, w.get("doc_label") or t(f"label_{typ}", lang), bold=True, align="center")
    if w.get("discipline") and typ in ("referat", "kursovaya", "doklad"):
        para(doc, f'{t("discipline", lang)} «{w["discipline"]}»', align="center")
    para(doc, f'{t("on_topic", lang)}: «{w.get("title", "")}»', bold=True, align="center")
    sp = w.get("specialty") or {}
    if typ == "magistr" and sp.get("code"):
        para(doc, f'{t("specialty", lang)}: {sp["code"]} – {sp.get("name", "")}', align="center")
    elif w.get("program"):
        para(doc, f'{t("program", lang)}: {w["program"]}', align="center")
    _blank(doc, 3)
    authors = w.get("authors") or [{}]
    para(doc, f'{t("author", lang)}: {authors[0].get("name", "")}', align="right", spacing=1.0)
    if w.get("supervisor", {}).get("name"):
        label = t("supervisor_short" if typ in ("referat", "kursovaya", "doklad") else "supervisor", lang)
        para(doc, f'{label}: {_person(w["supervisor"])}', align="right", spacing=1.0)
    if (w.get("consultant") or {}).get("name"):
        para(doc, f'{t("consultant", lang)}: {_person(w["consultant"])}', align="right", spacing=1.0)
    _blank(doc, 4)
    para(doc, f'{w.get("city", "")} {w.get("year", "")}'.strip(), align="center")
    page_break(doc)


def title_kandidat(doc, w: dict, lang: str, *, avtoreferat: bool = False) -> None:
    sp = w.get("specialty") or {}
    if not avtoreferat:
        para(doc, university_name(w, lang), align="center", spacing=1.0)
    else:
        d = (w.get("avtoreferat") or {}).get("defense") or {}
        if d.get("organization"):
            para(doc, d["organization"], align="center", spacing=1.0)
        if d.get("council"):
            para(doc, f'Диссертационный совет {d["council"]}' if lang == "ru" else d["council"],
                 align="center", spacing=1.0)
    _blank(doc, 1)
    para(doc, t("manuscript", lang), align="right")
    para(doc, "УДК " + w["udc"] if w.get("udc") else "", align="left", indent=False)
    _blank(doc, 2)
    para(doc, (w.get("authors") or [{}])[0].get("name", ""), bold=True, align="center")
    _blank(doc, 1)
    para(doc, (w.get("title") or "").upper(), bold=True, align="center")
    _blank(doc, 1)
    if sp.get("code"):
        para(doc, f'{sp["code"]} – {sp.get("name", "")}', align="center")
    _blank(doc, 1)
    label = t("label_avtoreferat" if avtoreferat else "label_kandidat", lang, science=sp.get("science", "___ наук"))
    para(doc, label, bold=avtoreferat, align="center")
    _blank(doc, 2)
    if not avtoreferat and w.get("supervisor", {}).get("name"):
        para(doc, t("supervisor", lang) + ":", align="right", spacing=1.0)
        para(doc, _person(w["supervisor"]), align="right", spacing=1.0)
    _blank(doc, 6 if avtoreferat else 4)
    para(doc, f'{w.get("city", "")} {w.get("year", "")}'.strip(), align="center")
    page_break(doc)


def avtoreferat_back(doc, w: dict, lang: str) -> None:
    """Оборот титула автореферата: где выполнена, руководитель, оппоненты, защита."""
    a = w.get("avtoreferat") or {}
    d = a.get("defense") or {}
    blank = "____________________"

    def line(label: str, value: str) -> None:
        p = para(doc, "", align="left", indent=False, spacing=1.0)
        p.add_run(label + " ").bold = True
        p.add_run(value or blank)
        para(doc, spacing=1.0)

    line(t("executed_at", lang), a.get("executed_at") or university_name(w, lang))
    line(t("supervisor", lang) + ":", _person(w.get("supervisor") or {}))
    opp = a.get("opponents") or []
    line(t("opponents", lang), "; ".join(opp) if opp else blank)
    line(t("lead_org", lang), a.get("lead_organization", ""))
    para(doc, t("defense", lang, date=d.get("date") or "«___» ____________ 20___ г.",
                council=d.get("council") or blank, org=d.get("organization") or blank,
                address=d.get("address") or blank), spacing=1.0)
    para(doc, spacing=1.0)
    para(doc, t("library", lang, org=d.get("organization") or blank), spacing=1.0)
    para(doc, spacing=1.0)
    para(doc, t("sent_out", lang), spacing=1.0)
    _blank(doc, 2)
    para(doc, t("secretary", lang) + f",\n{blank}", indent=False, align="left", spacing=1.0)
    page_break(doc)


def header_article(doc, w: dict, lang: str, ctx: Ctx) -> None:
    """Статья, тезисы, доклад на конференцию: УДК, заглавие, авторы, аннотации, ключевые слова."""
    if w.get("udc"):
        para(doc, f'{t("udc", lang)} {w["udc"]}', align="left", indent=False)
    para(doc, (w.get("title") or "").upper(), bold=True, align="center")
    for i, a in enumerate(w.get("authors") or [], 1):
        sup = f" {i}" if len(w.get("authors") or []) > 1 else ""
        para(doc, f'{a.get("name", "")}{"," if a.get("degree") else ""} {a.get("degree", "")}'.strip() + sup,
             bold=True, align="center", spacing=1.0)
        line = ", ".join(x for x in (a.get("affiliation"), a.get("city")) if x)
        if line:
            para(doc, line, italic=True, align="center", spacing=1.0)
        if a.get("email"):
            para(doc, a["email"], italic=True, align="center", spacing=1.0)
    para(doc)
    abstracts = w.get("abstract") or {}
    kws = w.get("keywords") or {}
    langs = [lang] + [x for x in ("ru", "ky", "en") if x != lang]
    for lg in langs:
        if not abstracts.get(lg) and not kws.get(lg):
            continue
        if lg != lang and w.get(f"title_{lg}"):
            para(doc, w[f"title_{lg}"].upper(), bold=True, align="center")
        if abstracts.get(lg):
            p = doc.add_paragraph()
            p.add_run(t("abstract", lg) + ". ").bold = True
            add_inline(p, abstracts[lg], ctx)
        if kws.get(lg):
            p = doc.add_paragraph()
            p.add_run(t("keywords", lg) + ": ").bold = True
            p.add_run(", ".join(kws[lg]) + ".")
        para(doc)


# ---------- текст ----------

def section_title(item: dict, lang: str) -> str:
    kind = item.get("kind", "chapter")
    if item.get("title"):
        return item["title"]
    return {"intro": t("intro", lang), "conclusion": t("conclusion", lang)}.get(kind, "")


def read_text(w: dict, item: dict) -> str | None:
    f = item.get("file")
    if not f:
        return None
    p = w["_dir"] / f
    return p.read_text(encoding="utf-8") if p.exists() else None


def skeleton(item: dict) -> str:
    """Глава ещё не написана: план как подзаголовки + метка, без выдуманного текста."""
    lines = []
    for pt in item.get("plan") or []:
        name = pt.split(" ", 1)[1] if pt.split(" ", 1)[0].replace(".", "").isdigit() and " " in pt else pt
        lines += [f"## {name}", "", "[ТЕКСТ НЕ НАПИСАН]", ""]
    return "\n".join(lines) or "[ТЕКСТ НЕ НАПИСАН]"


def build(work_yaml: str | Path, part: str | None = None) -> tuple[Path, list[str], dict]:
    w = load_work(work_yaml)
    lang = w.get("language", "ru")
    typ = w["type"]
    is_avto = typ == "avtoreferat" or part == "avtoreferat"
    if part == "avtoreferat" and typ != "kandidat":
        raise DataError("автореферат собирается из паспорта кандидатской (type: kandidat)")
    structure = (w.get("avtoreferat") or {}).get("structure") if (is_avto and typ == "kandidat") \
        else w.get("structure")
    if not structure:
        raise DataError("в work.yaml нет structure — нечего собирать")
    style = w["citation_style"]
    order = w.get("bib_order") or ("cite" if style == "ieee" else "alpha")
    numbering = w.get("numbering") or ("chapter" if typ in DISSERTATIONS else "through")
    short = typ in SHORT_TYPES and (typ != "doklad" or w.get("conference"))
    diss = typ in DISSERTATIONS or is_avto

    sources = load_sources(w)
    by_key = {s.get("key"): s for s in sources}
    texts = [(item, read_text(w, item)) for item in structure]
    cite_order: list[str] = []
    for _, txt in texts:
        for k in scan_citations(txt or ""):
            if k not in cite_order:
                cite_order.append(k)
    verified = [s for s in sources if s.get("verified") is True]
    bib = sort_entries(verified, order, cite_order)
    if is_avto:
        pub_keys = (w.get("avtoreferat") or {}).get("publications") or []
        bib = [by_key[k] for k in pub_keys if by_key.get(k, {}).get("verified") is True]
    num = {s["key"]: i for i, s in enumerate(bib, 1)}
    warnings: list[str] = []

    def cite_fn(inner: str) -> str:
        refs, bad = [], []
        for key, loc in parse_cite(inner):
            e = by_key.get(key)
            if e is None:
                bad.append(f"[НЕТ В СПИСКЕ: {key}]")
            elif e.get("verified") is not True:
                bad.append(f"[НЕ ПРОВЕРЕНО: {key}]")
            elif key not in num:
                bad.append(f"[НЕТ В СПИСКЕ: {key}]")
            else:
                refs.append((e, num[key], loc))
        return (cite(refs, style, lang) if refs else "") + (" " if refs and bad else "") + " ".join(bad)

    labels = {"figure": t("figure", lang), "table": t("table", lang)}
    ctx = Ctx(base_dir=w["_dir"], lang=lang, numbering=numbering, cite_fn=cite_fn, labels=labels,
              diss_style=diss)

    doc = new_document(first_page_number=short)
    doc.core_properties.author = ", ".join(a.get("name", "") for a in w.get("authors") or [])
    doc.core_properties.title = w.get("title", "")
    if is_avto:
        title_kandidat(doc, w, lang, avtoreferat=True)
        avtoreferat_back(doc, w, lang)
    elif typ == "kandidat":
        title_kandidat(doc, w, lang)
    elif short:
        header_article(doc, w, lang, ctx)
    else:
        title_student(doc, w, lang)
    if not short and not is_avto:
        toc(doc, t("toc_diss" if diss else "toc", lang), t("toc_hint", lang))

    appendices = [(it, tx) for it, tx in texts if it.get("kind") == "appendix"]
    chapter_no = 0
    first = True
    for item, txt in texts:
        kind = item.get("kind", "chapter")
        if kind == "appendix":
            continue
        title = section_title(item, lang)
        if txt is None:
            warnings.append(f"нет файла {item.get('file') or '(file не указан)'} — «{title}» собран как каркас")
            txt = skeleton(item)
        brk = not short and not first
        if kind == "chapter":
            chapter_no += 1
            ctx.start_chapter(chapter_no, numbered=True)
            label = (f'{t("chapter", lang)} {chapter_no}. {title.upper()}' if diss
                     else f"{chapter_no} {title}")
            heading(doc, label, 1, page_break=brk)
        else:
            ctx.start_chapter(0, numbered=False)
            if title:
                if short:
                    heading(doc, title, 2)
                else:
                    heading(doc, title.upper(), 1, page_break=brk)
        render_markdown(doc, txt, ctx)
        first = False

    ai = w.get("ai_disclosure") or {}
    if ai.get("include"):
        heading(doc, t("ai_title", lang), 2 if short else 1, page_break=not short)
        para(doc, ai.get("statement") or t("ai_default", lang, tools=ai.get("tools", "—"), how=ai.get("how", "—")))

    if is_avto:
        for k in (w.get("avtoreferat") or {}).get("publications") or []:
            if by_key.get(k, {}).get("verified") is not True:
                warnings.append(f"публикация {k} не найдена или не проверена — в список не вошла")
        heading(doc, t("publications", lang), 1, page_break=False)
        for i, e in enumerate(bib, 1):
            p = doc.add_paragraph()
            p.add_run(f"{i}. ")
            for txt_, it in format_entry(e, style):
                p.add_run(txt_).italic = it
        for lg in ("ky", "ru", "en"):
            ab = (w.get("abstract") or {}).get(lg)
            if not ab:
                continue
            heading(doc, t("summary", lg), 1, page_break=True)
            if w.get(f"title_{lg}") or lg == lang:
                para(doc, w.get(f"title_{lg}") or w.get("title", ""), bold=True, align="center")
            kws = (w.get("keywords") or {}).get(lg)
            if kws:
                p = doc.add_paragraph()
                p.add_run(t("keywords", lg) + ": ").bold = True
                p.add_run(", ".join(kws) + ".")
            add_inline(doc.add_paragraph(), ab, ctx)
    elif bib or cite_order:
        title = t("bibliography_short" if short else "bibliography", lang)
        heading(doc, title, 2 if short else 1, page_break=not short)
        for i, e in enumerate(bib, 1):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            if style == "apa":
                pf.first_line_indent = Cm(-1.25)
                pf.left_indent = Cm(1.25)
            elif style == "ieee":
                p.add_run(f"[{i}] ")
            else:
                p.add_run(f"{i}. ")
            for txt_, it in format_entry(e, style):
                p.add_run(txt_).italic = it
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for n, (item, txt) in enumerate(appendices):
        letters = "АБВГДЕЖИКЛМНПРСТУФХЦШЩЭЮЯ" if lang != "en" else "ABCDEFGHJKLMNPQRSTUVWXYZ"
        title = section_title(item, lang)
        h = heading(doc, f'{t("appendix", lang)} {letters[n % len(letters)]}', 1, page_break=True)
        if title:
            para(doc, title, bold=True, align="center")
        ctx.start_chapter(0, numbered=False)
        render_markdown(doc, txt if txt is not None else "[ТЕКСТ НЕ НАПИСАН]", ctx)
        h.paragraph_format.keep_with_next = True

    unverified = [s.get("key") for s in sources if s.get("verified") is not True]
    c = ctx.counts
    if c["data"]:
        warnings.append(f"меток «нужны ваши данные»: {c['data']} — без ваших данных эти места не заполнить")
    if c["source"]:
        warnings.append(f"утверждений без источника: {c['source']}")
    if c["draft"]:
        warnings.append(f"ненаписанных разделов: {c['draft']}")
    if c["unverified"]:
        warnings.append(f"ссылок на непроверенные источники: {c['unverified']} (в тексте [НЕ ПРОВЕРЕНО])")
    if c["unknown"]:
        warnings.append(f"ссылок на ключи, которых нет в sources.yaml: {c['unknown']}")
    if c["missing_fig"]:
        warnings.append(f"не найдено рисунков: {c['missing_fig']}")
    if unverified:
        warnings.append(f"не проверены и не вошли в список: {', '.join(map(str, unverified))} — verify_sources.py")

    name = "avtoreferat" if is_avto else typ
    out = out_dir(w) / f"{name}.docx"
    doc.save(str(out))
    return out, warnings, c


def main() -> None:
    utf8_stdout()
    ap = argparse.ArgumentParser(description="Сборка научной работы в docx по ГОСТ")
    ap.add_argument("work", help="путь к work.yaml")
    ap.add_argument("--part", choices=("avtoreferat",), help="собрать автореферат кандидатской")
    ap.add_argument("--strict", action="store_true", help="код ошибки, если остались метки или непроверенные ссылки")
    args = ap.parse_args()
    try:
        out, warnings, counts = build(args.work, args.part)
    except DataError as e:
        fail(str(e))
    print(f"Готово: {out}")
    for wmsg in warnings:
        print(f"  ! {wmsg}")
    if args.strict and any(counts[k] for k in ("data", "source", "draft", "unverified", "unknown")):
        sys.exit(2)


if __name__ == "__main__":
    main()
