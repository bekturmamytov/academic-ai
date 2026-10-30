"""docx по ГОСТ + разметка глав (markdown) → docx. Общий для build_docx и md_docx.

Оформление скопировано из academic-core/scripts/common.py и дополнено: оглавление,
подписи рисунков и таблиц, ссылки [@key], подсветка меток.
Зависимость: python-docx (объявляет вызывающий скрипт).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

MARKS = r"\[(?:НУЖНЫ ВАШИ ДАННЫЕ|НУЖЕН ИСТОЧНИК|ТЕКСТ НЕ НАПИСАН|НЕ ПРОВЕРЕНО|НЕТ В СПИСКЕ)[^\]]*\]"
MARK_RE = re.compile(MARKS, re.I)
CITE_RE = re.compile(r"\[(@[^\]]+)\]")
INLINE_RE = re.compile(
    r"(?P<cite>\[@[^\]]+\])|(?P<mark>(?i:" + MARKS + r"))"
    r"|\*\*(?P<b>.+?)\*\*|(?<![\w*])\*(?P<i>[^*\s][^*]*?)\*(?!\w)"
)
ALIGN = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
         "left": WD_ALIGN_PARAGRAPH.LEFT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


# ---------- документ ----------

def _field(run, instr: str, placeholder: str = "") -> None:
    def fc(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el

    run._r.append(fc("begin"))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run._r.append(it)
    run._r.append(fc("separate"))
    t = OxmlElement("w:t")
    t.text = placeholder
    run._r.append(t)
    run._r.append(fc("end"))


def _set_fonts(style, name: str = "Times New Roman") -> None:
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(a), None)


def new_document(*, first_page_number: bool = True):
    """Пустой docx: A4, Times New Roman 14, интервал 1,5, поля 30/15/20/20, номер страницы внизу по центру."""
    doc = Document()
    for s in doc.sections:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin, s.right_margin = Cm(3), Cm(1.5)
        s.top_margin, s.bottom_margin = Cm(2), Cm(2)
        s.different_first_page_header_footer = not first_page_number
        p = s.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        _field(p.add_run(), "PAGE", "1")

    # docDefaults: тема Calibri → Times
    rpr_default = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    fonts = rpr_default.find(qn("w:rFonts"))
    if fonts is not None:
        for a in list(fonts.attrib):
            del fonts.attrib[a]
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(a), "Times New Roman")

    st = doc.styles["Normal"]
    _set_fonts(st)
    st.font.size = Pt(14)
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(1.25)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True
    for lvl in (1, 2, 3):
        h = doc.styles[f"Heading {lvl}"]
        _set_fonts(h)
        h.font.size = Pt(14)
        h.font.bold = True
        h.font.italic = False
        h.font.color.rgb = None
        hp = h.paragraph_format
        hp.first_line_indent = Cm(0) if lvl == 1 else Cm(1.25)
        hp.line_spacing = 1.5
        hp.space_before = Pt(12 if lvl > 1 else 0)
        hp.space_after = Pt(12 if lvl == 1 else 6)
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER if lvl == 1 else WD_ALIGN_PARAGRAPH.LEFT
        hp.keep_with_next = True
    for name in ("Title", "Caption"):
        if name in [s.name for s in doc.styles]:
            _set_fonts(doc.styles[name])

    cp = doc.core_properties  # шаблон python-docx подписывает себя — документ принадлежит автору
    cp.author = cp.last_modified_by = cp.comments = cp.title = cp.subject = cp.keywords = ""

    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)
    return doc


def para(doc, text: str = "", *, bold=False, italic=False, align=None, indent=True, size=None,
         highlight=False, keep_next=False, spacing=None):
    p = doc.add_paragraph()
    if text:
        lines = text.split("\n")
        for i, line in enumerate(lines):
            r = p.add_run(line)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
            if highlight:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if i < len(lines) - 1:
                r.add_break()
    if align:
        p.alignment = ALIGN[align]
    if not indent or align in ("center", "right"):
        p.paragraph_format.first_line_indent = Cm(0)
    if keep_next:
        p.paragraph_format.keep_with_next = True
    if spacing:
        p.paragraph_format.line_spacing = spacing
    return p


def heading(doc, text: str, level: int, *, page_break: bool = False):
    h = doc.add_heading(text, level=level)
    if page_break:
        h.paragraph_format.page_break_before = True
    return h


def page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def toc(doc, title: str, hint: str) -> None:
    """Заголовок + поле TOC. Word обновит при открытии (updateFields) или по F9."""
    p = para(doc, title, bold=True, align="center")
    p.paragraph_format.space_after = Pt(12)
    p2 = doc.add_paragraph()
    p2.paragraph_format.first_line_indent = Cm(0)
    _field(p2.add_run(), 'TOC \\o "1-3" \\h \\z \\u', hint)
    page_break(doc)


def table(doc, header: list[str], rows: list[list], widths_cm: list[float] | None = None):
    """Таблица 12 pt, интервал 1,0, шапка жирная."""
    tbl = doc.add_table(rows=1, cols=len(header))
    tbl.style = "Table Grid"
    for i, h in enumerate(header):
        tbl.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = tbl.add_row().cells
        for i, v in enumerate(row[: len(header)]):
            cells[i].text = "" if v is None else str(v)
    for ri, r in enumerate(tbl.rows):
        for ci, c in enumerate(r.cells):
            if widths_cm:
                c.width = Cm(widths_cm[ci])
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size = Pt(12)
                    run.bold = ri == 0
    return tbl


# ---------- markdown глав ----------

@dataclass
class Ctx:
    """Состояние сборки: нумерация, ссылки, счётчики меток."""
    base_dir: Path
    lang: str = "ru"
    numbering: str = "through"           # through | chapter
    chapter: int = 0                     # номер текущей главы (0 — ненумеруемый раздел)
    numbered: bool = False               # нумеровать ли подзаголовки ##/###
    diss_style: bool = False             # «1.1.» вместо «1.1»
    cite_fn: Callable[[str], str] | None = None
    labels: dict = field(default_factory=dict)
    fig: int = 0
    tab: int = 0
    sec: int = 0
    sub: int = 0
    counts: dict = field(default_factory=lambda: {"data": 0, "source": 0, "draft": 0,
                                                  "unverified": 0, "unknown": 0, "missing_fig": 0})

    def number(self, kind: str) -> str:
        if kind == "fig":
            self.fig += 1
            n = self.fig
        else:
            self.tab += 1
            n = self.tab
        return f"{self.chapter}.{n}" if self.numbering == "chapter" and self.chapter else str(n)

    def start_chapter(self, n: int, numbered: bool) -> None:
        self.chapter, self.numbered = n, numbered
        self.sec = self.sub = 0
        if self.numbering == "chapter":
            self.fig = self.tab = 0


def _count_mark(ctx: Ctx, text: str) -> None:
    t = text.upper()
    key = ("data" if "ДАННЫЕ" in t else "source" if "ИСТОЧНИК" in t else "draft" if "НАПИСАН" in t
           else "unknown" if "НЕТ В СПИСКЕ" in t else "unverified")
    ctx.counts[key] += 1


def add_inline(p, text: str, ctx: Ctx, *, bold=False, italic=False) -> None:
    """Текст абзаца с **жирным**, *курсивом*, ссылками [@key] и подсвеченными метками."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            r.bold, r.italic = bold, italic
        if m.group("cite"):
            s = ctx.cite_fn(m.group("cite")[1:-1]) if ctx.cite_fn else m.group("cite")
            for part in re.split(r"(\[(?:НЕ ПРОВЕРЕНО|НЕТ В СПИСКЕ)[^\]]*\])", s):
                if not part:
                    continue
                r = p.add_run(part)
                r.bold, r.italic = bold, italic
                if MARK_RE.fullmatch(part):
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    _count_mark(ctx, part)
        elif m.group("mark"):
            r = p.add_run(m.group("mark"))
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            _count_mark(ctx, m.group("mark"))
        elif m.group("b") is not None:
            r = p.add_run(m.group("b"))
            r.bold, r.italic = True, italic
        else:
            r = p.add_run(m.group("i"))
            r.bold, r.italic = bold, True
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        r.bold, r.italic = bold, italic


def _caption(doc, text: str, align: str, *, keep_next=False):
    p = para(doc, "", align=align, indent=False, keep_next=keep_next)
    p.add_run(text)
    p.paragraph_format.line_spacing = 1.5
    return p


def _sec_heading(doc, title: str, level: int, ctx: Ctx) -> None:
    if ctx.numbered and ctx.chapter:
        title = re.sub(r"^\d+(\.\d+)*\.?\s+", "", title)  # номер ставит сборка
        if level == 2:
            ctx.sec += 1
            ctx.sub = 0
            num = f"{ctx.chapter}.{ctx.sec}"
        else:
            ctx.sub += 1
            num = f"{ctx.chapter}.{ctx.sec}.{ctx.sub}"
        title = f"{num}{'.' if ctx.diss_style else ''} {title}"
    heading(doc, title, level)


def render_markdown(doc, text: str, ctx: Ctx) -> None:
    """Глава в markdown → абзацы docx. Заголовок главы (# …) берётся из work.yaml и здесь пропускается."""
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    lines = text.splitlines()
    i, buf = 0, []
    pending_table_caption = None

    def flush():
        if buf:
            p = doc.add_paragraph()
            add_inline(p, " ".join(s.strip() for s in buf), ctx)
            buf.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            flush()
            i += 1
            continue
        if s.startswith("```"):
            flush()
            j = i + 1
            while j < len(lines) and not lines[j].strip().startswith("```"):
                p = para(doc, lines[j], indent=False, align="left", size=12, spacing=1.0)
                for r in p.runs:
                    r.font.name = "Courier New"
                j += 1
            i = j + 1
            continue
        if s.startswith(">>"):
            flush()
            p = para(doc, "", align="right", spacing=1.0)
            add_inline(p, s[2:].strip(), ctx)
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", s)
        if m:
            flush()
            if len(m.group(1)) > 1:
                _sec_heading(doc, m.group(2).strip(), len(m.group(1)), ctx)
            i += 1
            continue
        m = re.match(r"^(?:Таблица|Table|Таблица №)\s*:\s*(.+)$", s, re.I)
        if m:
            flush()
            pending_table_caption = m.group(1).strip()
            i += 1
            continue
        if s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            num = ctx.number("tab")
            cap = f"{ctx.labels.get('table', 'Таблица')} {num}"
            if pending_table_caption:
                cap += f" – {pending_table_caption}"
            _caption(doc, cap, "left", keep_next=True)
            if rows:
                table(doc, rows[0], rows[1:])
                doc.add_paragraph()
            pending_table_caption = None
            continue
        m = re.match(r"^!\[(.*?)\]\((.+?)\)$", s)
        if m:
            flush()
            cap_text, src = m.group(1).strip(), m.group(2).strip()
            img = (ctx.base_dir / src)
            if img.exists():
                doc.add_picture(str(img), width=Cm(15))
                pic = doc.paragraphs[-1]
                pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic.paragraph_format.first_line_indent = Cm(0)
                pic.paragraph_format.keep_with_next = True
            else:
                para(doc, f"[РИСУНОК НЕ НАЙДЕН: {src}]", align="center", highlight=True)
                ctx.counts["missing_fig"] += 1
            num = ctx.number("fig")
            _caption(doc, f"{ctx.labels.get('figure', 'Рисунок')} {num}" + (f" – {cap_text}" if cap_text else ""),
                     "center")
            i += 1
            continue
        m = re.match(r"^([-*+]|\d+[.)])\s+(.*)$", s)
        if m:
            flush()
            marker = "–" if m.group(1) in "-*+" else m.group(1).rstrip(".)") + ")"
            p = doc.add_paragraph()
            add_inline(p, f"{marker} {m.group(2)}", ctx)
            i += 1
            continue
        buf.append(s)
        i += 1
    flush()


def scan_citations(text: str) -> list[str]:
    """Ключи источников в порядке первого упоминания."""
    keys = []
    for m in CITE_RE.finditer(re.sub(r"<!--.*?-->", "", text, flags=re.S)):
        for part in m.group(1).split(";"):
            km = re.match(r"\s*@([\w:.\-]+)", part)
            if km and km.group(1) not in keys:
                keys.append(km.group(1))
    return keys


def parse_cite(inner: str) -> list[tuple[str, str]]:
    """«@a, с. 5; @b» → [(a, «с. 5»), (b, «»)]."""
    out = []
    for part in inner.split(";"):
        m = re.match(r"\s*@([\w:.\-]+)\s*(?:,\s*(.+))?$", part.strip())
        if m:
            out.append((m.group(1), (m.group(2) or "").strip()))
    return out
