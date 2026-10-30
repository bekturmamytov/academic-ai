# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1", "pyyaml>=6"]
# ///
"""Текст работы студента из docx — для рецензии и отзыва.

    uv run read_doc.py "<работа>.docx" [--out "<работа>.md"]

Заголовки помечаются #, таблицы — строками | … |. Внизу — объём: слова,
примерно страниц, таблиц, рисунков. PDF читайте инструментом Read напрямую.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).parent))
from common import fail, utf8_stdout  # noqa: E402

WORDS_PER_PAGE = 300  # Times 14, интервал 1,5 — около 1800 знаков


def extract(path: Path) -> tuple[str, dict]:
    doc = Document(str(path))
    lines, words = [], 0
    body = doc.element.body
    paras = {p._p: p for p in doc.paragraphs}
    tables = {t._tbl: t for t in doc.tables}
    for el in body.iterchildren():
        if el in paras:
            p = paras[el]
            txt = p.text.strip()
            if not txt:
                continue
            words += len(txt.split())
            name = (p.style.name or "").lower()
            lvl = next((int(c) for c in name if c.isdigit()), 1) if name.startswith(("heading", "заголовок")) else 0
            lines.append(("#" * lvl + " " if lvl else "") + txt)
        elif el in tables:
            for r in tables[el].rows:
                cells = [c.text.strip().replace("\n", " ") for c in r.cells]
                words += sum(len(c.split()) for c in cells)
                lines.append("| " + " | ".join(cells) + " |")
            lines.append("")
    stats = {"words": words, "pages": round(words / WORDS_PER_PAGE), "tables": len(doc.tables),
             "images": len(doc.inline_shapes)}
    return "\n".join(lines), stats


def main() -> None:
    utf8_stdout()
    ap = argparse.ArgumentParser(description="Текст из docx для рецензии")
    ap.add_argument("file")
    ap.add_argument("--out", help="сохранить текст в файл .md")
    args = ap.parse_args()
    path = Path(args.file).expanduser().resolve()
    if path.suffix.lower() != ".docx" or not path.exists():
        fail(f"нужен существующий .docx (PDF откройте через Read): {path}")
    text, st = extract(path)
    summary = (f"\n---\nОбъём: {st['words']} слов, ≈ {st['pages']} стр. (Times 14, 1,5), "
               f"таблиц {st['tables']}, рисунков {st['images']}")
    if args.out:
        Path(args.out).write_text(text + summary, encoding="utf-8")
        print(f"Текст: {Path(args.out).resolve()}{summary}")
    else:
        print(text + summary)


if __name__ == "__main__":
    main()
