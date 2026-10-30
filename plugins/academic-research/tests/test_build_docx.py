from docx import Document

from build_docx import build
from conftest import set_yaml


def all_text(path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join(parts)


def test_kandidat_from_example(workspace):
    out, warnings, counts = build(workspace / "work.yaml")
    assert out == workspace / "out" / "kandidat.docx" and out.exists()
    doc = Document(str(out))
    s = doc.sections[0]
    assert round(s.left_margin.cm, 1) == 3.0 and round(s.right_margin.cm, 1) == 1.5
    assert round(s.top_margin.cm, 1) == 2.0 and round(s.bottom_margin.cm, 1) == 2.0
    normal = doc.styles["Normal"]
    assert normal.font.name == "Times New Roman" and normal.font.size.pt == 14
    assert normal.paragraph_format.line_spacing == 1.5
    text = all_text(out)
    cp = doc.core_properties
    assert cp.author == "Иванов Иван Иванович" and "python-docx" not in (cp.comments or "")
    # титул кандидатской
    assert "На правах рукописи" in text
    assert "Диссертация на соискание ученой степени кандидата технических наук" in text
    assert "05.13.01 – Системный анализ" in text
    assert "Кыргызский государственный технический университет им. И. Раззакова" in text
    # оглавление — поле Word
    assert 'TOC \\o "1-3"' in doc.element.body.xml
    # главы, разделы, подписи — по главам
    assert "ГЛАВА 1. АНАЛИЗ МЕТОДОВ ВИБРОДИАГНОСТИКИ НАСОСНОГО ОБОРУДОВАНИЯ" in text
    assert "1.1. Виды отказов центробежных насосов" in text
    assert "2.1.1. Требования к данным" in text
    assert "Таблица 1.1 – Группы отказов и их признаки" in text
    assert "Рисунок 1.1 – Схема стенда для замеров" in text
    assert len(doc.inline_shapes) == 1
    # ссылки и список: по алфавиту he2016 (заглавие «Deep…») = 1, lecun2015 = 2
    assert "[1, с. 771]" in text and "[1; 2]" in text
    assert "СПИСОК ИСПОЛЬЗОВАННОЙ ЛИТЕРАТУРЫ" in text
    assert "2. LeCun, Y. Deep learning / Y. LeCun, Y. Bengio, G. Hinton // Nature." in text
    # непроверенный источник не попал в список
    assert "Основы вибродиагностики" not in text
    assert any("ivanov2020" in w for w in warnings)
    assert "ПРИЛОЖЕНИЕ А" in text
    assert counts["data"] == counts["unverified"] == counts["unknown"] == 0


def test_unverified_and_unknown_citations_are_marked(workspace):
    (workspace / "chapters" / "02.md").write_text(
        "## Обзор\n\nПо данным учебника [@ivanov2020], а также [@nobody2001].\n\n"
        "Средняя наработка на отказ составила [НУЖНЫ ВАШИ ДАННЫЕ: наработка по журналу ремонтов].\n",
        encoding="utf-8")
    out, warnings, counts = build(workspace / "work.yaml")
    text = all_text(out)
    assert "[НЕ ПРОВЕРЕНО: ivanov2020]" in text
    assert "[НЕТ В СПИСКЕ: nobody2001]" in text
    assert counts["unverified"] == 1 and counts["unknown"] == 1 and counts["data"] == 1
    assert any("нужны ваши данные" in w for w in warnings)


def test_missing_chapter_becomes_skeleton(workspace):
    (workspace / "chapters" / "02.md").unlink()
    out, warnings, counts = build(workspace / "work.yaml")
    assert counts["draft"] >= 1
    assert any("02.md" in w for w in warnings)
    assert "[ТЕКСТ НЕ НАПИСАН]" in all_text(out)


def test_statya_header_no_title_page(workspace):
    set_yaml(workspace / "work.yaml", type="statya", udc="621.67", numbering="through",
             abstract={"ru": "В статье рассмотрена постановка задачи.", "en": "The paper states the problem."},
             structure=[{"kind": "section", "title": "Введение", "file": "chapters/00-intro.md"},
                        {"kind": "section", "title": "Методы", "file": "chapters/01.md"},
                        {"kind": "section", "title": "Выводы", "file": "chapters/09-conclusion.md"}])
    out, _, _ = build(workspace / "work.yaml")
    assert out.name == "statya.docx"
    doc = Document(str(out))
    text = all_text(out)
    assert doc.paragraphs[0].text == "УДК 621.67"
    assert "На правах рукописи" not in text and "TOC \\o" not in doc.element.body.xml
    assert "Аннотация. В статье рассмотрена постановка задачи." in text
    assert "Abstract. The paper states the problem." in text
    assert "Ключевые слова: вибродиагностика" in text
    assert "Список литературы" in text
    assert "Рисунок 1 – Схема стенда для замеров" in text


def test_kursovaya_ieee_through_numbering(workspace):
    set_yaml(workspace / "work.yaml", type="kursovaya", citation_style="ieee", numbering="through",
             bib_order="cite", discipline="Техническая диагностика")
    out, _, _ = build(workspace / "work.yaml")
    text = all_text(out)
    assert "КУРСОВАЯ РАБОТА" in text and "по дисциплине «Техническая диагностика»" in text
    assert "СОДЕРЖАНИЕ" in text
    assert "1 Анализ методов вибродиагностики насосного оборудования" in text
    assert "Рисунок 1 – Схема стенда" in text
    assert "между агрегатами [1], [2]" in text and "[2, с. 771]" in text  # по порядку ссылок: lecun первый
    assert "[1] Y. LeCun, Y. Bengio, and G. Hinton" in text


def test_avtoreferat_part(workspace):
    set_yaml(workspace / "work.yaml", abstract={"ru": "Резюме работы.", "ky": "Иштин резюмеси.", "en": "Summary."})
    out, warnings, _ = build(workspace / "work.yaml", part="avtoreferat")
    text = all_text(out)
    assert out.name == "avtoreferat.docx"
    assert "АВТОРЕФЕРАТ" in text and "Официальные оппоненты:" in text
    assert "СПИСОК ОПУБЛИКОВАННЫХ РАБОТ ПО ТЕМЕ ДИССЕРТАЦИИ" in text
    assert text.count("РЕЗЮМЕ") == 2 and "SUMMARY" in text


def test_md_docx(tmp_path):
    from md_docx import build as md_build
    md = tmp_path / "rubric.md"
    md.write_text("# КРИТЕРИИ ОЦЕНИВАНИЯ КУРСОВОЙ РАБОТЫ\n\nТаблица: Критерии\n\n| Критерий | Баллы |\n|---|---|\n"
                  "| Актуальность | 10 |\n\n## Шкала\n\n- 87–100 — «отлично»;\n", encoding="utf-8")
    out = md_build(md, org=False)
    text = all_text(out)
    assert "КРИТЕРИИ ОЦЕНИВАНИЯ КУРСОВОЙ РАБОТЫ" in text and "Таблица 1 – Критерии" in text
    assert "– 87–100 — «отлично»;" in text
