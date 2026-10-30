from docx import Document
from docx.shared import Cm, Pt

from build_docx import build
from check_format import check, fix

LONG = ("Насосное оборудование работает в системах водоснабжения непрерывно, и внезапный отказ "
        "останавливает весь участок, поэтому раннее обнаружение неисправностей остаётся задачей. ")


def bad_docx(path):
    """Типичная «плохая» работа: шаблон Word, поля 2 см, Arial 12, одинарный интервал, метка осталась."""
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2)
    doc.add_paragraph("Введение")
    for i in range(12):
        p = doc.add_paragraph(LONG * 2)
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(12)
    doc.add_paragraph(LONG + "[НУЖНЫ ВАШИ ДАННЫЕ: результаты опроса]")
    doc.add_paragraph("Список литературы")
    doc.add_paragraph("Иванов И.И. Основы вибродиагностики. Бишкек, 2020.")
    doc.add_paragraph("Сайт статистики https://stat.gov.kg")
    doc.save(str(path))
    return path


def areas(probs, level="critical"):
    return {p.area for p in probs if p.level == level}


def test_built_docx_has_no_critical(workspace):
    out, _, _ = build(workspace / "work.yaml")
    probs, stats = check(out)
    assert areas(probs) == set(), [p.text for p in probs]
    assert stats["references"] == 2 and stats["images"] == 1 and stats["tables"] == 1


def test_bad_docx_problems_found(tmp_path):
    probs, _ = check(bad_docx(tmp_path / "bad.docx"))
    crit = areas(probs)
    assert {"поля", "номера страниц", "шрифт", "кегль", "интервал", "метки"} <= crit
    warn_text = " ".join(p.text for p in probs if p.level == "warning")
    assert "ГОСТ Р 7.0.100" in warn_text and "дата обращения" in warn_text


def test_fix_makes_safe_corrections(tmp_path):
    src = bad_docx(tmp_path / "bad.docx")
    out = fix(src)
    assert out.name == "bad_исправлено.docx" and src.exists()
    probs, _ = check(out)
    assert areas(probs) == {"метки"}  # метку автор убирает сам — данные не выдумываем
