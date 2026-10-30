"""Скрипты учебных материалов: лекции, слайды, СРС, УМКД, пособие, отчёты.

uv run --with pytest --with python-docx --with python-pptx --with openpyxl --with pyyaml \
    pytest plugins/academic-core/tests/test_content.py -q

Скрипты запускаются как у пользователя — `uv run <скрипт>` (зависимости из PEP 723).
"""
from __future__ import annotations

import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
FIX = Path(__file__).resolve().parent / "fixtures"
sys.path.insert(0, str(SCRIPTS))


def cmd(script: str) -> list[str]:
    uv = shutil.which("uv")
    return [uv, "run", "--quiet", str(SCRIPTS / script)] if uv else [sys.executable, str(SCRIPTS / script)]


def run(script: str, *args, cwd: Path) -> subprocess.CompletedProcess:
    r = subprocess.run([*cmd(script), *map(str, args)], cwd=cwd, capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0, f"{script} упал:\n{r.stdout}\n{r.stderr}"
    return r


@pytest.fixture()
def ws(tmp_path: Path) -> Path:
    """Рабочая папка: profile.yaml, university/ с логотипом, курс с лекцией, СРС и слайдами."""
    (tmp_path / "profile.yaml").write_text(
        'name: "Асанова Айгуль Бакытовна"\ndegree: "к.т.н., доцент"\nposition: "Доцент"\n'
        'department: "Программное обеспечение"\ncity: "Бишкек"\nemail: "a@example.kg"\nlanguage: ru\n',
        encoding="utf-8")
    uni = tmp_path / "university"
    uni.mkdir()
    from PIL import Image

    Image.new("RGB", (120, 120), (31, 78, 121)).save(uni / "logo.png")
    (uni / "brand.yaml").write_text('name: "Тестовый университет"\ncolor: "#7A1F3D"\nlogo: logo.png\n', encoding="utf-8")
    course = tmp_path / "course"
    course.mkdir()
    shutil.copyfile(ROOT / "templates" / "course.example.yaml", course / "course.yaml")
    shutil.copytree(FIX / "lectures", course / "lectures")
    shutil.copytree(FIX / "slides", course / "slides")
    shutil.copytree(FIX / "manual", course / "manual")
    shutil.copytree(FIX / "report", course / "report")
    shutil.copyfile(FIX / "srs.yaml", course / "srs.yaml")
    return tmp_path


def docx_text(path: Path) -> str:
    from docx import Document

    d = Document(str(path))
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    return "\n".join([p.text for p in d.paragraphs] + cells)


def test_markdown_parser():
    from content_kit import parse_md, runs_text

    blocks = parse_md("# Заголовок\n\nТекст **жирный** и *курсив*.\n\n- пункт\n  - вложенный\n1. первый\n\n"
                      "| A | B |\n|---|---|\n| 1 | 2 |\n")
    kinds = [b["t"] for b in blocks]
    assert kinds == ["h", "p", "ul", "ol", "table"]
    assert any(r.get("b") for r in blocks[1]["runs"]) and any(r.get("i") for r in blocks[1]["runs"])
    assert blocks[2]["items"][1]["level"] == 1
    assert runs_text(blocks[4]["rows"][0][1]) == "2"


def test_literature_status():
    from content_kit import literature_status

    data = {"literature": {"main": [{"text": "Дейт К. Дж. Введение в системы баз данных. 8-е изд.", "verified": True}],
                           "extra": [{"text": "Сомнительная книга. — 2020.", "verified": False}]}}
    assert literature_status("Дейт К. Дж. Введение в системы баз данных. 8-е изд.", data) == "ok"
    assert literature_status("Сомнительная книга. — 2020.", data) == "unverified"
    assert literature_status("Совсем другой источник о химии", data) == "absent"


def test_kr_scale():
    from content_kit import kr_grade

    assert kr_grade(87)[0] == "A" and kr_grade(86.5)[0] == "B" and kr_grade(61)[0] == "E" and kr_grade(60.9)[0] == "F"


def test_lecture_single_and_book(ws: Path):
    c = ws / "course"
    r = run("build_lecture.py", c / "lectures" / "01-vvedenie-v-bd.md", cwd=ws)
    out = c / "out" / "lectures" / "01.docx"
    assert out.exists()
    assert "нет в паспорте" in r.stdout  # выдуманный источник помечен
    text = docx_text(out)
    assert "[НЕ ПРОВЕРЕНО]" in text and "Ключевые понятия" in text
    run("build_lecture.py", c / "course.yaml", cwd=ws)
    book = c / "out" / "lectures" / "lectures.docx"
    assert book.exists()
    assert "TOC \\o" in zipfile.ZipFile(book).read("word/document.xml").decode()


def test_slides(ws: Path):
    from pptx import Presentation

    c = ws / "course"
    run("build_slides.py", c / "slides" / "01-slides.yaml", cwd=ws)
    out = c / "out" / "slides" / "01.pptx"
    assert out.exists()
    prs = Presentation(str(out))
    assert len(prs.slides) == 8  # титул + 6 + вопросы
    assert round(prs.slide_width / prs.slide_height, 2) == 1.78
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides]
    assert all(n.strip() for n in notes[1:-1])
    texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "Асанова" in texts and "Вопросы для обсуждения" in texts
    import yaml

    media = yaml.safe_load((c / "out" / "slides" / "01.yaml").read_text(encoding="utf-8"))
    assert media["accent"] == "#7A1F3D" and media["lang"] == "ru"
    kinds = [s["type"] for s in media["slides"]]
    assert kinds[0] == "title" and kinds[-1] == "closing" and {"section", "bullets", "table"} <= set(kinds)
    assert all(isinstance(b, str) and "**" not in b for s in media["slides"] for b in s.get("bullets", []))
    assert all(s["notes"] for s in media["slides"][1:-1])


def test_srs(ws: Path):
    c = ws / "course"
    run("build_srs.py", c / "srs.yaml", cwd=ws)
    out = c / "out" / "srs.docx"
    assert out.exists()
    text = docx_text(out)
    assert "Критерии оценивания" in text and "Максимум баллов: 10" in text


def test_umkd(ws: Path):
    c = ws / "course"
    run("build_srs.py", c / "srs.yaml", cwd=ws)
    shutil.copyfile(c / "out" / "srs.docx", c / "out" / "syllabus_ru.docx")  # заглушка силабуса для слияния
    r = run("build_umkd.py", c / "course.yaml", cwd=ws)
    out = c / "out" / "umkd.docx"
    assert out.exists()
    assert "Тестовые задания: нет" in r.stdout and "/academic-core:tests" in r.stdout
    assert "Силлабус: готово" in r.stdout
    text = docx_text(out)
    for section in ("Рабочая программа дисциплины", "Конспект лекций", "Глоссарий", "Лист согласования",
                    "Вопросы к экзамену", "Транзакция"):
        assert section in text, section


def test_manual(ws: Path):
    c = ws / "course"
    if not shutil.which("typst"):
        run("build_manual.py", c / "manual" / "manual.yaml", "--no-pdf", cwd=ws)
    else:
        run("build_manual.py", c / "manual" / "manual.yaml", cwd=ws)
        pdf = c / "out" / "manual.pdf"
        assert pdf.exists() and pdf.read_bytes()[:4] == b"%PDF"
    assert (c / "out" / "manual.docx").exists()
    assert "ПРОЕКТИРОВАНИЕ БАЗ ДАННЫХ" in docx_text(c / "out" / "manual.docx")


def test_report_plan(ws: Path):
    c = ws / "course"
    r = run("build_report.py", c / "report" / "plan.yaml", cwd=ws)
    assert (c / "out" / "report_plan_report.docx").exists()
    assert "нет отметки о выполнении" in r.stdout


def test_report_course(ws: Path):
    c = ws / "course"
    r = run("build_report.py", c / "report" / "course_report.yaml", cwd=ws)
    assert (c / "out" / "report_course_report.docx").exists()
    assert "graded=5" in r.stdout and "avg=75.5" in r.stdout and "pass=80.0" in r.stdout and "quality=60.0" in r.stdout


def test_course_report_requires_data(ws: Path):
    c = ws / "course"
    (c / "report" / "empty.yaml").write_text("type: course_report\ncourse: ../course.yaml\n", encoding="utf-8")
    r = subprocess.run([*cmd("build_report.py"), str(c / "report" / "empty.yaml")],
                       capture_output=True, text=True, encoding="utf-8")
    assert r.returncode != 0 and "файл оценок" in r.stderr
