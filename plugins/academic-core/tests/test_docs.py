"""Скрипты паспорта курса, силлабуса, РУП и тестов на примере «Базы данных».

    uv run --with pytest --with python-docx --with openpyxl --with pyyaml pytest plugins/academic-core/tests -q
"""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
import yaml

PLUGIN = Path(__file__).resolve().parents[1]
SCRIPTS = PLUGIN / "scripts"
EXAMPLE = PLUGIN / "templates" / "course.example.yaml"
TESTS_EXAMPLE = PLUGIN / "skills" / "tests" / "templates" / "tests.example.yaml"
sys.path.insert(0, str(SCRIPTS))

from common import check_hours, load_course  # noqa: E402
from course_check import bloom_verb, validate_course  # noqa: E402


def run(script: str, *args, ok: bool = True) -> subprocess.CompletedProcess:
    res = subprocess.run([sys.executable, str(SCRIPTS / script), *map(str, args)],
                         capture_output=True, text=True, encoding="utf-8")
    if ok:
        assert res.returncode == 0, res.stdout + res.stderr
    return res


@pytest.fixture()
def course_dir(tmp_path: Path) -> Path:
    d = tmp_path / "Мои курсы" / "ИС-201 Базы данных"
    d.mkdir(parents=True)
    shutil.copy(EXAMPLE, d / "course.yaml")
    shutil.copy(TESTS_EXAMPLE, d / "tests.yaml")
    (d.parent / "profile.yaml").write_text(yaml.safe_dump({
        "name": "Иванова Анна Сергеевна", "degree": "к.т.н.", "position": "доцент",
        "department": "Программная инженерия", "faculty": "Факультет ИТ",
        "university": "Тестовый университет", "city": "Бишкек", "language": "ru",
    }, allow_unicode=True), encoding="utf-8")
    return d


def edit(path: Path, fn) -> None:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    fn(data)
    path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")


def docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


# ---------- паспорт ----------

def test_example_hours_and_schema(course_dir):
    data = load_course(course_dir / "course.yaml")
    assert check_hours(data) == []
    errors, warnings = validate_course(data)
    assert errors == []
    assert len(data["topics"]) == data["weeks"] == 15
    assert data["teacher"]["name"] == "Иванова Анна Сергеевна"  # из profile.yaml


def test_validate_script_ok(course_dir):
    res = run("validate_course.py", course_dir / "course.yaml")
    assert "Ошибки" not in res.stdout


@pytest.mark.parametrize("mutate, needle", [
    (lambda d: d["hours"].__setitem__("srs", 61), "кр. × 30"),
    (lambda d: d["assessment"][0].__setitem__("weight", 25), "нужно 100"),
    (lambda d: d["outcomes"][0].__setitem__("text", "Знает реляционную модель"), "неизмеримый"),
    (lambda d: d["topics"][0].__setitem__("outcomes", ["РО9"]), "РО «РО9»"),
    (lambda d: d["topics"][0].__setitem__("labs", 3), "labs: по темам 31"),
])
def test_validate_catches_errors(course_dir, mutate, needle):
    edit(course_dir / "course.yaml", mutate)
    res = run("validate_course.py", course_dir / "course.yaml", ok=False)
    assert res.returncode == 1
    assert needle in res.stdout


@pytest.mark.parametrize("text, lang, level", [
    ("Объясняет реляционную модель", "ru", "understand"),
    ("Умеет проектировать базу данных", "ru", "create"),
    ("Designs a relational schema", "en", "create"),
    ("Реляциялык моделди түшүндүрөт", "ky", "understand"),
])
def test_bloom_verbs(text, lang, level):
    word, levels = bloom_verb(text, lang)
    assert levels and level in levels, (word, levels)


def test_bloom_non_measurable():
    assert bloom_verb("Знает основы SQL", "ru")[1] is None
    assert bloom_verb("Students will understand SQL", "en")[1] is None


# ---------- силлабус ----------

@pytest.mark.parametrize("lang", ["ru", "ky", "en"])
def test_syllabus(course_dir, lang):
    res = run("build_syllabus.py", course_dir / "course.yaml", "--lang", lang)
    out = course_dir / "out" / f"syllabus_{lang}.docx"
    assert out.exists() and str(out) in res.stdout
    text = docx_text(out)
    title = {"ru": "СИЛЛАБУС", "ky": "СИЛЛАБУС", "en": "SYLLABUS"}[lang]
    assert title in text
    assert "Иванова Анна Сергеевна" in text
    assert "150" in text  # всего часов
    for i in range(1, 14):
        assert re.search(rf"^{i}\. ", text, re.M), f"нет раздела {i}"
    assert "academicAI" not in text and "AI Lab" not in text


def test_syllabus_sections_ru(course_dir):
    run("build_syllabus.py", course_dir / "course.yaml")
    text = docx_text(course_dir / "out" / "syllabus_ru.docx")
    for s in ("Результаты обучения", "Тематический план", "балльно-рейтинговая", "Шкала оценок",
              "Политика курса", "Использование инструментов ИИ", "ЦУР 4", "Литература", "Протокол №"):
        assert s in text, s
    assert "[НЕ ПРОВЕРЕНО]" not in text


def test_syllabus_unverified_literature(course_dir):
    edit(course_dir / "course.yaml",
         lambda d: d["literature"]["extra"].append({"text": "Петров П. П. Книга. — 2020.", "verified": False}))
    res = run("build_syllabus.py", course_dir / "course.yaml")
    assert "НЕ ПРОВЕРЕНО" in res.stdout
    assert "[НЕ ПРОВЕРЕНО] Петров" in docx_text(course_dir / "out" / "syllabus_ru.docx")
    run("build_syllabus.py", course_dir / "course.yaml", "--final")
    assert "Петров" not in docx_text(course_dir / "out" / "syllabus_ru.docx")


def test_syllabus_stops_on_errors(course_dir):
    edit(course_dir / "course.yaml", lambda d: d.__setitem__("credits", 4))
    res = run("build_syllabus.py", course_dir / "course.yaml", ok=False)
    assert res.returncode == 1
    assert not (course_dir / "out" / "syllabus_ru.docx").exists()


# ---------- РУП ----------

def test_rup(course_dir):
    from openpyxl import load_workbook

    run("build_rup.py", course_dir / "course.yaml", "--docx")
    xlsx = course_dir / "out" / "rup.xlsx"
    assert xlsx.exists() and (course_dir / "out" / "rup_ru.docx").exists()
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Распределение часов", "Тематический план", "Календарный график контроля"]
    formulas = [c.value for ws in wb for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert any(f.startswith("=SUM(") for f in formulas)
    assert any("/30" in f for f in formulas)
    plan = wb["Тематический план"]
    weeks = [plan.cell(r, 1).value for r in range(5, 20)]
    assert weeks == list(range(1, 16))
    hours = wb["Распределение часов"]
    assert sum(hours.cell(r, 2).value for r in range(5, 10)) == 150
    text = docx_text(course_dir / "out" / "rup_ru.docx")
    for s in ("РАБОЧАЯ ПРОГРАММА ДИСЦИПЛИНЫ", "Место дисциплины", "Содержание дисциплины по темам",
              "Оценочные средства", "Лабораторная работа"):
        assert s in text, s


# ---------- тесты ----------

def test_tests_check(course_dir):
    res = run("build_tests.py", course_dir / "tests.yaml", "--check")
    assert "18 вопросов" in res.stdout
    assert "покрытие" not in res.stdout  # пример покрывает все темы и РО
    assert not (course_dir / "out").exists()


def test_tests_build(course_dir):
    from openpyxl import load_workbook

    run("build_tests.py", course_dir / "tests.yaml")
    out = course_dir / "out"
    for name in ("tests.xlsx", "tests_1.docx", "tests_2.docx", "tests_key.docx", "tests.gift", "tests_moodle.xml"):
        assert (out / name).exists(), name

    wb = load_workbook(out / "tests.xlsx")
    bank = wb["Банк вопросов"]
    assert bank.max_row == 19  # шапка + 18
    assert all(bank.cell(r, 11).value for r in range(2, 20)), "у каждого вопроса есть ключ"
    cov = [c.value for row in wb["Покрытие"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any(v.startswith("=COUNTIF(") for v in cov)

    # бланк без ответов, 15 вопросов, варианты различаются
    b1, b2 = docx_text(out / "tests_1.docx"), docx_text(out / "tests_2.docx")
    assert "Вариант 1" in b1 and "Названа гибкость" not in b1
    from docx import Document
    numbered = [p.text for p in Document(str(out / "tests_1.docx")).paragraphs if re.match(r"\d+\. ", p.text)]
    assert len(numbered) == 15
    assert b1.split("Вариант 1")[1] != b2.split("Вариант 2")[1]
    key = docx_text(out / "tests_key.docx")
    assert "Вариант 1" in key and "Вариант 2" in key and "Ключ ответов" in key


def gift_questions(text: str) -> list[str]:
    """Вопросы GIFT: блоки от ::имя:: до закрывающей неэкранированной }."""
    body = "\n".join(line for line in text.splitlines()
                     if not line.startswith("//") and not line.startswith("$CATEGORY"))
    blocks, depth, start = [], 0, None
    i = 0
    while i < len(body):
        ch = body[i]
        if ch == "\\":
            i += 2
            continue
        if ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            assert depth == 0, "вложенные скобки в GIFT"
            blocks.append(body[start:i + 1].strip())
            start = i + 1
        elif start is None:
            start = i
        i += 1
    assert depth == 0, "незакрытая скобка в GIFT"
    return blocks


def test_gift_structure(course_dir):
    run("build_tests.py", course_dir / "tests.yaml")
    gift = (course_dir / "out" / "tests.gift").read_text(encoding="utf-8")
    assert gift.count("$CATEGORY: $course$/top/") == 15
    blocks = gift_questions(gift)
    assert len(blocks) == 18
    for b in blocks:
        assert re.match(r"^::Q\d+::", b), b[:40]
    assert "1\\:1 -> " in gift  # двоеточие в соответствии экранировано
    assert re.search(r"\{TRUE\}|\{FALSE\}", gift)
    assert re.search(r"~%-?\d+(\.\d+)?%", gift)  # веса multiple
    assert "{}" in gift  # эссе


def test_gift_escape():
    from build_tests import gift_escape

    assert gift_escape("a=b {c} ~d #e :f \\g") == "a\\=b \\{c\\} \\~d \\#e \\:f \\\\g"


def test_moodle_xml(course_dir):
    run("build_tests.py", course_dir / "tests.yaml")
    root = ET.parse(course_dir / "out" / "tests_moodle.xml").getroot()
    qs = [q for q in root.findall("question") if q.get("type") != "category"]
    assert len(qs) == 18
    types = {q.get("type") for q in qs}
    assert types == {"multichoice", "truefalse", "matching", "shortanswer", "essay"}
    for q in qs:
        if q.get("type") == "multichoice":
            fr = [float(a.get("fraction")) for a in q.findall("answer")]
            assert len(fr) >= 4 and abs(sum(f for f in fr if f > 0) - 100) < 0.01


def test_tests_errors(course_dir):
    def bad(d):
        d["questions"][0]["options"] = ["СУБД", "ОС", "IDE"]
        d["questions"][1]["answer"] = "Нет такого"
        d["questions"][2]["text"] = d["questions"][3]["text"]
        d["questions"][4]["topic"] = 99
    edit(course_dir / "tests.yaml", bad)
    res = run("build_tests.py", course_dir / "tests.yaml", ok=False)
    assert res.returncode == 1
    for needle in ("нужно не меньше 4", "не совпадает", "дублирует", "недели 99"):
        assert needle in res.stdout, needle


def test_tests_scope(course_dir):
    edit(course_dir / "tests.yaml", lambda d: (d.__setitem__("topics", [1, 2, 3, 4, 5, 6, 7, 16]),
                                                d.__setitem__("questions", d["questions"][:3]),
                                                d.__setitem__("per_variant", 3)))
    res = run("build_tests.py", course_dir / "tests.yaml", "--check")
    assert "темам недель 3, 4, 5, 6, 7" in res.stdout
