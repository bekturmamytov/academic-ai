"""Настройки вуза university/settings.yaml: загрузка, проверка и как они меняют документы.

    uv run --with pytest --with python-docx --with openpyxl --with pyyaml pytest plugins/academic-core/tests -q
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import pytest
import yaml

PLUGIN = Path(__file__).resolve().parents[1]
SCRIPTS = PLUGIN / "scripts"
TEMPLATE = PLUGIN / "templates" / "settings.example.yaml"
sys.path.insert(0, str(SCRIPTS))

from common import check_hours, load_course  # noqa: E402
from settings import DEFAULTS, DataError, diff_from_defaults, load_settings  # noqa: E402
from test_content import run as run_uv  # noqa: E402
from test_docs import docx_text, run  # noqa: E402 — фикстуры course_dir и ws — в conftest.py


def write_settings(root: Path, data: dict | str) -> Path:
    uni = root / "university"
    uni.mkdir(exist_ok=True)
    text = data if isinstance(data, str) else yaml.safe_dump(data, allow_unicode=True, sort_keys=False)
    (uni / "settings.yaml").write_text(text, encoding="utf-8")
    return uni


# ---------- загрузчик ----------

def test_no_file_gives_kr_defaults(tmp_path):
    s = load_settings(tmp_path)
    assert s["_file"] is None and s["_errors"] == []
    assert diff_from_defaults(s) == []
    assert s["document"]["margins_mm"] == {"left": 30, "right": 15, "top": 20, "bottom": 20}
    assert s["grading"]["credit_hours"] == 30


def test_template_equals_defaults(tmp_path):
    """Образец settings.example.yaml без правок = нормы КР: ни ошибок, ни отличий."""
    uni = tmp_path / "university"
    uni.mkdir()
    shutil.copy(TEMPLATE, uni / "settings.yaml")
    s = load_settings(tmp_path)
    assert s["_errors"] == []
    assert diff_from_defaults(s) == []
    assert not [w for w in s["_warnings"] if "неизвестный" in w]


def test_core_and_research_copies_identical():
    research = PLUGIN.parent / "academic-research" / "scripts" / "settings.py"
    assert (SCRIPTS / "settings.py").read_bytes() == research.read_bytes(), \
        "settings.py в academic-core и academic-research должны совпадать — скопируйте файл"


def test_merge_partial_and_lang_dicts(tmp_path):
    write_settings(tmp_path, {
        "document": {"margins_mm": {"left": 25}, "font": ""},
        "approval": {"role": {"ru": "Декан факультета"}},
        "research": {"thesis_pages": {"bachelor": [50, 70]}},
    })
    s = load_settings(tmp_path / "курс")
    assert s["document"]["margins_mm"] == {"left": 25, "right": 15, "top": 20, "bottom": 20}
    assert s["document"]["font"] == "Times New Roman"  # пусто — по умолчанию
    assert s["approval"]["role"] == {"ru": "Декан факультета"}  # {ru, ky, en} заменяется целиком
    assert s["research"]["thesis_pages"] == {"bachelor": [50, 70], "master": [60, 90]}
    assert any("approval.role: нет текста на ky" in w for w in s["_warnings"])


@pytest.mark.parametrize("data, key", [
    ({"grading": {"scale": [{"min": 61, "max": 100, "grade": "зачёт"}, {"min": 0, "max": 59, "grade": "незачёт"}]}},
     "grading.scale"),
    ({"grading": {"scale": [{"min": 50, "max": 100, "grade": "зачёт"}, {"min": 0, "max": 60, "grade": "незачёт"}]}},
     "grading.scale"),
    ({"grading": {"assessment_template": [{"name": "Экзамен", "weight": 60, "type": "final"}]}},
     "grading.assessment_template"),
    ({"document": {"margins_mm": {"left": -5}}}, "document.margins_mm.left"),
    ({"document": {"page_numbers": "внизу"}}, "document.page_numbers"),
    ({"research": {"citation_style": "mla"}}, "research.citation_style"),
    ({"manual": {"paper": "a3"}}, "manual.paper"),
    ({"brand": {"color": "синий"}}, "brand.color"),
    ({"syllabus": {"hide_sections": ["sdg", "цели"]}}, "syllabus.hide_sections"),
    ({"tests": {"formats": ["pdf"]}}, "tests.formats"),
])
def test_validation_errors_name_the_key(tmp_path, data, key):
    write_settings(tmp_path, data)
    with pytest.raises(DataError) as e:
        load_settings(tmp_path)
    assert key in str(e.value)
    s = load_settings(tmp_path, strict=False)  # неверное значение заменено нормой КР
    assert any(err.startswith(key) for err in s["_errors"])


def test_scale_gap_message(tmp_path):
    write_settings(tmp_path, {"grading": {"scale": [
        {"min": 61, "max": 100, "grade": "зачёт"}, {"min": 0, "max": 59, "grade": "незачёт"}]}})
    s = load_settings(tmp_path, strict=False)
    assert "нет оценки для баллов 60" in s["_errors"][0]


def test_unknown_keys_warn(tmp_path):
    write_settings(tmp_path, {"documents": {"font": "Arial"}, "document": {"fonts": "Arial"}})
    s = load_settings(tmp_path)
    assert s["_errors"] == []
    assert {w.split(":")[0] for w in s["_warnings"]} >= {"documents", "document.fonts"}
    assert s["document"]["font"] == "Times New Roman"


def test_brand_yaml_legacy(tmp_path):
    uni = tmp_path / "university"
    uni.mkdir()
    (uni / "brand.yaml").write_text('name: "Старый вуз"\nname_ky: "Эски ЖОЖ"\ncolor: "#AA0000"\nlogo: logo.png\n',
                                    encoding="utf-8")
    (uni / "logo.png").write_bytes(b"")
    s = load_settings(tmp_path)
    assert s["university"]["name"] == "Старый вуз" and s["brand"]["color"] == "#AA0000"
    assert s["document"]["title_page_logo"] is True  # без settings.yaml логотип на титуле, как раньше
    write_settings(tmp_path, {"university": {"name": "", "short": "СВ"}, "brand": {"color": "#00AA00"}})
    s = load_settings(tmp_path)
    assert s["university"]["name"] == "Старый вуз"  # пусто в settings.yaml — берётся из brand.yaml
    assert s["brand"]["color"] == "#00AA00"  # settings.yaml главнее
    assert s["document"]["title_page_logo"] is False


# ---------- документы ----------

UNI = {
    "university": {"name": "Университет Настроек", "city": "Ош",
                   "ministry": {"ru": "Министерство образования и науки КР"}},
    "document": {"font": "Arial", "size": 12, "table_size": 10, "line_spacing": 1.0, "indent_cm": 1.0,
                 "margins_mm": {"left": 25, "right": 10, "top": 15, "bottom": 15}, "page_numbers": "bottom-right"},
    "approval": {"role": {"ru": "Проректор по учебной работе"}, "name": "А.А. Асанов",
                 "agreed": [{"role": {"ru": "Декан факультета"}, "name": "Б.Б. Бекова"}], "protocol": False},
    "grading": {"scale": [
        {"min": 90, "max": 100, "letter": "A", "grade": {"ru": "превосходно"}},
        {"min": 55, "max": 89, "letter": "B", "grade": {"ru": "зачтено"}},
        {"min": 0, "max": 54, "letter": "F", "grade": {"ru": "не зачтено"}},
    ]},
    "syllabus": {"hide_sections": ["sdg", "methods", "srs"],
                 "policies": {"ai_use": {"ru": "Политика ИИ вуза: только с указанием."}},
                 "extra_sections": [{"title": {"ru": "Материально-техническое обеспечение"},
                                     "text": {"ru": "Компьютерный класс 301."}}]},
}


def test_syllabus_follows_settings(course_dir):
    from docx import Document

    write_settings(course_dir.parent, UNI)
    edit_course = yaml.safe_load((course_dir / "course.yaml").read_text(encoding="utf-8"))
    edit_course.get("policies", {}).pop("ai_use", None)
    (course_dir / "course.yaml").write_text(yaml.safe_dump(edit_course, allow_unicode=True, sort_keys=False),
                                            encoding="utf-8")
    run("build_syllabus.py", course_dir / "course.yaml")
    path = course_dir / "out" / "syllabus_ru.docx"
    doc = Document(str(path))
    sec = doc.sections[0]
    assert round(sec.left_margin.mm) == 25 and round(sec.right_margin.mm) == 10
    assert round(sec.top_margin.mm) == 15 and round(sec.bottom_margin.mm) == 15
    assert sec.footer.paragraphs[0].alignment == 2  # номер справа
    normal = doc.styles["Normal"]
    assert normal.font.name == "Arial" and normal.font.size.pt == 12
    assert normal.paragraph_format.line_spacing == 1.0
    assert round(normal.paragraph_format.first_line_indent.cm, 2) == 1.0
    sizes = {r.font.size.pt for t in doc.tables for row in t.rows for c in row.cells
             for p in c.paragraphs for r in p.runs if r.font.size}
    assert 10 in sizes and 12 not in sizes
    text = docx_text(path)
    assert "Министерство образования и науки КР" in text
    assert "Университет Настроек" in text and "Ош 20" in text
    assert "Проректор по учебной работе" in text and "А.А. Асанов" in text
    assert "«СОГЛАСОВАНО»" in text and "Декан факультета" in text and "Б.Б. Бекова" in text
    assert "Протокол №" not in text  # approval.protocol: false
    # шкала вуза
    assert "превосходно" in text and "55–89" in text and "87–100" not in text
    # скрытые разделы и свои разделы
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert not any("устойчивого" in h or "самостоятельной" in h for h in headings)
    assert any(h.endswith("Материально-техническое обеспечение") for h in headings)
    assert "Компьютерный класс 301." in text
    assert "Политика ИИ вуза: только с указанием." in text
    nums = [int(h.split(".")[0]) for h in headings if h[:1].isdigit()]
    assert nums == list(range(1, len(nums) + 1))  # нумерация без дыр


def test_course_overrides_settings(course_dir):
    write_settings(course_dir.parent, UNI)
    data = load_course(course_dir / "course.yaml")
    assert data["policies"]["attendance"]  # из course.yaml, не из настроек
    data["grading_scale"] = [{"min": 0, "max": 100, "letter": "P", "grade": "pass"}]
    from course_check import grading_scale

    assert grading_scale(data)[0]["letter"] == "P"


def test_credit_hours_from_settings(course_dir):
    data = load_course(course_dir / "course.yaml")
    assert check_hours(data) == []
    write_settings(course_dir.parent, {"grading": {"credit_hours": 36}})
    data = load_course(course_dir / "course.yaml")
    assert any("× 36" in p for p in check_hours(data))


def test_tests_without_moodle(course_dir):
    bank = yaml.safe_load((course_dir / "tests.yaml").read_text(encoding="utf-8"))
    bank.pop("variants", None)  # в tests.yaml не задано — берётся из настроек вуза
    (course_dir / "tests.yaml").write_text(yaml.safe_dump(bank, allow_unicode=True, sort_keys=False), encoding="utf-8")
    write_settings(course_dir.parent, {"tests": {"moodle": False, "variants": 3}})
    run("build_tests.py", course_dir / "tests.yaml")
    out = course_dir / "out"
    assert (out / "tests_3.docx").exists() and (out / "tests.xlsx").exists()
    assert not (out / "tests.gift").exists() and not (out / "tests_moodle.xml").exists()
    run("build_tests.py", course_dir / "tests.yaml", "--formats", "gift")  # просьба явно — выгружаем
    assert (out / "tests.gift").exists()


def test_settings_error_stops_build(course_dir):
    write_settings(course_dir.parent, {"document": {"size": 100}})
    res = run("build_syllabus.py", course_dir / "course.yaml", ok=False)
    assert res.returncode != 0 and "document.size" in res.stderr


def test_university_check_script(tmp_path):
    res = run("university_check.py", tmp_path)
    assert "нормы КР" in res.stdout
    write_settings(tmp_path, {"university": {"name": "КНУ"}, "document": {"margins_mm": {"left": 25}}})
    res = run("university_check.py", tmp_path / "university")
    assert "левое поле, мм [document.margins_mm.left]: 25 (по умолчанию: 30)" in res.stdout
    write_settings(tmp_path, {"research": {"citation_style": "mla"}})
    res = run("university_check.py", tmp_path, ok=False)
    assert res.returncode == 1 and "research.citation_style" in res.stdout
    res = run("university_check.py", tmp_path, "--yaml", ok=False)
    assert yaml.safe_load(res.stdout)["research"]["citation_style"] == "gost"


def test_defaults_keys_match_template():
    tpl = yaml.safe_load(TEMPLATE.read_text(encoding="utf-8"))

    def keys(d, prefix=""):
        out = set()
        for k, v in d.items():
            out.add(prefix + k)
            if isinstance(v, dict) and prefix + k not in ("research.thesis_pages",):
                out |= keys(v, prefix + k + ".")
        return out

    assert keys(tpl) == keys(DEFAULTS)


def test_manual_and_report_follow_settings(ws):
    import json

    write_settings(ws, {
        "manual": {"paper": "a4", "approval_body": {"ru": "Учебно-методический совет КГТУ"}},
        "grading": {"scale": [
            {"min": 90, "max": 100, "letter": "A", "grade": {"ru": "отлично"}},
            {"min": 70, "max": 89, "letter": "B", "grade": {"ru": "хорошо"}},
            {"min": 50, "max": 69, "letter": "C", "grade": {"ru": "удовлетворительно"}},
            {"min": 0, "max": 49, "letter": "F", "grade": {"ru": "неудовлетворительно"}},
        ]},
    })
    c = ws / "course"
    spec = yaml.safe_load((c / "manual" / "manual.yaml").read_text(encoding="utf-8"))
    spec.pop("paper")  # не задано в manual.yaml — формат из настроек вуза
    (c / "manual" / "manual.yaml").write_text(yaml.safe_dump(spec, allow_unicode=True), encoding="utf-8")
    run_uv("build_manual.py", c / "manual" / "manual.yaml", "--no-pdf", cwd=ws)
    meta = json.loads((c / "out" / "manual-build" / "data.json").read_text(encoding="utf-8"))["meta"]
    assert meta["paper"] == "a4"
    assert meta["recommended"].startswith("Рекомендовано к изданию: Учебно-методический совет КГТУ")
    assert "Учебно-методический совет КГТУ" in docx_text(c / "out" / "manual.docx")
    r = run_uv("build_report.py", c / "report" / "course_report.yaml", cwd=ws)
    assert "pass=100.0" in r.stdout and "quality=80.0" in r.stdout
    text = docx_text(c / "out" / "report_course_report.docx")
    assert "Успеваемость (балл 50 и выше)" in text and "50–69" in text
