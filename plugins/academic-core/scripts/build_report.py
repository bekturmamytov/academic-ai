# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1", "pyyaml>=6", "openpyxl>=3.1"]
# ///
"""Отчёты преподавателя: report.yaml → out/report_<тип>.docx.

    uv run build_report.py <report.yaml> [--course course.yaml]

type: individual_plan — индивидуальный план; plan_report — отчёт о выполнении плана;
course_report — отчёт по дисциплине (успеваемость только из файла оценок преподавателя).
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from common import DataError, course_title, fail, new_document, out_dir, para, read_yaml, table, university_name  # noqa: E402
from content_kit import KR_SCALE, add_runs, kr_grade, load_context, parse_inline  # noqa: E402
from i18n import LABELS, t  # noqa: E402

LABELS.update({
    "r_plan": {"ru": "Индивидуальный план работы преподавателя", "ky": "Окутуучунун жеке иш планы",
               "en": "Individual work plan"},
    "r_plan_report": {"ru": "Отчёт о выполнении индивидуального плана работы преподавателя",
                      "ky": "Окутуучунун жеке иш планынын аткарылышы жөнүндө отчёт", "en": "Work plan completion report"},
    "r_course": {"ru": "Отчёт по дисциплине", "ky": "Сабак боюнча отчёт", "en": "Course report"},
    "r_year": {"ru": "на {} учебный год", "ky": "{} окуу жылына", "en": "for the {} academic year"},
    "r_teaching": {"ru": "Учебная работа", "ky": "Окуу иши", "en": "Teaching"},
    "r_methodical": {"ru": "Учебно-методическая работа", "ky": "Окуу-методикалык иш", "en": "Course development"},
    "r_research": {"ru": "Научно-исследовательская работа", "ky": "Илимий-изилдөө иши", "en": "Research"},
    "r_educational": {"ru": "Воспитательная работа", "ky": "Тарбиялык иш", "en": "Student mentoring"},
    "r_other": {"ru": "Организационная и общественная работа", "ky": "Уюштуруу жана коомдук иш", "en": "Service"},
    "r_name": {"ru": "Наименование работы", "ky": "Иштин аталышы", "en": "Activity"},
    "r_deadline": {"ru": "Срок", "ky": "Мөөнөтү", "en": "Deadline"},
    "r_result": {"ru": "Отметка о выполнении", "ky": "Аткарылышы", "en": "Status"},
    "r_discipline": {"ru": "Дисциплина", "ky": "Сабак", "en": "Course"},
    "r_group": {"ru": "Группа", "ky": "Топ", "en": "Group"},
    "r_plan_h": {"ru": "План, ч", "ky": "План, с", "en": "Plan, h"},
    "r_fact_h": {"ru": "Факт, ч", "ky": "Факт, с", "en": "Actual, h"},
    "r_other_h": {"ru": "Др.", "ky": "Б.", "en": "Other"},
    "r_sem": {"ru": "Сем.", "ky": "Сем.", "en": "Sem."},
    "r_teacher_sign": {"ru": "Преподаватель", "ky": "Окутуучу", "en": "Instructor"},
    "r_head_sign": {"ru": "Заведующий кафедрой", "ky": "Кафедра башчысы", "en": "Head of Department"},
    "r_students": {"ru": "Студентов в ведомости", "ky": "Ведомосттогу студенттер", "en": "Students listed"},
    "r_graded": {"ru": "Аттестовано (есть балл)", "ky": "Аттестацияланган", "en": "Graded"},
    "r_avg": {"ru": "Средний балл", "ky": "Орточо упай", "en": "Mean score"},
    "r_pass": {"ru": "Успеваемость (61 балл и выше)", "ky": "Өздөштүрүү (61 жана жогору)", "en": "Pass rate (61+)"},
    "r_quality": {"ru": "Качество знаний (74 балла и выше)", "ky": "Билимдин сапаты (74 жана жогору)", "en": "Quality (74+)"},
    "r_distribution": {"ru": "Распределение оценок", "ky": "Баалардын бөлүштүрүлүшү", "en": "Grade distribution"},
    "r_count": {"ru": "Студентов", "ky": "Студенттер", "en": "Students"},
    "r_conclusions": {"ru": "Выводы и предложения", "ky": "Корутунду жана сунуштар", "en": "Conclusions"},
    "r_source": {"ru": "Источник данных: ведомость преподавателя «{}».", "ky": "Маалымат булагы: окутуучунун ведомосту «{}».",
                 "en": "Data source: instructor's grade sheet “{}”."},
})
SECTIONS = ("methodical", "research", "educational", "other")
SCORE_WORDS = ("итог", "total", "жыйынтык", "балл", "score", "упай", "сумма")


def header_block(doc, data: dict, spec: dict, lang: str, title: str, approve: bool):
    prof = data.get("_profile", {})
    para(doc, university_name(data, lang), bold=True, align="center")
    if prof.get("department"):
        para(doc, f'{t("department", lang)} «{prof["department"]}»', align="center")
    if approve:
        para(doc)
        for line in (t("approve", lang), t("r_head_sign", lang), "________________________",
                     "«___» ____________ 20__ г."):
            para(doc, line, align="right")
    para(doc)
    para(doc, title.upper(), bold=True, align="center")
    year = spec.get("academic_year") or (data.get("course", {}) or {}).get("academic_year", "")
    if year:
        para(doc, t("r_year", lang).format(year), align="center")
    teacher = {**(data.get("teacher") or {}), **(spec.get("teacher") or {})}
    para(doc)
    for label, val in (("ФИО", teacher.get("name")), ("Должность", teacher.get("position")),
                       ("Учёная степень, звание", teacher.get("degree")), ("Ставка", spec.get("rate"))):
        if val:
            p = para(doc, indent=False)
            add_runs(p, [{"text": f"{label}: ", "b": True}, {"text": str(val)}])
    return teacher


def signatures(doc, lang: str, teacher: dict):
    para(doc)
    para(doc, f'{t("r_teacher_sign", lang)}  ________________  {teacher.get("name", "")}', indent=False)
    para(doc)
    para(doc, f'{t("r_head_sign", lang)}  ________________  ________________', indent=False)
    para(doc, "«___» ____________ 20__ г.", indent=False)


def plan(doc, data: dict, spec: dict, lang: str, report: bool, warns: list[str]):
    title = t("r_plan_report" if report else "r_plan", lang)
    teacher = header_block(doc, data, spec, lang, title, approve=not report)
    n = 0
    teaching = spec.get("teaching") or []
    n += 1
    doc.add_heading(f'{n}. {t("r_teaching", lang)}', level=2)
    if teaching:
        kinds = ("lectures", "practicals", "labs", "other")
        header = [t("r_discipline", lang), t("r_group", lang), t("r_sem", lang),
                  "Лек", "Пр", "Лаб", t("r_other_h", lang), t("r_plan_h", lang)]
        rows, plan_sum, fact_sum = [], 0, 0
        for it in teaching:
            total = sum(int(it.get(k) or 0) for k in kinds)
            plan_sum += total
            row = [it.get("discipline", ""), it.get("group", ""), it.get("semester", "")] + \
                  [it.get(k, "") for k in kinds] + [total]
            if report:
                fact = it.get("fact")
                if fact is None:
                    warns.append(f"учебная работа «{it.get('discipline', '')}»: нет факта часов (fact)")
                fact_sum += int(fact or 0)
                row.append("" if fact is None else fact)
            rows.append(row)
        total_row = [t("total", lang), "", "", "", "", "", "", plan_sum] + ([fact_sum] if report else [])
        rows.append(total_row)
        widths = [4.0, 2.2, 1.3, 1.3, 1.3, 1.3, 1.3, 1.6] + ([1.6] if report else [])
        if report:
            header.append(t("r_fact_h", lang))
            widths = [3.3, 2.0, 1.2, 1.3, 1.3, 1.3, 1.2, 1.6, 1.6]
        table(doc, header, rows, widths_cm=widths)
    else:
        warns.append("нет учебной нагрузки (teaching) — раздел пустой")
        para(doc, "[Нагрузка не указана]", italic=True)
    for key in SECTIONS:
        items = spec.get(key) or []
        if not items and key == "other":
            continue
        n += 1
        doc.add_heading(f'{n}. {t("r_" + key, lang)}', level=2)
        if not items:
            warns.append(f"раздел «{t('r_' + key, lang)}» пустой ({key})")
        rows = []
        for i, it in enumerate(items, 1):
            it = it if isinstance(it, dict) else {"name": str(it)}
            result = it.get("result", "") if report else ""
            if report and not result:
                warns.append(f"{t('r_' + key, lang)}: «{str(it.get('name', ''))[:40]}» — нет отметки о выполнении")
            rows.append([i, it.get("name", ""), it.get("deadline", ""), result])
        table(doc, ["№", t("r_name", lang), t("r_deadline", lang), t("r_result", lang)],
                   rows or [["", "", "", ""]], widths_cm=[1, 8.5, 3, 4])
    if report and spec.get("conclusions"):
        n += 1
        doc.add_heading(f'{n}. {t("r_conclusions", lang)}', level=2)
        for chunk in str(spec["conclusions"]).strip().split("\n\n"):
            add_runs(para(doc), parse_inline(" ".join(chunk.split())))
    signatures(doc, lang, teacher)


def _num(v) -> float | None:
    if v is None:
        return None
    s = str(v).strip().replace(",", ".").replace("\xa0", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def read_grades(path: Path) -> tuple[list[str], list[list]]:
    if not path.exists():
        raise DataError(f"не найден файл оценок {path}")
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
    elif path.suffix.lower() in (".csv", ".txt"):
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1251"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=";,\t")
        rows = [r for r in csv.reader(text.splitlines(), dialect) if any(c.strip() for c in r)]
    else:
        raise DataError(f"файл оценок {path.name}: нужен .xlsx или .csv")
    if len(rows) < 2:
        raise DataError(f"в файле оценок {path.name} нет строк со студентами")
    return [str(c or "").strip() for c in rows[0]], rows[1:]


def score_column(header: list[str], rows: list[list], wanted: str | None) -> int:
    low = [h.lower() for h in header]
    if wanted:
        if wanted.lower() not in low:
            raise DataError(f"в файле оценок нет столбца «{wanted}»; есть: {', '.join(header)}")
        return low.index(wanted.lower())
    for word in SCORE_WORDS:
        for i, h in enumerate(low):
            if word in h:
                return i
    raise DataError("не найден столбец с итоговым баллом — укажите score_column в report.yaml "
                    f"(столбцы: {', '.join(header)})")


def course_report(doc, data: dict, spec: dict, lang: str, base: Path, warns: list[str]):
    if not spec.get("grades"):
        raise DataError("для отчёта по дисциплине нужен файл оценок: grades: ведомость.xlsx (или .csv)")
    gpath = (base / spec["grades"]).resolve()
    header, rows = read_grades(gpath)
    col = score_column(header, rows, spec.get("score_column"))
    scores, missing = [], 0
    for r in rows:
        v = _num(r[col] if col < len(r) else None)
        if v is None:
            missing += 1
            continue
        if not 0 <= v <= 100:
            raise DataError(f"в ведомости балл {v} вне шкалы 0–100 — проверьте файл")
        scores.append(v)
    if not scores:
        raise DataError(f"в столбце «{header[col]}» нет чисел")
    if missing:
        warns.append(f"у {missing} студентов нет балла в столбце «{header[col]}» — учтены как не аттестованные")
    title = f'{t("r_course", lang)} «{course_title(data, lang) or spec.get("discipline", "")}»'
    teacher = header_block(doc, data, spec, lang, title, approve=False)
    c = data.get("course", {}) or {}
    for label, val in ((t("r_group", lang), spec.get("group")), (t("semester", lang).capitalize(), spec.get("semester") or c.get("semester")),
                       (t("credits", lang), data.get("credits"))):
        if val:
            add_runs(para(doc, indent=False), [{"text": f"{label}: ", "b": True}, {"text": str(val)}])
    counts = [0] * len(KR_SCALE)
    for s in scores:
        counts[kr_grade(s)[1]] += 1
    n = len(scores)
    li = {"ru": 3, "ky": 4, "en": 5}.get(lang, 3)
    doc.add_heading(f'1. {t("r_distribution", lang)}', level=2)
    table(doc, [t("letter", lang), t("points_range", lang), t("grade", lang), t("r_count", lang), "%"],
               [[r[2], f"{r[0]}–{r[1]}", r[li], counts[i], f"{counts[i] * 100 / n:.1f}"] for i, r in enumerate(KR_SCALE)],
               widths_cm=[2, 2.8, 5.2, 3, 3.5])
    passed = sum(1 for s in scores if s >= 61)
    quality = sum(1 for s in scores if s >= 74)
    doc.add_heading(f'2. {t("assessment", lang)}', level=2)
    table(doc, ["Показатель", "Значение"], [
        [t("r_students", lang), len(rows)],
        [t("r_graded", lang), n],
        [t("r_avg", lang), f"{sum(scores) / n:.1f}"],
        [t("r_pass", lang), f"{passed * 100 / n:.1f} % ({passed} из {n})"],
        [t("r_quality", lang), f"{quality * 100 / n:.1f} % ({quality} из {n})"],
    ], widths_cm=[10, 6.5])
    para(doc, t("r_source", lang).format(gpath.name), italic=True, size=12)
    doc.add_heading(f'3. {t("r_conclusions", lang)}', level=2)
    if spec.get("conclusions"):
        for chunk in str(spec["conclusions"]).strip().split("\n\n"):
            add_runs(para(doc), parse_inline(" ".join(chunk.split())))
    else:
        warns.append("нет выводов (conclusions) — раздел оставлен для заполнения")
        para(doc, "[Выводы и предложения преподавателя]", italic=True)
    signatures(doc, lang, teacher)
    return {"students": len(rows), "graded": n, "avg": round(sum(scores) / n, 1),
            "pass": round(passed * 100 / n, 1), "quality": round(quality * 100 / n, 1)}


def main() -> None:
    ap = argparse.ArgumentParser(description="Отчёты преподавателя в docx")
    ap.add_argument("report", help="путь к report.yaml")
    ap.add_argument("--course", help="путь к course.yaml (для отчёта по дисциплине)")
    a = ap.parse_args()
    src = Path(a.report).expanduser().resolve()
    warns: list[str] = []
    stats = None
    try:
        spec = read_yaml(src)
        course = a.course or (str((src.parent / spec["course"]).resolve()) if spec.get("course") else None)
        data = load_context(src, course)
        lang = spec.get("language") or data.get("language", "ru")
        kind = spec.get("type")
        doc = new_document()
        if kind in ("individual_plan", "plan_report"):
            plan(doc, data, spec, lang, kind == "plan_report", warns)
        elif kind == "course_report":
            stats = course_report(doc, data, spec, lang, src.parent, warns)
        else:
            raise DataError("type в report.yaml: individual_plan, plan_report или course_report")
        d = out_dir(data) if data.get("course") else src.parent / "out"
        d.mkdir(parents=True, exist_ok=True)
        path = d / f"report_{kind}.docx"
        doc.save(path)
    except DataError as e:
        fail(str(e))
    print(path)
    if stats:
        print("Итоги по ведомости: " + ", ".join(f"{k}={v}" for k, v in stats.items()))
    for w in warns:
        print(f"Предупреждение: {w}")


if __name__ == "__main__":
    main()
