# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0", "pyyaml>=6"]
# ///
"""Презентация лекции: slides.yaml → out/slides/NN.pptx (16:9, белый фон, один акцентный цвет).

    uv run build_slides.py <slides.yaml> [--course course.yaml]
"""
from __future__ import annotations

import argparse
import math
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from common import (  # noqa: E402
    DataError, city_name, course_title, fail, find_up, load_course, out_dir, read_yaml, university_context,
    university_name,
)
from i18n import LABELS, t  # noqa: E402

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

LABELS.update({
    "questions_title": {"ru": "Вопросы для обсуждения", "ky": "Талкуу үчүн суроолор", "en": "Questions for discussion"},
    "thanks": {"ru": "Спасибо за внимание", "ky": "Көңүл бурганыңыз үчүн рахмат", "en": "Thank you"},
    "image": {"ru": "Изображение", "ky": "Сүрөт", "en": "Image"},
    "diagram": {"ru": "Схема", "ky": "Схема", "en": "Diagram"},
})

FONT = "Arial"
SW, SH, M = 13.333, 7.5, 0.6
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MAX_BULLETS = 6
MAX_TITLE_WORDS = 8


def hex_color(s: str | None) -> RGBColor:
    s = (s or "#1F4E79").strip().lstrip("#")
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", s):
        raise DataError(f"цвет «{s}» в brand.yaml — нужен формат #RRGGBB")
    return RGBColor.from_string(s.upper())


def tint(c: RGBColor, k: float) -> RGBColor:
    return RGBColor(*(round(v + (255 - v) * k) for v in c))


def fits(texts: list[str], size: float, w: float, h: float, gap: float = 0.45, line: float = 1.2) -> bool:
    """Грубая оценка: помещается ли текст кеглем size в рамку w×h дюймов."""
    cpl = max(1, int(w * 72 / (size * 0.55)))
    lines = sum(max(1, math.ceil(len(s) / cpl)) for s in texts)
    need = lines * size * line / 72 + max(0, len(texts) - 1) * size * gap / 72
    return need <= h


def pick_size(texts, sizes, w, h, **kw) -> tuple[int, bool]:
    for s in sizes:
        if fits(texts, s, w, h, **kw):
            return s, True
    return sizes[-1], False


class Deck:
    def __init__(self, data: dict | None, spec: dict, base: Path):
        self.data, self.spec, self.base = data, spec, base
        ctx = data or university_context(base)  # без курса — цвет и логотип из university/ рядом
        brand = ctx.get("_brand", {}) or {}
        self.accent = hex_color(brand.get("color"))
        self.light = tint(self.accent, 0.88)
        logo = brand.get("logo")
        bdir = ctx.get("_brand_dir")
        self.logo = bdir / logo if logo and bdir and (bdir / logo).exists() else None
        self.lang = spec.get("language") or (data or {}).get("language", "ru")
        self.course = course_title(data, self.lang) if data else spec.get("course", "")
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SW), Inches(SH)
        self.layout = self.prs.slide_layouts[6]
        self.warns: list[str] = []
        self.n = 0

    # ---------- примитивы ----------
    def text(self, slide, x, y, w, h, text="", size=18, *, bold=False, color=DARK, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, italic=False):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.alignment = align
        self.run(p, text, size, bold=bold, color=color, italic=italic)
        return tb

    @staticmethod
    def run(p, text, size, *, bold=False, color=DARK, italic=False):
        r = p.add_run()
        r.text = str(text)
        f = r.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
        f.color.rgb = color
        return r

    def rect(self, slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None):
        s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill is None:
            s.fill.background()
        else:
            s.fill.solid()
            s.fill.fore_color.rgb = fill
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = line
            s.line.width = Pt(1.25)
        style = s._element.find(qn("p:style"))
        if style is not None:
            s._element.remove(style)
        return s

    def bullet(self, p, level: int, size: int):
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.38 + 0.4 * level))))
        pPr.set("indent", str(-Emu(Inches(0.38))))
        clr = pPr.makeelement(qn("a:buClr"), {})
        clr.append(clr.makeelement(qn("a:srgbClr"), {"val": str(self.accent)}))
        pPr.append(clr)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": FONT}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•" if level == 0 else "–"}))

    def new_slide(self, notes: str = "", *, chrome=True):
        s = self.prs.slides.add_slide(self.layout)
        self.n += 1
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = WHITE
        s.notes_slide.notes_text_frame.text = flow(notes)
        if chrome:
            self.text(s, M, SH - 0.5, 9, 0.3, self.course, 10, color=GREY)
            self.text(s, SW - M - 1, SH - 0.5, 1, 0.3, str(self.n), 10, color=GREY, align=PP_ALIGN.RIGHT)
            if self.logo:
                s.shapes.add_picture(str(self.logo), Inches(SW - M - 1.2), Inches(0.4), height=Inches(0.55))
        return s

    def title(self, s, text: str, where: str):
        if len(str(text).split()) > MAX_TITLE_WORDS:
            self.warns.append(f"{where}: заголовок длиннее {MAX_TITLE_WORDS} слов — сократите")
        w = SW - 2 * M - (1.5 if self.logo else 0)
        size, ok = pick_size([text], [32, 28, 24], w, 1.0, line=1.15)
        self.text(s, M, 0.4, w, 1.0, text, size, bold=True, color=self.accent, anchor=MSO_ANCHOR.BOTTOM)
        self.rect(s, M, 1.55, 1.0, 0.06, self.accent)
        if not ok:
            self.warns.append(f"{where}: заголовок не помещается в две строки")

    def bullets(self, s, items: list, x, y, w, h, where: str):
        flat = []
        for it in items:
            if isinstance(it, dict):
                flat.append((0, str(it.get("text", ""))))
                flat += [(1, str(x_)) for x_ in it.get("sub", []) or []]
            else:
                flat.append((0, str(it)))
        if sum(1 for lvl, _ in flat if lvl == 0) > MAX_BULLETS:
            self.warns.append(f"{where}: больше {MAX_BULLETS} пунктов — разделите слайд")
        size, ok = pick_size([txt for _, txt in flat], [28, 26, 24, 22, 20, 18, 16], w - 0.4, h)
        if not ok:
            self.warns.append(f"{where}: текст не помещается даже кеглем 16 — сократите пункты")
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, (lvl, txt) in enumerate(flat):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(0 if i == 0 else size * 0.45)
            p.line_spacing = 1.1
            self.bullet(p, lvl, size)
            for j, part in enumerate(re.split(r"\*\*(.+?)\*\*", txt)):
                if part:
                    self.run(p, part, size if lvl == 0 else size - 2, bold=bool(j % 2), color=DARK if lvl == 0 else GREY)
        return tb

    def box_text(self, shape, title: str, body: str, w: float, h: float):
        tf = shape.text_frame
        tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
        for side in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{side}", Inches(0.12))
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        texts = [x for x in (title, body) if x]
        size, _ = pick_size(texts, [20, 18, 16, 14, 12], w - 0.24, h - 0.24, gap=0.3)
        first = True
        for txt, bold, col in ((title, True, self.accent), (body, False, DARK)):
            if not txt:
                continue
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            self.run(p, txt, size if bold else max(size - 2, 11), bold=bold, color=col)
            first = False

    # ---------- визуальные блоки ----------
    def diagram(self, s, spec: dict, x, y, w, h, where: str):
        kind = spec.get("kind", "steps")
        items = [it if isinstance(it, dict) else {"title": str(it)} for it in spec.get("items", [])]
        if not items:
            return self.placeholder(s, f'{t("diagram", self.lang)}: {spec.get("description", "")}', x, y, w, h)
        if len(items) > 6:
            self.warns.append(f"{where}: в схеме больше 6 элементов — упростите")
            items = items[:6]
        n = len(items)
        if kind == "cards":
            cols = 2 if n > 1 else 1
            rows = math.ceil(n / cols)
            gap = 0.25
            cw, ch = (w - gap * (cols - 1)) / cols, min(1.9, (h - gap * (rows - 1)) / rows)
            for i, it in enumerate(items):
                cx, cy = x + (i % cols) * (cw + gap), y + (i // cols) * (ch + gap)
                box = self.rect(s, cx, cy, cw, ch, self.light, MSO_SHAPE.ROUNDED_RECTANGLE)
                box.adjustments[0] = 0.08
                self.box_text(box, it.get("title", ""), it.get("text", ""), cw, ch)
            return
        horizontal = w >= 8
        gap = 0.45
        if horizontal:
            bw, bh = (w - gap * (n - 1)) / n, min(2.0, h)
            by = y + (h - bh) / 2
            for i, it in enumerate(items):
                bx = x + i * (bw + gap)
                box = self.rect(s, bx, by, bw, bh, self.light, MSO_SHAPE.ROUNDED_RECTANGLE, line=self.accent)
                self.box_text(box, it.get("title", ""), it.get("text", ""), bw, bh)
                if i < n - 1:
                    self.rect(s, bx + bw + 0.08, by + bh / 2 - 0.15, gap - 0.16, 0.3, self.accent, MSO_SHAPE.RIGHT_ARROW)
        else:
            bh = min(1.1, (h - gap * (n - 1)) / n)
            for i, it in enumerate(items):
                by = y + i * (bh + gap)
                box = self.rect(s, x, by, w, bh, self.light, MSO_SHAPE.ROUNDED_RECTANGLE, line=self.accent)
                self.box_text(box, it.get("title", ""), it.get("text", ""), w, bh)
                if i < n - 1:
                    self.rect(s, x + w / 2 - 0.15, by + bh + 0.07, 0.3, gap - 0.14, self.accent, MSO_SHAPE.DOWN_ARROW)

    def placeholder(self, s, label: str, x, y, w, h):
        box = self.rect(s, x, y, w, h, None, MSO_SHAPE.RECTANGLE, line=GREY)
        box.line.dash_style = MSO_LINE.DASH
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        self.run(p, label, 14, color=GREY, italic=True)

    def image(self, s, spec, x, y, w, h):
        spec = spec if isinstance(spec, dict) else {"path": str(spec)}
        path = self.base / spec.get("path", "") if spec.get("path") else None
        if path and path.is_file():
            from PIL import Image

            with Image.open(path) as im:
                iw, ih = im.size
            k = min(w / iw, h / ih)
            pw, ph = iw * k, ih * k
            s.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
        else:
            label = spec.get("description") or spec.get("caption") or spec.get("path", "")
            self.placeholder(s, f'{t("image", self.lang)}: {label}', x, y, w, h)

    def table(self, s, spec: dict, x, y, w, h, where: str):
        header = [str(c) for c in spec.get("header", [])]
        rows = [[("" if c is None else str(c)) for c in r] for r in spec.get("rows", [])]
        if not header:
            raise DataError(f"{where}: у таблицы нет header")
        if len(rows) > 8:
            self.warns.append(f"{where}: в таблице больше 8 строк — на слайде будет мелко")
        cols = len(header)
        size = 24 if len(rows) <= 3 else 20 if len(rows) <= 5 else 16 if len(rows) <= 8 else 13
        row_h = size * 2.1 / 72
        gt = s.shapes.add_table(len(rows) + 1, cols, Inches(x), Inches(y), Inches(w),
                                Inches(row_h * (len(rows) + 1))).table
        for r, vals in enumerate([header] + rows):
            for c in range(cols):
                cell = gt.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.accent if r == 0 else (self.light if r % 2 == 0 else WHITE)
                cell.margin_left = cell.margin_right = Inches(0.1)
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = ""
                self.run(p, vals[c] if c < len(vals) else "", size, bold=r == 0, color=WHITE if r == 0 else DARK)
        if row_h * (len(rows) + 1) > h:
            self.warns.append(f"{where}: таблица выше слайда — разделите")

    # ---------- слайды ----------
    def title_slide(self):
        spec, data = self.spec, self.data or {}
        s = self.new_slide(spec.get("title_notes", ""), chrome=False)
        top = 0.6
        if self.logo:
            s.shapes.add_picture(str(self.logo), Inches(M), Inches(top), height=Inches(0.9))
        uni = university_name(data, self.lang) if data else ""
        if uni:
            self.text(s, M + (1.2 if self.logo else 0), top + 0.15, 10, 0.6, uni, 14, color=GREY, anchor=MSO_ANCHOR.MIDDLE)
        self.text(s, M, 2.15, SW - 2 * M, 0.4, self.course.upper(), 16, bold=True, color=self.accent)
        title = spec.get("title", "")
        size, _ = pick_size([title], [40, 36, 32, 28], SW - 2 * M, 1.6, line=1.15)
        self.text(s, M, 2.6, SW - 2 * M, 1.6, title, size, bold=True, color=DARK)
        if spec.get("subtitle"):
            self.text(s, M, 4.25, SW - 2 * M, 0.5, spec["subtitle"], 20, color=GREY)
        self.rect(s, M, 4.95, 1.4, 0.07, self.accent)
        teacher = data.get("teacher", {}) or {}
        who = ", ".join(x for x in (teacher.get("name"), teacher.get("degree")) if x)
        if who:
            self.text(s, M, 5.25, SW - 2 * M, 0.45, who, 18, color=DARK)
        if teacher.get("position"):
            self.text(s, M, 5.7, SW - 2 * M, 0.4, teacher["position"], 14, color=GREY)
        city = city_name(data)
        year = str(data.get("course", {}).get("academic_year", ""))
        footer = " · ".join(x for x in (city, year) if x)
        if footer:
            self.text(s, M, SH - 0.7, 6, 0.3, footer, 12, color=GREY)

    def section_slide(self, sl: dict, no: int):
        s = self.new_slide(sl.get("notes", ""), chrome=False)
        self.rect(s, 0, 0, 0.25, SH, self.accent)
        self.text(s, 1.2, 2.0, 3, 1.2, f"{no:02d}", 66, bold=True, color=tint(self.accent, 0.55))
        size, _ = pick_size([sl.get("title", "")], [40, 36, 32], SW - 2.4, 2.0, line=1.15)
        self.text(s, 1.2, 3.2, SW - 2.4, 2.0, sl.get("title", ""), size, bold=True, color=self.accent)
        if sl.get("subtitle"):
            self.text(s, 1.2, 5.3, SW - 2.4, 0.6, sl["subtitle"], 20, color=GREY)

    def content_slide(self, sl: dict, where: str):
        s = self.new_slide(sl.get("notes", ""))
        self.title(s, sl.get("title", ""), where)
        top, h = 1.9, SH - 1.9 - 0.75
        visual = next((k for k in ("diagram", "image", "placeholder") if sl.get(k)), None)
        bullets = sl.get("bullets") or []
        if bullets and visual:
            bw = (SW - 2 * M) * 0.52
            self.bullets(s, bullets, M, top, bw, h, where)
            vx, vw = M + bw + 0.4, SW - 2 * M - bw - 0.4
        elif bullets:
            self.bullets(s, bullets, M, top, SW - 2 * M, h, where)
            return
        else:
            vx, vw = M, SW - 2 * M
        if visual == "diagram":
            self.diagram(s, sl["diagram"], vx, top, vw, h, where)
        elif visual == "image":
            self.image(s, sl["image"], vx, top, vw, h)
        elif visual == "placeholder":
            self.placeholder(s, str(sl["placeholder"]), vx, top, vw, h)
        else:
            self.warns.append(f"{where}: пустой слайд — нет bullets, diagram, image или table")

    def table_slide(self, sl: dict, where: str):
        s = self.new_slide(sl.get("notes", ""))
        self.title(s, sl.get("title", ""), where)
        self.table(s, sl.get("table") or {}, M, 1.9, SW - 2 * M, SH - 1.9 - 0.75, where)

    def closing_slide(self):
        qs = self.spec.get("questions") or []
        s = self.new_slide(self.spec.get("closing_notes", ""))
        if qs:
            self.title(s, t("questions_title", self.lang), "заключительный слайд")
            tb = s.shapes.add_textbox(Inches(M), Inches(1.9), Inches(SW - 2 * M), Inches(SH - 2.65))
            tf = tb.text_frame
            tf.word_wrap = True
            size, ok = pick_size([str(q) for q in qs], [26, 24, 22, 20, 18], SW - 2 * M - 0.6, SH - 2.65)
            for i, q in enumerate(qs):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_before = Pt(0 if i == 0 else size * 0.5)
                self.run(p, f"{i + 1}. ", size, bold=True, color=self.accent)
                self.run(p, str(q), size)
            if not ok:
                self.warns.append("заключительный слайд: вопросы не помещаются — оставьте 3–5")
        else:
            self.text(s, M, 2.8, SW - 2 * M, 1.2, t("thanks", self.lang), 40, bold=True, color=self.accent,
                      align=PP_ALIGN.CENTER)
        email = ((self.data or {}).get("teacher") or {}).get("email")
        if email:
            self.text(s, M, SH - 1.1, SW - 2 * M, 0.4, email, 14, color=GREY, align=PP_ALIGN.CENTER)

    def build(self) -> Presentation:
        slides = self.spec.get("slides") or []
        if not slides:
            raise DataError("в slides.yaml нет списка slides")
        self.title_slide()
        section = 0
        for i, sl in enumerate(slides, start=1):
            where = f"слайд {i + 1} «{str(sl.get('title', ''))[:40]}»"
            kind = sl.get("type") or ("table" if sl.get("table") else "content")
            if not str(sl.get("notes") or "").strip():
                self.warns.append(f"{where}: нет заметок докладчика")
            if kind == "section":
                section += 1
                self.section_slide(sl, section)
            elif kind == "table":
                self.table_slide(sl, where)
            elif kind == "content":
                self.content_slide(sl, where)
            else:
                raise DataError(f"{where}: неизвестный type «{kind}» — section, content или table")
        self.closing_slide()
        return self.prs


def flow(text) -> str:
    """Заметки: переносы строк внутри абзаца → пробел, абзацы — через пустую строку."""
    return "\n\n".join(" ".join(par.split()) for par in str(text or "").strip().split("\n\n") if par.strip())


def plain(text) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", str(text))


def media_export(deck: Deck) -> dict:
    """Та же презентация в простой схеме для видео-лекции (academic-media): out/slides/NN.yaml."""
    spec, data = deck.spec, deck.data or {}
    teacher = data.get("teacher") or {}
    slides = [{"type": "title", "title": spec.get("title", ""), "subtitle": spec.get("subtitle", ""),
               "notes": flow(spec.get("title_notes"))}]
    for sl in spec.get("slides") or []:
        kind = sl.get("type") or ("table" if sl.get("table") else "content")
        out = {"type": "bullets" if kind == "content" else kind, "title": plain(sl.get("title", ""))}
        if kind == "section" and sl.get("subtitle"):
            out["subtitle"] = sl["subtitle"]
        if kind == "table":
            tb = sl.get("table") or {}
            out["table"] = {"header": [plain(c) for c in tb.get("header", [])],
                            "rows": [[plain(c) for c in r] for r in tb.get("rows", [])]}
        if kind == "content":
            bullets = []
            for it in sl.get("bullets") or []:
                if isinstance(it, dict):
                    bullets.append(plain(it.get("text", "")))
                    bullets += [f"– {plain(x)}" for x in it.get("sub", []) or []]
                else:
                    bullets.append(plain(it))
            if not bullets and sl.get("diagram"):
                for it in sl["diagram"].get("items", []):
                    it = it if isinstance(it, dict) else {"title": str(it)}
                    bullets.append(" — ".join(x for x in (it.get("title"), it.get("text")) if x))
            out["bullets"] = bullets
        out["notes"] = flow(sl.get("notes"))
        slides.append(out)
    qs = [plain(q) for q in spec.get("questions") or []]
    slides.append({"type": "closing", "title": t("questions_title" if qs else "thanks", deck.lang),
                   "subtitle": teacher.get("email", ""), "bullets": qs,
                   "notes": flow(spec.get("closing_notes"))})
    return {"title": spec.get("title", ""), "course": deck.course,
            "teacher": ", ".join(x for x in (teacher.get("name"), teacher.get("degree")) if x),
            "accent": f"#{deck.accent}", "lang": deck.lang, "slides": slides}


def main() -> None:
    ap = argparse.ArgumentParser(description="Сборка презентации из slides.yaml")
    ap.add_argument("slides", help="путь к slides.yaml")
    ap.add_argument("--course", help="путь к course.yaml (по умолчанию ищется выше по папкам)")
    a = ap.parse_args()
    src = Path(a.slides).expanduser().resolve()
    try:
        spec = read_yaml(src)
        course = Path(a.course).expanduser().resolve() if a.course else find_up(src.parent, "course.yaml")
        data = load_course(course) if course else None
        deck = Deck(data, spec, src.parent)
        prs = deck.build()
        m = re.match(r"(\d+)", src.stem)
        no = int(spec.get("week") or (m.group(1) if m else 1))
        d = (out_dir(data) if data else src.parent / "out") / "slides"
        d.mkdir(parents=True, exist_ok=True)
        path = d / f"{no:02d}.pptx"
        prs.save(path)
        import yaml

        (d / f"{no:02d}.yaml").write_text(yaml.safe_dump(media_export(deck), allow_unicode=True, sort_keys=False),
                                          encoding="utf-8")
    except DataError as e:
        fail(str(e))
    print(path)
    if not data:
        print("Предупреждение: course.yaml не найден — на титуле нет курса и преподавателя")
    for w in deck.warns:
        print(f"Предупреждение: {w}")


if __name__ == "__main__":
    main()
