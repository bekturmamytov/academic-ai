# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1", "python-docx>=1.1", "pyyaml>=6"]
# ///
"""Банк тестов → out/tests.xlsx, бланки out/tests_<N>.docx, ключ out/tests_key.docx,
Moodle out/tests.gift и out/tests_moodle.xml.

    uv run build_tests.py "<папка курса>/tests.yaml" [--course course.yaml] [--check]

--check — только проверка банка (покрытие тем и РО, ключи, дубли, дистракторы).
Схема tests.yaml — skills/tests/references/tests-schema.md.
"""
from __future__ import annotations

import argparse
import html
import random
import re
import sys
import xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from common import DataError, course_title, fail, load_course, new_document, para, read_yaml, table  # noqa: E402
from course_check import BLOOM_LEVELS, report  # noqa: E402
from i18n import bloom_name, t, tv  # noqa: E402

TYPES = ("single", "multiple", "truefalse", "matching", "short", "open")
DIFFICULTY = ("easy", "medium", "hard")
LETTERS = {"ru": "АБВГДЕЖЗИК", "ky": "АБВГДЕЖЗИК", "en": "ABCDEFGHIJ"}
CATCH_ALL = re.compile(r"все (вышеперечисленн|перечисленн|ответы верн)|нет (правильного|верного)|"
                       r"all of the above|none of the above|баары туура", re.I)


# ---------- загрузка и проверка ----------

def norm(s) -> str:
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def as_list(v) -> list:
    if v is None:
        return []
    return list(v) if isinstance(v, (list, tuple)) else [v]


def load(tests_yaml: Path, course_arg: str | None) -> tuple[dict, dict | None]:
    bank = read_yaml(tests_yaml)
    bank["_dir"] = tests_yaml.parent
    course_path = None
    if course_arg:
        course_path = Path(course_arg)
    elif bank.get("course"):
        course_path = tests_yaml.parent / bank["course"]
    elif (tests_yaml.parent / "course.yaml").exists():
        course_path = tests_yaml.parent / "course.yaml"
    course = load_course(course_path) if course_path else None
    return bank, course


def validate_bank(bank: dict, course: dict | None) -> tuple[list[str], list[str]]:
    err: list[str] = []
    warn: list[str] = []
    qs = bank.get("questions") or []
    if not qs:
        return ["questions: банк пуст"], []
    ids, texts = set(), {}
    for i, q in enumerate(qs, 1):
        qid = str(q.get("id") or "")
        w = f"вопрос {qid or i}"
        if not qid:
            err.append(f"{w}: нет id")
        elif qid in ids:
            err.append(f"{w}: id повторяется")
        ids.add(qid)
        typ = q.get("type")
        if typ not in TYPES:
            err.append(f"{w}: type «{typ}» — допустимо {', '.join(TYPES)}")
            continue
        text = q.get("text")
        if not text:
            err.append(f"{w}: нет текста вопроса")
        else:
            key = norm(text)
            if key in texts:
                err.append(f"{w}: дублирует вопрос {texts[key]}")
            texts[key] = qid
        if q.get("difficulty") not in DIFFICULTY:
            err.append(f"{w}: difficulty «{q.get('difficulty')}» — допустимо {', '.join(DIFFICULTY)}")
        if q.get("bloom") not in BLOOM_LEVELS:
            err.append(f"{w}: bloom «{q.get('bloom')}» — допустимо {', '.join(BLOOM_LEVELS)}")
        if not isinstance(q.get("topic"), int):
            err.append(f"{w}: topic — номер недели темы (целое число)")
        if not q.get("outcome"):
            err.append(f"{w}: не указан РО (outcome)")
        pts = q.get("points", 1)
        if not isinstance(pts, (int, float)) or pts <= 0:
            err.append(f"{w}: points — число больше нуля")

        ans = q.get("answer")
        if typ in ("single", "multiple"):
            opts = [str(o) for o in as_list(q.get("options"))]
            keys = [str(a) for a in as_list(ans)]
            if len(opts) < 4:
                err.append(f"{w}: вариантов ответа {len(opts)} — нужно не меньше 4")
            if len({norm(o) for o in opts}) != len(opts):
                err.append(f"{w}: варианты ответа повторяются")
            if not keys:
                err.append(f"{w}: нет ключа (answer)")
            for k in keys:
                if k not in opts:
                    err.append(f"{w}: ответ «{k}» не совпадает ни с одним вариантом")
            if typ == "single" and len(keys) > 1:
                err.append(f"{w}: type single, а правильных ответов {len(keys)} — используйте multiple")
            if typ == "multiple" and keys and len(keys) >= len(opts):
                err.append(f"{w}: все варианты верные — нет дистракторов")
            if typ == "multiple" and len(keys) == 1:
                warn.append(f"{w}: multiple с одним верным ответом — может, single?")
            if any(CATCH_ALL.search(o) for o in opts):
                warn.append(f"{w}: вариант «все/ни один из перечисленных» — слабый дистрактор")
            distract = [o for o in opts if o not in keys]
            if keys and distract and typ == "single":
                avg = sum(len(o) for o in distract) / len(distract)
                if len(keys[0]) > 20 and len(keys[0]) > 2 * avg:
                    warn.append(f"{w}: правильный ответ заметно длиннее остальных — подсказка студенту")
        elif typ == "truefalse":
            if not isinstance(ans, bool):
                err.append(f"{w}: answer — true или false")
        elif typ == "matching":
            pairs = q.get("pairs") or []
            if len(pairs) < 3:
                err.append(f"{w}: пар для соответствия {len(pairs)} — нужно не меньше 3")
            if any(not isinstance(p, (list, tuple)) or len(p) != 2 or not all(p) for p in pairs):
                err.append(f"{w}: pairs — список пар [левое, правое]")
            else:
                if len({norm(p[0]) for p in pairs}) != len(pairs) or len({norm(p[1]) for p in pairs}) != len(pairs):
                    err.append(f"{w}: в парах есть повторы")
        elif typ == "short":
            if not [a for a in as_list(ans) if str(a).strip()]:
                err.append(f"{w}: нет допустимых ответов (answer)")
        elif typ == "open":
            if not ans:
                err.append(f"{w}: нет эталона или критериев оценки (answer)")

    # покрытие
    if course:
        weeks = {tp.get("week"): tp for tp in course.get("topics", []) or []}
        outs = {str(o.get("id")) for o in course.get("outcomes", []) or []}
        for q in qs:
            if isinstance(q.get("topic"), int) and q["topic"] not in weeks:
                err.append(f"вопрос {q.get('id')}: темы недели {q['topic']} нет в course.yaml")
            if q.get("outcome") and str(q["outcome"]) not in outs:
                err.append(f"вопрос {q.get('id')}: РО «{q['outcome']}» нет в course.yaml")
        scope = bank.get("topics")  # тест на часть курса: недели, которые должны быть покрыты
        if scope is not None and not (isinstance(scope, list) and all(isinstance(w, int) for w in scope)):
            err.append("topics: список номеров недель, например [1, 2, 3, 4, 5, 6, 7]")
            scope = None
        want_t = [w for w in sorted(weeks) if scope is None or w in scope]
        want_o = outs if scope is None else {str(o) for w in want_t for o in weeks[w].get("outcomes", []) or []}
        got_t = {q.get("topic") for q in qs}
        got_o = {str(q.get("outcome")) for q in qs}
        miss_t = [str(w) for w in want_t if w not in got_t]
        miss_o = sorted(want_o - got_o)
        if miss_t:
            warn.append(f"покрытие: нет вопросов по темам недель {', '.join(miss_t)}")
        if miss_o:
            warn.append(f"покрытие: нет вопросов по РО {', '.join(miss_o)}")
    else:
        warn.append("course.yaml не найден — покрытие тем и РО не проверено")

    n = len(qs)
    diff = Counter(q.get("difficulty") for q in qs)
    if n >= 10 and not diff.get("hard"):
        warn.append("сложность: нет вопросов повышенного уровня (hard)")
    if n >= 10 and max(diff.values()) / n > 0.7:
        warn.append(f"сложность: {max(diff.values())} из {n} вопросов одного уровня — разнообразьте")
    bl = Counter(q.get("bloom") for q in qs)
    if n >= 10 and bl.get("remember", 0) / n > 0.5:
        warn.append("Блум: больше половины вопросов на запоминание — добавьте применение и анализ")
    pv = bank.get("per_variant")
    if pv is not None and (not isinstance(pv, int) or not 1 <= pv <= n):
        err.append(f"per_variant: от 1 до {n}")
    nv = bank.get("variants", 2)
    if not isinstance(nv, int) or not 1 <= nv <= 30:
        err.append("variants: от 1 до 30")
    return err, warn


# ---------- варианты ----------

def make_variants(bank: dict) -> list[list[dict]]:
    """Варианты: вопросы перемешаны, при per_variant < банка — поровну по темам."""
    qs = bank["questions"]
    per = bank.get("per_variant") or len(qs)
    seed = bank.get("seed", 0)
    res = []
    for v in range(1, int(bank.get("variants", 2)) + 1):
        rng = random.Random(f"{seed}-{v}")
        by_topic: dict = {}
        for q in qs:
            by_topic.setdefault(q.get("topic"), []).append(q)
        pools = list(by_topic.values())
        for p in pools:
            rng.shuffle(p)
        rng.shuffle(pools)
        picked = []
        while len(picked) < per:
            for p in pools:
                if p and len(picked) < per:
                    picked.append(p.pop())
        rng.shuffle(picked)
        items = []
        for q in picked:
            item = {"q": q}
            if q["type"] in ("single", "multiple"):
                opts = [str(o) for o in q["options"]]
                rng.shuffle(opts)
                item["options"] = opts
            elif q["type"] == "matching":
                rights = [p[1] for p in q["pairs"]]
                rng.shuffle(rights)
                item["rights"] = rights
            items.append(item)
        res.append(items)
    return res


def key_text(item: dict, lang: str) -> str:
    q = item["q"]
    L = LETTERS.get(lang, LETTERS["ru"])
    typ = q["type"]
    if typ in ("single", "multiple"):
        keys = {str(a) for a in as_list(q["answer"])}
        return ", ".join(L[i] for i, o in enumerate(item["options"]) if o in keys)
    if typ == "truefalse":
        return t("true" if q["answer"] else "false", lang)
    if typ == "matching":
        return ", ".join(f"{i}–{L[item['rights'].index(p[1])]}" for i, p in enumerate(q["pairs"], 1))
    if typ == "short":
        return " / ".join(str(a) for a in as_list(q["answer"]))
    return str(q["answer"])


def bank_key(q: dict, lang: str) -> str:
    if q["type"] in ("single", "multiple"):
        return "; ".join(str(a) for a in as_list(q["answer"]))
    if q["type"] == "truefalse":
        return t("true" if q["answer"] else "false", lang)
    if q["type"] == "matching":
        return "; ".join(f"{a} — {b}" for a, b in q["pairs"])
    if q["type"] == "short":
        return " / ".join(str(a) for a in as_list(q["answer"]))
    return str(q["answer"])


# ---------- docx ----------

def test_document():
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    doc = new_document()
    pf = doc.styles["Normal"].paragraph_format
    doc.styles["Normal"].font.size = Pt(12)
    pf.line_spacing = 1.0
    pf.first_line_indent = Cm(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return doc


def header_block(doc, bank: dict, course: dict | None, lang: str, subtitle: str) -> None:
    from docx.shared import Pt

    if course:
        from common import university_name
        uni = university_name(course, lang)
        if uni:
            para(doc, uni, align="center")
        para(doc, course_title(course, lang), bold=True, align="center")
    p = para(doc, f'{bank.get("title") or t("test", lang)} · {subtitle}', bold=True, align="center")
    p.paragraph_format.space_after = Pt(6)


def blank(bank: dict, course: dict | None, items: list[dict], v: int, lang: str):
    from docx.shared import Pt

    doc = test_document()
    L = LETTERS.get(lang, LETTERS["ru"])
    header_block(doc, bank, course, lang, f'{t("variant", lang)} {v}')
    para(doc, f'{t("student", lang)}: ______________________________   {t("group", lang)}: ________   '
              f'{t("date", lang)}: ___________')
    note = t("instructions", lang)
    if bank.get("time_minutes"):
        note += {"ru": f" Время — {bank['time_minutes']} мин.", "ky": f" Убакыт — {bank['time_minutes']} мүн.",
                 "en": f" Time: {bank['time_minutes']} min."}.get(lang, "")
    para(doc, note, italic=True).paragraph_format.space_after = Pt(6)
    for n, item in enumerate(items, 1):
        q = item["q"]
        p = para(doc)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True
        p.add_run(f"{n}. ").bold = True
        p.add_run(str(q["text"]))
        hint = p.add_run(f' ({t("h_" + q["type"], lang)}; {q.get("points", 1)} б.)' if lang != "en"
                         else f' ({t("h_" + q["type"], lang)}; {q.get("points", 1)} pt)')
        hint.italic = True
        typ = q["type"]
        if typ in ("single", "multiple"):
            for i, o in enumerate(item["options"]):
                op = para(doc, f"    {L[i]}) {o}")
                op.paragraph_format.keep_with_next = i < len(item["options"]) - 1
        elif typ == "truefalse":
            para(doc, f'    ☐ {t("true", lang)}     ☐ {t("false", lang)}')
        elif typ == "matching":
            rows = []
            for i in range(len(q["pairs"])):
                rows.append([f"{i + 1}. {q['pairs'][i][0]}", f"{L[i]}) {item['rights'][i]}"])
            tbl = table(doc, ["", ""], rows, widths_cm=[7.5, 9.0])
            tbl._tbl.remove(tbl.rows[0]._tr)
            para(doc, f'{t("answer", lang)}: ' + "   ".join(f"{i}–___" for i in range(1, len(q["pairs"]) + 1)))
        elif typ == "short":
            para(doc, f'{t("answer", lang)}: ________________________________________')
        else:
            for _ in range(6):
                para(doc, "_" * 78)
    return doc


def key_doc(bank: dict, course: dict | None, variants: list[list[dict]], lang: str):
    doc = test_document()
    header_block(doc, bank, course, lang, t("answer_key", lang))
    for v, items in enumerate(variants, 1):
        para(doc, f'{t("variant", lang)} {v}', bold=True).paragraph_format.keep_with_next = True
        rows = [[n, it["q"]["id"], key_text(it, lang), it["q"].get("points", 1)] for n, it in enumerate(items, 1)]
        rows.append(["", "", t("total_points", lang), sum(it["q"].get("points", 1) for it in items)])
        table(doc, ["№", "ID", t("answer", lang), t("points", lang)], rows, widths_cm=[1.2, 1.8, 11.5, 2.0])
    return doc


# ---------- xlsx ----------

def build_xlsx(bank: dict, course: dict | None, variants: list[list[dict]], lang: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L

    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    font, bold = Font(name="Times New Roman", size=11), Font(name="Times New Roman", size=11, bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    def grid(ws, r1, r2, ncol):
        for r in range(r1, r2 + 1):
            for c in range(1, ncol + 1):
                cell = ws.cell(r, c)
                cell.border, cell.alignment = border, wrap
                cell.font = bold if r == r1 else font
                if r == r1:
                    cell.fill = head_fill

    topics = {tp.get("week"): tp for tp in (course or {}).get("topics", []) or []}
    Lt = LETTERS.get(lang, LETTERS["ru"])
    wb = Workbook()
    ws = wb.active
    ws.title = t("bank", lang)
    head = ["ID", t("week", lang), t("topic", lang), "LO" if lang == "en" else "РО", t("q_type", lang),
            t("difficulty", lang), t("bloom_level", lang), t("points", lang), t("question", lang),
            t("options", lang), t("answer", lang)]
    ws.append(head)
    for q in bank["questions"]:
        if q["type"] in ("single", "multiple"):
            opts = "\n".join(f"{Lt[i]}) {o}" for i, o in enumerate(q["options"]))
        elif q["type"] == "matching":
            opts = "\n".join(f"{a} → {b}" for a, b in q["pairs"])
        else:
            opts = ""
        ws.append([q["id"], q["topic"], (topics.get(q["topic"]) or {}).get("title", ""), str(q["outcome"]),
                   tv(q["type"], lang), tv(q["difficulty"], lang), bloom_name(q["bloom"], lang),
                   q.get("points", 1), q["text"], opts, bank_key(q, lang)])
    n = len(bank["questions"])
    grid(ws, 1, n + 1, len(head))
    for i, w in enumerate([7, 8, 28, 7, 16, 11, 13, 7, 50, 40, 40], 1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{L(len(head))}{n + 1}"
    bank_sheet = ws.title

    # Покрытие — формулами СЧЁТЕСЛИ по банку
    ws = wb.create_sheet(t("coverage", lang))
    r = 1

    def block(title_row: list[str], keys: list, col: str, labels: list[str] | None = None):
        nonlocal r
        start = r
        ws.append(title_row + [t("count", lang)])
        for i, k in enumerate(keys):
            r += 1
            label = labels[i] if labels else k
            row = [k, label] if labels else [k]
            ws.append(row + [f"=COUNTIF('{bank_sheet}'!${col}:${col},{'A' + str(r)})"])
        grid(ws, start, r, len(title_row) + 1)
        ws.append([])
        r += 2

    tks = sorted(topics) if topics else sorted({q["topic"] for q in bank["questions"]})
    block([t("week", lang), t("topic", lang)], tks, "B", [(topics.get(k) or {}).get("title", "") for k in tks])
    outs = [str(o.get("id")) for o in (course or {}).get("outcomes", []) or []] or \
        sorted({str(q["outcome"]) for q in bank["questions"]})
    block(["LO" if lang == "en" else "РО"], outs, "D")
    block([t("q_type", lang)], [tv(x, lang) for x in TYPES], "E")
    block([t("difficulty", lang)], [tv(x, lang) for x in DIFFICULTY], "F")
    block([t("bloom_level", lang)], [bloom_name(x, lang) for x in BLOOM_LEVELS], "G")
    for i, w in enumerate([22, 40, 12], 1):
        ws.column_dimensions[L(i)].width = w

    # Варианты: порядок вопросов и ключ
    ws = wb.create_sheet(t("variants", lang))
    head = ["№"]
    for v in range(1, len(variants) + 1):
        head += [f'{t("variant", lang)} {v}: ID', t("answer", lang)]
    ws.append(head)
    for i in range(max(len(v) for v in variants)):
        row = [i + 1]
        for items in variants:
            row += [items[i]["q"]["id"], key_text(items[i], lang)] if i < len(items) else ["", ""]
        ws.append(row)
    grid(ws, 1, ws.max_row, len(head))
    ws.column_dimensions["A"].width = 5
    for c in range(2, len(head) + 1):
        ws.column_dimensions[L(c)].width = 14 if c % 2 == 0 else 30
    return wb


# ---------- Moodle ----------

def gift_escape(s) -> str:
    s = str(s).replace("\\", "\\\\")
    for ch in "~=#{}:":
        s = s.replace(ch, "\\" + ch)
    return s.replace("\r", "").replace("\n", "\\n")


def frac(x: float) -> str:
    return f"{x:.5f}".rstrip("0").rstrip(".")


def category_path(bank: dict, course: dict | None, q: dict, lang: str) -> str:
    root = course_title(course, lang) if course else (bank.get("title") or "tests")
    tp = next((x for x in (course or {}).get("topics", []) or [] if x.get("week") == q["topic"]), None)
    leaf = f'{t("topic", lang)} {int(q["topic"]):02d}' + (f'. {tp["title"]}' if tp else "")
    return "$course$/top/" + "/".join(p.replace("/", "//") for p in (root, leaf))


def build_gift(bank: dict, course: dict | None, lang: str) -> str:
    out = [f"// {bank.get('title') or t('test', lang)}" + (f" — {course_title(course, lang)}" if course else ""),
           "// Moodle GIFT, UTF-8", ""]
    current = None
    for q in sorted(bank["questions"], key=lambda x: (x["topic"], str(x["id"]))):
        cat = category_path(bank, course, q, lang)
        if cat != current:
            out += [f"$CATEGORY: {cat}", ""]
            current = cat
        head = f"::{gift_escape(q['id'])}:: {gift_escape(q['text'])}"
        typ = q["type"]
        if typ == "single":
            body = [f"\t={gift_escape(o)}" if o == str(q["answer"]) else f"\t~{gift_escape(o)}"
                    for o in map(str, q["options"])]
        elif typ == "multiple":
            keys = {str(a) for a in as_list(q["answer"])}
            k, n = len(keys), len(q["options"])
            body = [f"\t~%{frac(100 / k)}%{gift_escape(o)}" if o in keys
                    else f"\t~%-{frac(100 / (n - k))}%{gift_escape(o)}" for o in map(str, q["options"])]
        elif typ == "truefalse":
            out += [f"{head} {{{'TRUE' if q['answer'] else 'FALSE'}}}", ""]
            continue
        elif typ == "matching":
            body = [f"\t={gift_escape(a)} -> {gift_escape(b)}" for a, b in q["pairs"]]
        elif typ == "short":
            body = [f"\t={gift_escape(a)}" for a in as_list(q["answer"])]
        else:
            out += [f"{head} {{}}", ""]
            continue
        out += [head + " {", *body, "}", ""]
    return "\n".join(out)


def build_moodle_xml(bank: dict, course: dict | None, lang: str) -> bytes:
    quiz = ET.Element("quiz")

    def text_el(parent, tag, value, fmt="html"):
        el = ET.SubElement(parent, tag, {"format": fmt} if fmt else {})
        ET.SubElement(el, "text").text = html.escape(str(value)) if fmt == "html" else str(value)
        return el

    def answer(parent, fraction, value, fmt="html"):
        a = ET.SubElement(parent, "answer", {"fraction": fraction, "format": fmt})
        ET.SubElement(a, "text").text = html.escape(str(value)) if fmt == "html" else str(value)
        fb = ET.SubElement(a, "feedback", {"format": "html"})
        ET.SubElement(fb, "text").text = ""

    current = None
    mtype = {"single": "multichoice", "multiple": "multichoice", "truefalse": "truefalse",
             "matching": "matching", "short": "shortanswer", "open": "essay"}
    for q in sorted(bank["questions"], key=lambda x: (x["topic"], str(x["id"]))):
        cat = category_path(bank, course, q, lang)
        if cat != current:
            c = ET.SubElement(quiz, "question", {"type": "category"})
            ET.SubElement(ET.SubElement(c, "category"), "text").text = cat
            current = cat
        typ = q["type"]
        el = ET.SubElement(quiz, "question", {"type": mtype[typ]})
        ET.SubElement(ET.SubElement(el, "name"), "text").text = str(q["id"])
        text_el(el, "questiontext", q["text"])
        text_el(el, "generalfeedback", "")
        ET.SubElement(el, "defaultgrade").text = str(q.get("points", 1))
        ET.SubElement(el, "penalty").text = "0" if typ in ("truefalse", "open") else "0.3333333"
        ET.SubElement(el, "hidden").text = "0"
        if typ in ("single", "multiple"):
            keys = {str(a) for a in as_list(q["answer"])}
            k, n = len(keys), len(q["options"])
            ET.SubElement(el, "single").text = "true" if typ == "single" else "false"
            ET.SubElement(el, "shuffleanswers").text = "true"
            ET.SubElement(el, "answernumbering").text = "abc"
            for o in map(str, q["options"]):
                if typ == "single":
                    answer(el, "100" if o in keys else "0", o)
                else:
                    answer(el, frac(100 / k) if o in keys else "-" + frac(100 / (n - k)), o)
        elif typ == "truefalse":
            answer(el, "100" if q["answer"] else "0", "true", "moodle_auto_format")
            answer(el, "0" if q["answer"] else "100", "false", "moodle_auto_format")
        elif typ == "matching":
            ET.SubElement(el, "shuffleanswers").text = "true"
            for a, b in q["pairs"]:
                sq = text_el(el, "subquestion", a)
                ET.SubElement(ET.SubElement(sq, "answer"), "text").text = str(b)
        elif typ == "short":
            ET.SubElement(el, "usecase").text = "0"
            for a in as_list(q["answer"]):
                answer(el, "100", a, "moodle_auto_format")
        else:
            ET.SubElement(el, "responseformat").text = "editor"
            ET.SubElement(el, "responserequired").text = "1"
            ET.SubElement(el, "responsefieldlines").text = "15"
            ET.SubElement(el, "attachments").text = "0"
            text_el(el, "graderinfo", q["answer"])
            text_el(el, "responsetemplate", "")
        tags = ET.SubElement(el, "tags")
        for tag in (str(q["outcome"]), f"bloom-{q['bloom']}", f"difficulty-{q['difficulty']}"):
            ET.SubElement(ET.SubElement(tags, "tag"), "text").text = tag
    ET.indent(quiz)
    return b'<?xml version="1.0" encoding="UTF-8"?>\n' + ET.tostring(quiz, encoding="utf-8")


# ---------- запуск ----------

def main() -> None:
    ap = argparse.ArgumentParser(description="Тесты из tests.yaml: xlsx, бланки docx, Moodle GIFT/XML")
    ap.add_argument("tests", help="путь к tests.yaml")
    ap.add_argument("--course", help="путь к course.yaml (по умолчанию — рядом с tests.yaml)")
    ap.add_argument("--check", action="store_true", help="только проверить банк")
    args = ap.parse_args()
    path = Path(args.tests).expanduser().resolve()
    try:
        bank, course = load(path, args.course)
    except DataError as e:
        fail(str(e))
    errors, warnings = validate_bank(bank, course)
    if errors:
        report(errors, warnings)
        fail("тесты не собраны — исправьте ошибки в tests.yaml")
    q = bank["questions"]
    print(f"Банк: {len(q)} вопросов; типы: " + ", ".join(f"{k} {v}" for k, v in Counter(x['type'] for x in q).items()))
    if args.check:
        report([], warnings)
        return
    lang = bank.get("language") or (course or {}).get("language", "ru")
    out = (course["_dir"] / ((course.get("files") or {}).get("out", "out"))) if course else path.parent / "out"
    out.mkdir(parents=True, exist_ok=True)
    variants = make_variants(bank)
    files = []

    def save(obj, p: Path):
        try:
            if isinstance(obj, (str, bytes)):
                p.write_bytes(obj if isinstance(obj, bytes) else obj.encode("utf-8"))
            else:
                obj.save(p)
        except PermissionError:
            fail(f"не удалось записать {p} — закройте файл и запустите снова")
        files.append(p)

    save(build_xlsx(bank, course, variants, lang), out / "tests.xlsx")
    for v, items in enumerate(variants, 1):
        save(blank(bank, course, items, v, lang), out / f"tests_{v}.docx")
    save(key_doc(bank, course, variants, lang), out / "tests_key.docx")
    save(build_gift(bank, course, lang), out / "tests.gift")
    save(build_moodle_xml(bank, course, lang), out / "tests_moodle.xml")
    if warnings:
        report([], warnings)
    for p in files:
        print(f"Готово: {p}")


if __name__ == "__main__":
    main()
