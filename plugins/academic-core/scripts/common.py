"""Общий код скриптов academic-core: данные курса, профиль, настройки вуза, оформление docx.

Скрипты импортируют его так:
    sys.path.insert(0, str(Path(__file__).parent)); from common import ...
Зависимости (python-docx, pyyaml) объявляет сам скрипт в PEP 723.
"""
from __future__ import annotations

import sys
from pathlib import Path

import yaml

from i18n import t, tv
from settings import DataError, lang_text, load_settings, logo_path
from settings import university_name as settings_university_name

HOUR_KINDS = ("lectures", "practicals", "labs", "srsp", "srs")

# Windows: вывод в канал по умолчанию в cp1251 — печатаем всегда в UTF-8.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass


def fail(msg: str) -> None:
    print(f"Ошибка: {msg}", file=sys.stderr)
    sys.exit(1)


def read_yaml(path: Path) -> dict:
    try:
        return yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except FileNotFoundError:
        raise DataError(f"не найден файл {path}")
    except yaml.YAMLError as e:
        raise DataError(f"в {path.name} ошибка разметки YAML: {e}")


def find_up(start: Path, name: str) -> Path | None:
    """Ищет файл или папку name в start и выше по дереву (рабочая папка)."""
    for d in [start, *start.parents]:
        if (d / name).exists():
            return d / name
    return None


def university_context(start: Path) -> dict:
    """Настройки вуза для папки start: _settings, _brand (name, name_ky, name_en, logo, color), _brand_dir."""
    s = load_settings(start)
    u, b = s["university"], s["brand"]
    brand = {k: u[k] for k in ("name", "name_ky", "name_en") if u.get(k)}
    brand.update({k: b[k] for k in ("logo", "color") if b.get(k)})
    return {"_settings": s, "_brand": brand, "_brand_dir": s["_dir"]}


def settings(data: dict | None) -> dict:
    """Настройки вуза из данных курса; нет — нормы КР."""
    s = (data or {}).get("_settings")
    return s if s is not None else load_settings(None)


def load_course(course_yaml: str | Path) -> dict:
    """Паспорт курса + профиль + настройки вуза в одном словаре.

    Возвращает dict с ключами course-паспорта и служебными:
    _dir (папка курса), _profile, _settings (university/settings.yaml поверх норм КР),
    _brand, _brand_dir.
    """
    path = Path(course_yaml).expanduser().resolve()
    data = read_yaml(path)
    data["_dir"] = path.parent
    prof = find_up(path.parent, "profile.yaml")
    data["_profile"] = read_yaml(prof) if prof else {}
    data.update(university_context(path.parent))
    s = data["_settings"]
    teacher = data.setdefault("teacher", {}) or {}
    for k in ("name", "degree", "position", "email"):
        if not teacher.get(k):
            teacher[k] = data["_profile"].get(k, "")
    data["teacher"] = teacher
    data.setdefault("language", data["_profile"].get("language") or s["environment"]["languages"][0])
    pol = data.get("policies")
    if isinstance(pol, dict) or pol is None:  # политика вуза — где курс её не задал
        pol = dict(pol or {})
        for k, v in s["syllabus"]["policies"].items():
            if pol.get(k) or not v:
                continue
            if isinstance(v, dict):
                pol.update({f"{k}_{lg}": x for lg, x in v.items()
                            if x and lg != data["language"] and not pol.get(f"{k}_{lg}")})
                v = v.get(data["language"]) or v.get("ru") or next(x for x in v.values() if x)
            pol[k] = v
        if pol:
            data["policies"] = pol
    return data


def credit_hours(data: dict | None) -> int:
    """Часов в одном кредите: settings.yaml → grading.credit_hours, иначе 30."""
    return int(settings(data)["grading"]["credit_hours"])


def check_hours(data: dict) -> list[str]:
    """Сверка часов: сумма видов = кредиты × часы кредита, сумма по темам = итог по видам."""
    problems = []
    hours = data.get("hours", {}) or {}
    credits = data.get("credits")
    ch = credit_hours(data)
    total = sum(int(hours.get(k, 0) or 0) for k in HOUR_KINDS)
    if credits and total != credits * ch:
        problems.append(
            f"часы по видам дают {total}, а {credits} кр. × {ch} = {credits * ch}"
        )
    topics = data.get("topics", []) or []
    kinds = ["lectures", "practicals", "labs", "srs"]
    if any("srsp" in tp for tp in topics):
        kinds.insert(3, "srsp")
    for k in kinds:
        if k in hours and topics:
            s = sum(int(tp.get(k, 0) or 0) for tp in topics)
            if s != int(hours[k] or 0):
                problems.append(f"{k}: по темам {s} ч, в паспорте {hours[k]} ч")
    weights = sum(int(a.get("weight", 0) or 0) for a in data.get("assessment", []) or [])
    if data.get("assessment") and weights != 100:
        problems.append(f"веса оценивания дают {weights}, а нужно 100")
    return problems


def university_name(data: dict, lang: str) -> str:
    """Вуз: settings.yaml / brand.yaml на языке lang, иначе profile.yaml → university."""
    if data.get("_settings") is not None:
        name = settings_university_name(data["_settings"], lang)
    else:
        b = data.get("_brand", {})
        name = b.get({"ru": "name", "ky": "name_ky", "en": "name_en"}.get(lang, "name")) or b.get("name")
    return name or data.get("_profile", {}).get("university", "")


def city_name(data: dict) -> str:
    """Город на титуле: settings.yaml → university.city, иначе profile.yaml → city."""
    return settings(data)["university"].get("city") or (data.get("_profile") or {}).get("city", "")


def ministry_name(data: dict, lang: str) -> str:
    """Первая строка титула: settings.yaml → university.ministry, иначе Минобрнауки КР."""
    return lang_text(settings(data)["university"].get("ministry"), lang, t("ministry", lang))


def approval_role(data: dict, lang: str) -> str:
    """Кто утверждает: settings.yaml → approval.role, иначе заведующий кафедрой."""
    return lang_text(settings(data)["approval"].get("role"), lang, t("approve_role", lang))


def loc(obj: dict, key: str, lang: str, data: dict | None = None):
    """Поле на языке lang: key_<lang>, иначе key (язык паспорта).

    Если язык документа не совпадает с языком паспорта и перевода нет —
    запоминает поле в data["_missing_tr"], чтобы скрипт предупредил.
    """
    if not isinstance(obj, dict):
        return obj
    base = (data or {}).get("language", "ru")
    if lang != base:
        v = obj.get(f"{key}_{lang}")
        if v:
            return v
        if obj.get(key) and data is not None:
            data.setdefault("_missing_tr", set()).add(key)
    return obj.get(key)


def course_title(data: dict, lang: str) -> str:
    c = data.get("course", {})
    return c.get(f"title_{lang}") or c.get("title", "")


def out_dir(data: dict) -> Path:
    d = data["_dir"] / (data.get("files", {}) or {}).get("out", "out")
    d.mkdir(parents=True, exist_ok=True)
    return d


# ---------- docx по ГОСТ ----------

def new_document(data: dict | None = None):
    """Пустой docx по settings.yaml → document; по умолчанию Times New Roman 14, интервал 1,5,
    поля 30/15/20/20, номера страниц внизу по центру. data — данные курса или настройки вуза."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Mm, Pt

    ds = (data if data and {"document", "grading"} <= set(data) else settings(data))["document"]
    font, size, m = ds["font"], ds["size"], ds["margins_mm"]
    doc = Document()
    doc._table_size = ds["table_size"]
    for s in doc.sections:
        s.left_margin, s.right_margin = Mm(m["left"]), Mm(m["right"])
        s.top_margin, s.bottom_margin = Mm(m["top"]), Mm(m["bottom"])
        if ds["page_numbers"] == "none":
            continue
        p = (s.header if ds["page_numbers"].startswith("top") else s.footer).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if ds["page_numbers"].endswith("right") else WD_ALIGN_PARAGRAPH.CENTER
        # поле PAGE: begin, код, separate, текущее значение, end — каждое в своём run
        for kind, text in (("begin", None), (None, "PAGE"), ("separate", None), ("value", "1"), ("end", None)):
            run = p.add_run("1" if kind == "value" else None)
            if kind in ("begin", "separate", "end"):
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
                run._r.append(el)
            elif kind is None:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
                run._r.append(el)

    st = doc.styles["Normal"]
    st.font.name = font
    st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        st.element.rPr.rFonts.attrib.pop(qn(attr), None)
    pf = st.paragraph_format
    pf.line_spacing = ds["line_spacing"]
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(ds["indent_cm"])
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for lvl in (1, 2, 3):
        h = doc.styles[f"Heading {lvl}"]
        h.font.name = font
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = None
        rf = h.element.rPr.rFonts
        rf.set(qn("w:eastAsia"), font)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(attr), None)  # иначе Word/LibreOffice берут шрифт темы (Calibri Light)
        h.paragraph_format.first_line_indent = Cm(0)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if lvl == 1 else WD_ALIGN_PARAGRAPH.LEFT
        )
        h.paragraph_format.keep_with_next = True
    return doc


def para(doc, text: str = "", *, bold=False, italic=False, align=None, indent=True, size=None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
    if align:
        p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                       "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    if not indent or align:
        p.paragraph_format.first_line_indent = Cm(0)
    return p


def table(doc, header: list[str], rows: list[list], widths_cm: list[float] | None = None):
    """Таблица 12 pt (settings.yaml → document.table_size), интервал 1,0, шапка жирная, текст — влево."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    tbl = doc.add_table(rows=1, cols=len(header))
    tbl.style = "Table Grid"
    for i, h in enumerate(header):
        tbl.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = tbl.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    hdr = OxmlElement("w:tblHeader")  # шапка повторяется на каждой странице
    hdr.set(qn("w:val"), "true")
    tbl.rows[0]._tr.get_or_add_trPr().append(hdr)
    if widths_cm:
        tbl.autofit = False  # без этого Word и LibreOffice делят ширину поровну
        for ci, col in enumerate(tbl.columns):
            col.width = Cm(widths_cm[ci])
    for ri, r in enumerate(tbl.rows):
        r._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))  # строка не рвётся между страницами
        for ci, c in enumerate(r.cells):
            if widths_cm:
                c.width = Cm(widths_cm[ci])
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size = Pt(getattr(doc, "_table_size", 12))
                    run.bold = ri == 0
    doc.add_paragraph()
    return tbl


def approval_block(doc, data: dict, lang: str) -> None:
    """«УТВЕРЖДАЮ» справа (settings.yaml → approval); с approval.agreed — «СОГЛАСОВАНО» слева."""
    from docx.shared import Cm

    a = settings(data)["approval"]
    date = f'«___» ____________ {str(data.get("course", {}).get("academic_year", ""))[:4]} г.'
    approve = [t("approve", lang), approval_role(data, lang),
               f'________________ {a["name"]}' if a.get("name") else "________________________", date]
    if not a.get("agreed"):
        for line in approve:
            para(doc, line, align="right")
        return
    agreed = [t("agreed", lang)]
    for x in a["agreed"]:
        agreed += [lang_text(x.get("role"), lang), f'________________ {x.get("name") or ""}'.rstrip(), date]
    tbl = doc.add_table(rows=1, cols=2)
    tbl.autofit = False
    for cell, lines in zip(tbl.rows[0].cells, (agreed, approve)):
        cell.width = Cm(8.25)
        for i, line in enumerate(lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.add_run(line)
            p.paragraph_format.first_line_indent = Cm(0)
            p.alignment = 0


def title_page(doc, data: dict, lang: str, doc_title: str, *, approve: bool = True):
    """Титульный лист: министерство, вуз, кафедра, «УТВЕРЖДАЮ», название, город, год.

    Министерство, вуз, город, логотип и подписи — из settings.yaml (university, approval,
    document.title_page_logo), иначе — нормы КР и profile.yaml.
    """
    from docx.enum.text import WD_BREAK
    from docx.shared import Cm

    prof = data.get("_profile", {})
    s = settings(data)
    logo = logo_path(s) if s["document"]["title_page_logo"] else None
    has_logo = logo is not None
    if has_logo:
        doc.add_picture(str(logo), width=Cm(3))
        doc.paragraphs[-1].alignment = 1

    def gap(n: int) -> None:
        """Пустые строки титула — одинарным интервалом; с логотипом их меньше, чтобы титул влез на лист."""
        for _ in range(max(1, n // 2) if has_logo else n):
            para(doc).paragraph_format.line_spacing = 1.0
    para(doc, ministry_name(data, lang), align="center")
    para(doc, university_name(data, lang), bold=True, align="center")
    if prof.get("faculty"):
        para(doc, prof["faculty"], align="center")
    if prof.get("department"):
        para(doc, f'{t("department", lang)} «{prof["department"]}»', align="center")
    para(doc)
    if approve and s["approval"]["enabled"]:
        approval_block(doc, data, lang)
    gap(4)
    para(doc, doc_title.upper(), bold=True, align="center")
    para(doc, f'{t("discipline", lang)}: «{course_title(data, lang)}»', bold=True, align="center")
    c = data.get("course", {})
    meta = [c.get("program"), f'{tv(c.get("level", ""), lang)}, {t("course_year", lang)} {c.get("year", "")}, '
            f'{t("semester", lang)} {c.get("semester", "")}', f'{t("credits", lang)}: {data.get("credits", "")}']
    for m in meta:
        if m:
            para(doc, m, align="center")
    gap(6)
    city = city_name(data)
    para(doc, f'{city} {str(c.get("academic_year", ""))[:4]}'.strip(), align="center")
    doc.paragraphs[-1].runs[-1].add_break(WD_BREAK.PAGE)
