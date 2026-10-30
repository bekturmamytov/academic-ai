# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1", "python-docx>=1.1", "pyyaml>=6"]
# ///
"""РУП дисциплины из паспорта курса → out/rup.xlsx (+ out/rup_<lang>.docx с --docx).

    uv run build_rup.py "<папка курса>/course.yaml" [--docx] [--lang ru|ky|en] [--final]

rup.xlsx: «Распределение часов», «Тематический план», «Календарный график контроля».
Итоги — формулами Excel (СУММ), чтобы правка часов в таблице сразу пересчитывалась.
Рабочая программа (docx) — структура в skills/rup/references/structure.md.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from common import CREDIT_HOURS, HOUR_KINDS, course_title, loc, new_document, out_dir, para, title_page  # noqa: E402
from course_check import assessment_plan, report  # noqa: E402
from course_doc import (  # noqa: E402
    DASH, heading, kv_table, load_checked, save, sec_assessment, sec_goals, sec_literature, sec_outcomes, sec_plan,
    sec_signatures, sec_srs, topics_sorted, translation_warning,
)
from i18n import LABELS, t, tv  # noqa: E402

LABELS.update({
    "sh_hours": {"ru": "Распределение часов", "ky": "Сааттарды бөлүштүрүү", "en": "Workload"},
    "sh_plan": {"ru": "Тематический план", "ky": "Тематикалык план", "en": "Course schedule"},
    "sh_control": {"ru": "Календарный график контроля", "ky": "Көзөмөлдүн календардык графиги",
                   "en": "Assessment calendar"},
    "kind": {"ru": "Вид учебной работы", "ky": "Окуу ишинин түрү", "en": "Type of work"},
    "share": {"ru": "Доля", "ky": "Үлүшү", "en": "Share"},
    "credits_calc": {"ru": "Кредиты (итого / 30)", "ky": "Кредиттер (бардыгы / 30)", "en": "Credits (total / 30)"},
    "credits_passport": {"ru": "Кредиты по учебному плану", "ky": "Окуу планы боюнча кредиттер",
                         "en": "Credits in the curriculum"},
    "check": {"ru": "Сверка", "ky": "Салыштыруу", "en": "Check"},
    "ok": {"ru": "сходится", "ky": "дал келет", "en": "OK"},
    "bad": {"ru": "НЕ сходится", "ky": "дал келбейт", "en": "MISMATCH"},
    "by_passport": {"ru": "По паспорту", "ky": "Паспорт боюнча", "en": "Course passport"},
    "lab_or_practice": {"ru": "Лабораторная / практическое", "ky": "Лабораториялык / практикалык",
                        "en": "Lab / seminar"},
    "control": {"ru": "Контроль", "ky": "Көзөмөл", "en": "Assessment"},
    "nn": {"ru": "№", "ky": "№", "en": "No."},
})


# ---------- xlsx ----------

def build_xlsx(data: dict, lang: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    total_fill = PatternFill("solid", fgColor="F2F2F2")
    bold = Font(name="Times New Roman", bold=True, size=11)
    normal = Font(name="Times New Roman", size=11)
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="top", wrap_text=True)

    def style_range(ws, r1, r2, c1, c2, *, head=False, total=False):
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                cell = ws.cell(r, c)
                cell.border = border
                cell.font = bold if (head or total) else normal
                cell.alignment = center if (head or c > 2 or isinstance(cell.value, (int, float))) else wrap
                if head:
                    cell.fill = head_fill
                elif total:
                    cell.fill = total_fill

    def title(ws, text, ncols):
        ws.cell(1, 1, text).font = Font(name="Times New Roman", bold=True, size=13)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        c = data.get("course") or {}
        sub = f'{c.get("code", "")} {course_title(data, lang)} · {t("credits", lang)}: {data.get("credits")} · ' \
              f'{c.get("academic_year", "")}'
        ws.cell(2, 1, sub.strip()).font = normal
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)

    def page_setup(ws, landscape=False):
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb = Workbook()
    hours = data.get("hours") or {}

    # 1. Распределение часов
    ws = wb.active
    ws.title = t("sh_hours", lang)[:31]
    title(ws, t("sh_hours", lang), 3)
    ws.append([])
    hdr = 4
    for col, v in enumerate([t("kind", lang), t("hours", lang), t("share", lang)], 1):
        ws.cell(hdr, col, v)
    first = hdr + 1
    for i, k in enumerate(HOUR_KINDS):
        ws.cell(first + i, 1, t(k, lang))
        ws.cell(first + i, 2, int(hours.get(k, 0) or 0))
    last = first + len(HOUR_KINDS) - 1
    r_total = last + 3
    for i in range(len(HOUR_KINDS)):
        ws.cell(first + i, 3, f"=IF($B${r_total}=0,0,B{first + i}/$B${r_total})").number_format = "0%"
    r_aud = last + 1
    ws.cell(r_aud, 1, t("contact_hours", lang))
    ws.cell(r_aud, 2, f"=SUM(B{first}:B{first + 2})")
    ws.cell(r_aud, 3, f"=IF($B${r_total}=0,0,B{r_aud}/$B${r_total})").number_format = "0%"
    ws.cell(r_aud + 1, 1, t("self_hours", lang))
    ws.cell(r_aud + 1, 2, f"=SUM(B{first + 3}:B{last})")
    ws.cell(r_aud + 1, 3, f"=IF($B${r_total}=0,0,B{r_aud + 1}/$B${r_total})").number_format = "0%"
    ws.cell(r_total, 1, t("total", lang))
    ws.cell(r_total, 2, f"=SUM(B{first}:B{last})")
    ws.cell(r_total + 1, 1, t("credits_calc", lang))
    ws.cell(r_total + 1, 2, f"=B{r_total}/{CREDIT_HOURS}").number_format = "0.##"
    ws.cell(r_total + 2, 1, t("credits_passport", lang))
    ws.cell(r_total + 2, 2, int(data.get("credits") or 0))
    ws.cell(r_total + 3, 1, t("check", lang))
    ws.cell(r_total + 3, 2, f'=IF(B{r_total}=B{r_total + 2}*{CREDIT_HOURS},"{t("ok", lang)}","{t("bad", lang)}")')
    style_range(ws, hdr, hdr, 1, 3, head=True)
    style_range(ws, first, r_total + 3, 1, 3)
    style_range(ws, r_total, r_total, 1, 3, total=True)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 10
    page_setup(ws)
    hours_row = {k: first + i for i, k in enumerate(HOUR_KINDS)}
    hours_sheet = ws.title

    # 2. Тематический план
    ws = wb.create_sheet(t("sh_plan", lang)[:31])
    kinds = list(HOUR_KINDS)
    head = [t("week", lang), t("topic", lang)] + [t(f"a_{k}", lang) for k in kinds] + \
           [t("total", lang), "LO" if lang == "en" else "РО", t("lab_or_practice", lang), t("control", lang)]
    ncol = len(head)
    title(ws, t("sh_plan", lang), ncol)
    hdr = 4
    for col, v in enumerate(head, 1):
        ws.cell(hdr, col, v)
    c_first, c_last = 3, 2 + len(kinds)
    L = get_column_letter
    r = hdr
    for tp in topics_sorted(data):
        r += 1
        ws.cell(r, 1, tp.get("week"))
        ws.cell(r, 2, loc(tp, "title", lang, data))
        for j, k in enumerate(kinds):
            ws.cell(r, c_first + j, int(tp.get(k, 0) or 0))
        ws.cell(r, c_last + 1, f"=SUM({L(c_first)}{r}:{L(c_last)}{r})")
        ws.cell(r, c_last + 2, ", ".join(map(str, tp.get("outcomes", []) or [])))
        ws.cell(r, c_last + 3, loc(tp, "lab", lang, data) or loc(tp, "practical", lang, data) or "")
        ws.cell(r, c_last + 4, loc(tp, "control", lang, data) or "")
    data_first, data_last = hdr + 1, r
    r_tot = r + 1
    ws.cell(r_tot, 2, t("total", lang))
    for c in range(c_first, c_last + 2):
        ws.cell(r_tot, c, f"=SUM({L(c)}{data_first}:{L(c)}{data_last})")
    r_pas = r_tot + 1
    ws.cell(r_pas, 2, t("by_passport", lang))
    for j, k in enumerate(kinds):
        ws.cell(r_pas, c_first + j, f"='{hours_sheet}'!B{hours_row[k]}")
    ws.cell(r_pas, c_last + 1, f"=SUM({L(c_first)}{r_pas}:{L(c_last)}{r_pas})")
    r_chk = r_pas + 1
    ws.cell(r_chk, 2, t("check", lang))
    for c in range(c_first, c_last + 2):
        col = L(c)
        # СРСП бывает не расписана по темам — тогда сверка по ней не нужна
        ws.cell(r_chk, c, f'=IF({col}{r_tot}={col}{r_pas},"✓",IF({col}{r_tot}=0,"—","≠"))')
    style_range(ws, hdr, hdr, 1, ncol, head=True)
    style_range(ws, data_first, r_chk, 1, ncol)
    style_range(ws, r_tot, r_tot, 1, ncol, total=True)
    for c in (c_last + 3, c_last + 4):
        for rr in range(data_first, data_last + 1):
            ws.cell(rr, c).alignment = wrap
    widths = [8, 46] + [7] * len(kinds) + [8, 12, 34, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = ws.cell(hdr + 1, 3)
    page_setup(ws, landscape=True)

    # 3. Календарный график контроля
    ws = wb.create_sheet(t("sh_control", lang)[:31])
    n_weeks = int(data.get("weeks", 15) or 15)
    head = [t("control_kind", lang), t("q_type", lang), t("control_form", lang)] + \
           [str(w) for w in range(1, n_weeks + 1)] + [t("session", lang), t("total", lang)]
    ncol = len(head)
    title(ws, t("sh_control", lang), ncol)
    hdr = 4
    ws.cell(hdr - 1, 4, t("week", lang)).font = bold
    for col, v in enumerate(head, 1):
        ws.cell(hdr, col, v)
    w_first, w_sess = 4, 4 + n_weeks
    r = hdr
    for row in assessment_plan(data):
        r += 1
        a = row["item"]
        ws.cell(r, 1, loc(a, "name", lang, data))
        ws.cell(r, 2, tv(row["type"], lang))
        ws.cell(r, 3, loc(a, "form", lang, data) or "")
        for wk, pts in row["weeks"].items():
            ws.cell(r, w_first + wk - 1, pts)
        if row["session"]:
            ws.cell(r, w_sess, row["session"])
        ws.cell(r, w_sess + 1, f"=SUM({L(w_first)}{r}:{L(w_sess)}{r})")
    data_first, data_last = hdr + 1, r
    r_tot = r + 1
    ws.cell(r_tot, 1, t("total", lang))
    for c in range(w_first, w_sess + 2):
        ws.cell(r_tot, c, f"=SUM({L(c)}{data_first}:{L(c)}{data_last})")
    r_chk = r_tot + 1
    ws.cell(r_chk, 1, t("check", lang))
    tot_cell = f"{L(w_sess + 1)}{r_tot}"
    ws.cell(r_chk, w_sess + 1, f'=IF(ROUND({tot_cell},2)=100,"{t("ok", lang)}","{t("bad", lang)}")')
    style_range(ws, hdr, hdr, 1, ncol, head=True)
    style_range(ws, data_first, r_chk, 1, ncol)
    style_range(ws, r_tot, r_tot, 1, ncol, total=True)
    for rr in range(data_first, r_chk + 1):
        for c in range(w_first, w_sess + 2):
            ws.cell(rr, c).number_format = "0.##"
    for i, w in enumerate([30, 12, 24] + [5.5] * n_weeks + [9, 8], 1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = ws.cell(hdr + 1, 4)
    page_setup(ws, landscape=True)
    return wb


# ---------- docx: рабочая программа дисциплины ----------

def build_docx(data: dict, lang: str, final: bool):
    doc = new_document()
    title_page(doc, data, lang, t("rup", lang))
    c = data.get("course") or {}
    hours = data.get("hours") or {}

    def place():
        para(doc, t("place_text", lang).format(cycle=tv(c.get("cycle"), lang) or DASH,
                                               type=tv(c.get("type"), lang) or DASH))
        prog = loc(c, "program", lang, data)
        if prog:
            para(doc, f'{t("f_program", lang)}: {prog}.')
        pre = loc(c, "prerequisites", lang, data) or []
        post = loc(c, "postrequisites", lang, data) or []
        para(doc, f'{t("prerequisites", lang)}: {", ".join(pre) or DASH}.')
        para(doc, f'{t("postrequisites", lang)}: {", ".join(post) or DASH}.')

    def volume():
        total = sum(int(hours.get(k, 0) or 0) for k in HOUR_KINDS)
        aud = sum(int(hours.get(k, 0) or 0) for k in ("lectures", "practicals", "labs"))
        kv_table(doc, [(t(k, lang), str(hours.get(k, 0) or 0)) for k in HOUR_KINDS] + [
            (t("contact_hours", lang), str(aud)),
            (t("all_hours", lang), str(total)),
            (t("credits", lang), str(data.get("credits", ""))),
        ])

    def content():
        for tp in topics_sorted(data):
            p = para(doc, indent=False)
            p.add_run(f'{t("topic", lang)} {tp.get("week")}. {loc(tp, "title", lang, data)}').bold = True
            p.paragraph_format.keep_with_next = True
            if tp.get("content"):
                para(doc, str(loc(tp, "content", lang, data)))
            for key, lab in (("lab", "lab_work"), ("practical", "practical_work")):
                if tp.get(key):
                    p = para(doc)
                    p.add_run(f'{t(lab, lang)}: ').italic = True
                    p.add_run(str(loc(tp, key, lang, data)))

    sections = [
        ("sec_goals", lambda: sec_goals(doc, data, lang)),
        ("sec_place", place),
        ("sec_outcomes", lambda: sec_outcomes(doc, data, lang)),
        ("sec_volume", volume),
        ("sec_content", content),
        ("sec_plan", lambda: sec_plan(doc, data, lang)),
        ("sec_support", lambda: sec_srs(doc, data, lang)),
        ("sec_assess_tools", lambda: sec_assessment(doc, data, lang)),
        ("literature", lambda: sec_literature(doc, data, lang, final)),
        ("sec_signatures", lambda: sec_signatures(doc, data, lang)),
    ]
    for n, (key, render) in enumerate(sections, 1):
        heading(doc, n, key, lang)
        render()
    return doc


def main() -> None:
    ap = argparse.ArgumentParser(description="РУП дисциплины из course.yaml")
    ap.add_argument("course", help="путь к course.yaml")
    ap.add_argument("--docx", action="store_true", help="также рабочая программа дисциплины в docx")
    ap.add_argument("--lang", choices=["ru", "ky", "en"], help="язык (по умолчанию — язык паспорта)")
    ap.add_argument("--final", action="store_true", help="чистовик: без непроверенной литературы (для --docx)")
    args = ap.parse_args()
    data, warnings = load_checked(args.course, "РУП")
    lang = args.lang or data.get("language", "ru")
    out = out_dir(data)
    xlsx = out / "rup.xlsx" if lang == data.get("language", "ru") else out / f"rup_{lang}.xlsx"
    save(build_xlsx(data, lang), xlsx)
    paths = [xlsx]
    if args.docx:
        docx_path = out / f"rup_{lang}.docx"
        save(build_docx(data, lang, args.final), docx_path)
        paths.append(docx_path)
    translation_warning(data, lang, warnings)
    if warnings:
        report([], warnings)
    for p in paths:
        print(f"Готово: {p}")


if __name__ == "__main__":
    main()
