# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1", "pypdf>=4", "pyyaml>=6"]
# ///
"""Текст старого силлабуса / программы (docx, pdf) для переноса в course.yaml.

    uv run extract_text.py "<файл.docx|.pdf>" [--out текст.md] [--layout]

--layout — для образца вуза (docx): поля, шрифт, кегль, интервал, отступ, номера страниц
перед текстом — чтобы заполнить university/settings.yaml → document.

Абзацы и таблицы docx — в порядке документа, таблицы строками «ячейка | ячейка».
Скан PDF без текстового слоя не читается — скрипт об этом скажет.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from common import fail  # noqa: E402


def docx_text(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    out = []
    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, doc)
            text = p.text.strip()
            if text:
                style = (p.style.name or "") if p.style is not None else ""
                out.append(("## " if style.startswith("Heading") else "") + text)
        elif tag == "tbl":
            out.append("")
            for row in Table(el, doc).rows:
                cells, prev = [], None
                for c in row.cells:  # объединённые ячейки python-docx повторяет — убираем дубли
                    if c._tc is not prev:
                        cells.append(c.text.strip().replace("\n", " "))
                    prev = c._tc
                out.append(" | ".join(cells))
            out.append("")
    return "\n".join(out)


def docx_layout(path: Path) -> str:
    """Оформление образца: поля первого раздела, стиль «Обычный», где номер страницы."""
    import re

    from docx import Document

    doc = Document(str(path))
    s = doc.sections[0]
    mm = lambda v: f"{v / 36000:.0f}" if v is not None else "?"  # noqa: E731
    st = doc.styles["Normal"]
    pf = st.paragraph_format
    fonts = {r.font.name for p in doc.paragraphs for r in p.runs if r.font.name and r.text.strip()}
    sizes = {r.font.size.pt for p in doc.paragraphs for r in p.runs if r.font.size and r.text.strip()}
    pos = "нет"
    for where, part in (("bottom", s.footer), ("top", s.header)):
        for p in part.paragraphs:
            if re.search(r"instrText[^>]*>\s*PAGE\b|w:instr=\"\s*PAGE", p._p.xml):
                side = {1: "center", 2: "right"}.get(int(p.alignment or 0), "left")
                pos = f"{where}-{side}"
    lines = [
        f"поля, мм: левое {mm(s.left_margin)}, правое {mm(s.right_margin)}, "
        f"верхнее {mm(s.top_margin)}, нижнее {mm(s.bottom_margin)}",
        f"стиль «Обычный»: {st.font.name or '?'} {st.font.size.pt if st.font.size else '?'} pt, "
        f"интервал {pf.line_spacing or '?'}, отступ {f'{pf.first_line_indent.cm:.2f} см' if pf.first_line_indent else '?'}",
        f"шрифты в тексте: {', '.join(sorted(fonts)) or 'как в стиле'}; кегли: "
        f"{', '.join(f'{x:g}' for x in sorted(sizes)) or 'как в стиле'}",
        f"номер страницы: {pos}",
    ]
    return "<!-- оформление образца -->\n" + "\n".join(lines) + "\n<!-- конец оформления -->\n"


def pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    pages = [(p.extract_text() or "").strip() for p in PdfReader(str(path)).pages]
    if not any(pages):
        fail(f"в {path.name} нет текстового слоя (скан). Откройте PDF через Read или пришлите docx")
    return "\n\n".join(f"<!-- стр. {i} -->\n{t}" for i, t in enumerate(pages, 1))


def main() -> None:
    ap = argparse.ArgumentParser(description="Текст из docx/pdf")
    ap.add_argument("file")
    ap.add_argument("--out", help="записать в файл вместо вывода на экран")
    ap.add_argument("--layout", action="store_true", help="docx: сначала оформление (поля, шрифт, интервал)")
    args = ap.parse_args()
    path = Path(args.file).expanduser()
    if not path.exists():
        fail(f"не найден файл {path}")
    ext = path.suffix.lower()
    if ext == ".docx":
        text = (docx_layout(path) + "\n" if args.layout else "") + docx_text(path)
    elif ext == ".pdf":
        text = pdf_text(path)
    elif ext == ".doc":
        fail("старый формат .doc не читается — сохраните файл в Word как .docx")
    else:
        fail(f"формат {ext} не поддерживается — нужен .docx или .pdf")
    if args.out:
        Path(args.out).write_text(text, encoding="utf-8")
        print(f"Готово: {args.out}")
    else:
        print(text)


if __name__ == "__main__":
    main()
