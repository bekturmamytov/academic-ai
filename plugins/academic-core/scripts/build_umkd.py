# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1", "pyyaml>=6", "docxcompose>=1.4"]
# ///
"""УМКД дисциплины: собирает готовые материалы курса в один docx и дописывает простые разделы.

    uv run build_umkd.py <course.yaml>

Берёт: out/syllabus_<язык>.docx, lectures/*.md, srs.yaml, out/tests*.docx, umkd.yaml (глоссарий,
вопросы, методические указания — необязательно). Печатает, каких разделов нет и чем их сделать.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from common import (  # noqa: E402
    CREDIT_HOURS, DataError, check_hours, course_title, fail, load_course, new_document, out_dir, para, table,
    read_yaml, title_page,
)
from content_kit import (  # noqa: E402
    KR_SCALE, add_runs, add_toc, parse_md, passport_literature,
    render_docx, runs_text, section_items, split_frontmatter,
)
from i18n import LABELS, t  # noqa: E402

LABELS.update({
    "umkd": {"ru": "Учебно-методический комплекс дисциплины", "ky": "Сабактын окуу-методикалык комплекси",
             "en": "Course teaching and learning package"},
    "u_syllabus": {"ru": "Силлабус", "ky": "Силлабус", "en": "Syllabus"},
    "u_program": {"ru": "Рабочая программа дисциплины", "ky": "Сабактын жумушчу программасы", "en": "Course programme"},
    "u_lectures": {"ru": "Конспект лекций", "ky": "Лекциялардын конспекти", "en": "Lecture notes"},
    "u_practice": {"ru": "Планы практических и лабораторных занятий", "ky": "Практикалык жана лабораториялык сабактардын пландары",
                   "en": "Seminar and lab plans"},
    "u_srs": {"ru": "Самостоятельная работа студентов", "ky": "Студенттердин өз алдынча иши", "en": "Self-study"},
    "u_assessment": {"ru": "Оценочные средства", "ky": "Баалоо каражаттары", "en": "Assessment materials"},
    "u_tests": {"ru": "Тестовые задания", "ky": "Тест тапшырмалары", "en": "Tests"},
    "u_midterm": {"ru": "Вопросы к рубежному контролю", "ky": "Аралык текшерүүнүн суроолору", "en": "Midterm questions"},
    "u_exam": {"ru": "Вопросы к экзамену", "ky": "Экзамендин суроолору", "en": "Exam questions"},
    "u_glossary": {"ru": "Глоссарий", "ky": "Глоссарий", "en": "Glossary"},
    "u_guidelines": {"ru": "Методические указания для студентов", "ky": "Студенттер үчүн методикалык көрсөтмөлөр",
                     "en": "Guidance for students"},
    "u_approval": {"ru": "Лист согласования", "ky": "Макулдашуу баракчасы", "en": "Approval sheet"},
    "u_goals": {"ru": "Цели дисциплины", "ky": "Сабактын максаттары", "en": "Course aims"},
    "place": {"ru": "Место дисциплины в образовательной программе", "ky": "Сабактын билим берүү программасындагы орду",
              "en": "Place in the programme"},
    "volume": {"ru": "Объём дисциплины", "ky": "Сабактын көлөмү", "en": "Workload"},
    "plan": {"ru": "Тематический план", "ky": "Тематикалык план", "en": "Topic plan"},
    "u_methods": {"ru": "Образовательные технологии и устойчивое развитие", "ky": "Билим берүү технологиялары жана туруктуу өнүгүү",
                "en": "Teaching methods and sustainable development"},
    "policy": {"ru": "Политика курса", "ky": "Курстун саясаты", "en": "Course policy"},
    "scale": {"ru": "Шкала оценок", "ky": "Баалардын шкаласы", "en": "Grading scale"},
    "kind_hours": {"ru": "Вид работы", "ky": "Иштин түрү", "en": "Activity"},
    "term": {"ru": "Термин", "ky": "Термин", "en": "Term"},
    "definition": {"ru": "Определение", "ky": "Аныктама", "en": "Definition"},
    "weight": {"ru": "Вес, %", "ky": "Салмагы, %", "en": "Weight, %"},
    "prereq": {"ru": "Пререквизиты", "ky": "Пререквизиттер", "en": "Prerequisites"},
    "postreq": {"ru": "Постреквизиты", "ky": "Постреквизиттер", "en": "Follow-on courses"},
    "u_competencies": {"ru": "Компетенции", "ky": "Компетенциялар", "en": "Competences"},
    "sdg_full": {"ru": "Цели устойчивого развития ООН", "ky": "БУУнун туруктуу өнүгүү максаттары", "en": "UN SDGs"},
    "q_remember": {"ru": "Перечислите основные понятия темы «{}».", "ky": "«{}» темасынын негизги түшүнүктөрүн атаңыз.",
                   "en": "List the key concepts of “{}”."},
    "q_understand": {"ru": "Объясните суть темы «{}» и приведите пример.", "ky": "«{}» темасынын маңызын түшүндүрүп, мисал келтириңиз.",
                     "en": "Explain “{}” and give an example."},
    "q_apply": {"ru": "Покажите на примере применение материала темы «{}».", "ky": "«{}» темасынын колдонулушун мисал менен көрсөтүңүз.",
                "en": "Show how “{}” is applied, using an example."},
    "q_analyze": {"ru": "Сравните подходы, рассмотренные в теме «{}».", "ky": "«{}» темасындагы ыкмаларды салыштырыңыз.",
                  "en": "Compare the approaches covered in “{}”."},
    "q_evaluate": {"ru": "Оцените достоинства и ограничения решений по теме «{}».",
                   "ky": "«{}» темасы боюнча чечимдердин артыкчылыктарын жана чектөөлөрүн баалаңыз.",
                   "en": "Evaluate the strengths and limits of solutions in “{}”."},
    "q_create": {"ru": "Предложите решение практической задачи по теме «{}».", "ky": "«{}» темасы боюнча практикалык маселенин чечимин сунуштаңыз.",
                 "en": "Propose a solution to a practical problem on “{}”."},
})
BLOOM = {"remember": "Знание", "understand": "Понимание", "apply": "Применение",
         "analyze": "Анализ", "evaluate": "Оценка", "create": "Создание"}
SKILL = {
    "syllabus": "/academic-core:syllabus", "lectures": "/academic-core:lecture", "srs": "/academic-core:srs",
    "practice": "/academic-core:srs", "tests": "/academic-core:tests", "glossary": "/academic-core:lecture (раздел «Ключевые понятия») или umkd.yaml",
    "questions": "umkd.yaml (skill umkd)", "guidelines": "umkd.yaml (skill umkd)",
}


class Umkd:
    def __init__(self, data: dict):
        self.data = data
        self.lang = data.get("language", "ru")
        self.dir: Path = data["_dir"]
        self.out = out_dir(data)
        ufile = self.dir / "umkd.yaml"
        self.extra = read_yaml(ufile) if ufile.exists() else {}
        sfile = self.dir / "srs.yaml"
        self.srs = read_yaml(sfile) if sfile.exists() else {}
        self.lectures = sorted(p for p in (self.dir / "lectures").glob("*.md") if not p.name.startswith("_")) \
            if (self.dir / "lectures").is_dir() else []
        self.doc = new_document()
        self.composer = None
        self.status: list[tuple[str, str, str]] = []
        self.warns: list[str] = []
        self.sec = 0

    # ---------- служебное ----------
    def heading(self, key: str):
        self.sec += 1
        h = self.doc.add_heading(f"{self.sec}. {t(key, self.lang)}", level=1)
        h.paragraph_format.page_break_before = True

    def mark(self, key: str, state: str, how: str = ""):
        self.status.append((t(key, self.lang), state, how))

    def md(self, text: str, shift: int = 1):
        render_docx(self.doc, parse_md(str(text)), self.dir, shift=shift)

    def merge(self, path: Path):
        from docx import Document
        from docxcompose.composer import Composer

        sub = Document(str(path))
        body = sub.element.body
        first = " ".join(p.text for p in sub.paragraphs[:4]).lower()
        if any(word in first for word in ("министерство", "министрлиги", "ministry")):
            for el in list(body):
                if el.tag.endswith("sectPr"):
                    break
                has_break = bool(el.xpath(".//w:br[@w:type='page']"))
                body.remove(el)
                if has_break:
                    break
        if self.composer is None:
            self.composer = Composer(self.doc)
        self.composer.append(sub)

    def find_out(self, *patterns: str) -> Path | None:
        for pat in patterns:
            hits = sorted(self.out.glob(pat)) + sorted(self.dir.glob(pat))
            if hits:
                return hits[0]
        return None

    # ---------- разделы ----------
    def syllabus(self):
        self.heading("u_syllabus")
        f = self.find_out(f"syllabus_{self.lang}.docx", "syllabus*.docx")
        if f:
            self.merge(f)
            self.mark("u_syllabus", "готово", f.name)
        else:
            para(self.doc, "[Раздел не собран: нет файла силлабуса]", italic=True)
            self.mark("u_syllabus", "нет", SKILL["syllabus"])

    def program(self):
        d, lang, doc = self.data, self.lang, self.doc
        self.heading("u_program")
        c = d.get("course", {}) or {}
        doc.add_heading(f'{self.sec}.1. {t("u_goals", lang)}', level=2)
        for g in d.get("goals", []) or []:
            para(doc, f"– {g}")
        if d.get("description"):
            para(doc, " ".join(str(d["description"]).split()))
        doc.add_heading(f'{self.sec}.2. {t("place", lang)}', level=2)
        for label, val in ((t("prereq", lang), c.get("prerequisites")), (t("postreq", lang), c.get("postrequisites")),
                           (t("u_competencies", lang), d.get("competencies"))):
            if val:
                p = para(doc, indent=False)
                add_runs(p, [{"text": f"{label}: ", "b": True}, {"text": ", ".join(map(str, val))}])
        doc.add_heading(f'{self.sec}.3. {t("outcomes", lang)}', level=2)
        table(doc, ["Код", t("outcomes", lang), "Уровень (Блум)"],
                   [[o.get("id", ""), o.get("text", ""), BLOOM.get(o.get("bloom", ""), o.get("bloom", ""))]
                    for o in d.get("outcomes", []) or []], widths_cm=[2, 10.5, 4])
        doc.add_heading(f'{self.sec}.4. {t("volume", lang)}', level=2)
        hours = d.get("hours", {}) or {}
        rows = [[t(k, lang), hours.get(k, 0)] for k in ("lectures", "practicals", "labs", "srsp", "srs") if hours.get(k)]
        rows.append([t("total", lang), sum(int(v or 0) for v in hours.values())])
        rows.append([t("credits", lang), f'{d.get("credits", "")} (× {CREDIT_HOURS} ч)'])
        table(doc, [t("kind_hours", lang), t("hours", lang)], rows, widths_cm=[10, 6.5])
        for p in check_hours(d):
            self.warns.append(f"часы: {p}")
        doc.add_heading(f'{self.sec}.5. {t("plan", lang)}', level=2)
        topics = d.get("topics", []) or []
        rows = [[tp.get("week", ""), tp.get("title", ""), tp.get("lectures", ""), tp.get("practicals", ""),
                 tp.get("labs", ""), tp.get("srs", ""), ", ".join(map(str, tp.get("outcomes", []) or []))]
                for tp in topics]
        rows.append(["", t("total", lang)] + [sum(int(tp.get(k, 0) or 0) for tp in topics)
                                              for k in ("lectures", "practicals", "labs", "srs")] + [""])
        table(doc, [t("week", lang), t("topic", lang), "Лек", "Пр", "Лаб", t("srs", lang), "РО"], rows,
                   widths_cm=[1.4, 7.2, 1.3, 1.3, 1.3, 1.6, 2.4])
        doc.add_heading(f'{self.sec}.6. {t("u_methods", lang)}', level=2)
        if d.get("methods"):
            para(doc, "Активные методы: " + ", ".join(d["methods"]) + ".")
        if d.get("sdg"):
            para(doc, f'{t("sdg_full", lang)}: ' + ", ".join(f"ЦУР {x}" for x in d["sdg"]) + ".")
        doc.add_heading(f'{self.sec}.7. {t("assessment", lang)}', level=2)
        assess = d.get("assessment", []) or []
        if assess:
            table(doc, [t("assessment", lang), t("week", lang), t("weight", lang)],
                       [[a.get("name", ""), a.get("week", ""), a.get("weight", "")] for a in assess],
                       widths_cm=[10.5, 2.5, 3.5])
        para(doc, f'{t("scale", lang)}:', bold=True, indent=False)
        li = {"ru": 3, "ky": 4, "en": 5}.get(lang, 3)
        table(doc, [t("points_range", lang), t("letter", lang), t("grade", lang)],
                   [[f"{r[0]}–{r[1]}", r[2], r[li]] for r in KR_SCALE], widths_cm=[4, 3, 9.5])
        pol = d.get("policies", {}) or {}
        if pol:
            doc.add_heading(f'{self.sec}.8. {t("policy", lang)}', level=2)
            for v in pol.values():
                para(doc, f"– {v}")
        self.mark("u_program", "готово", "из course.yaml")

    def lectures_part(self):
        self.heading("u_lectures")
        if not self.lectures:
            para(self.doc, "[Раздел не собран: нет лекций в папке lectures/]", italic=True)
            self.mark("u_lectures", "нет", SKILL["lectures"])
            return
        topics = len(self.data.get("topics", []) or [])
        from build_lecture import check_literature

        for i, md in enumerate(self.lectures):
            _, body = split_frontmatter(md.read_text(encoding="utf-8"))
            blocks = parse_md(body)
            self.warns += check_literature(blocks, self.data, md.name)
            if render_docx(self.doc, blocks, md.parent, shift=1, new_page=bool(i)):
                self.warns.append(f"{md.name}: есть пометки [нужен источник] — снимите до утверждения")
        state = "готово" if not topics or len(self.lectures) >= topics else f"частично: {len(self.lectures)} из {topics}"
        self.mark("u_lectures", state, "" if state == "готово" else SKILL["lectures"])

    def _tasks(self, practice: bool) -> list[dict]:
        words = ("практ", "лаб", "семинар", "practic", "lab", "seminar")
        return [tk for tk in self.srs.get("tasks", []) or []
                if any(w in str(tk.get("kind", "")).lower() for w in words) == practice]

    def practice(self):
        from build_srs import task_section

        self.heading("u_practice")
        tasks = self._tasks(True)
        if tasks:
            for i, tk in enumerate(tasks, 1):
                task_section(self.doc, i, tk, self.data, self.lang, self.warns)
            self.mark("u_practice", "готово", "srs.yaml")
            return
        topics = [tp for tp in self.data.get("topics", []) or [] if tp.get("practicals") or tp.get("labs")]
        table(self.doc, [t("week", self.lang), t("topic", self.lang), "Работа", "Пр", "Лаб"],
                   [[tp.get("week", ""), tp.get("title", ""), tp.get("lab", ""), tp.get("practicals", ""), tp.get("labs", "")]
                    for tp in topics], widths_cm=[1.4, 5.6, 6.9, 1.3, 1.3])
        self.mark("u_practice", "черновик: только темы и часы", SKILL["practice"])

    def srs_part(self):
        from build_srs import task_section

        self.heading("u_srs")
        tasks = self._tasks(False)
        if tasks:
            for i, tk in enumerate(tasks, 1):
                task_section(self.doc, i, tk, self.data, self.lang, self.warns)
            self.mark("u_srs", "готово", "srs.yaml")
            return
        topics = self.data.get("topics", []) or []
        table(self.doc, [t("week", self.lang), t("topic", self.lang), "Задание СРС", "Форма", "Срок", t("hours", self.lang)],
                   [[tp.get("week", ""), tp.get("title", ""), tp.get("srs_task", ""), tp.get("srs_form", ""),
                     tp.get("srs_due", ""), tp.get("srs", "")] for tp in topics],
                   widths_cm=[1.3, 4.2, 5.4, 2.5, 1.5, 1.6])
        self.mark("u_srs", "черновик: задания из паспорта без критериев", SKILL["srs"])

    def _auto_questions(self, topics: list[dict]) -> list[str]:
        bloom = {str(o.get("id")): o.get("bloom", "understand") for o in self.data.get("outcomes", []) or []}
        qs = []
        for tp in topics:
            levels = list(dict.fromkeys(bloom.get(str(o), "understand") for o in tp.get("outcomes", []) or [])) or ["understand"]
            for lv in levels:
                qs.append(t(f"q_{lv}", self.lang).format(tp.get("title", "")))
        return qs

    def assessment(self):
        self.heading("u_assessment")
        lang, doc = self.lang, self.doc
        n = 0
        tests = self.find_out("tests_1.docx", "tests*.docx", "tests/*.docx")
        n += 1
        doc.add_heading(f"{self.sec}.{n}. {t('u_tests', lang)}", level=2)
        if tests:
            self.merge(tests)
            self.mark("u_tests", f"готово ({tests.name}; ключ tests_key.docx не включён)", "")
        else:
            para(doc, "[Тесты не собраны]", italic=True)
            self.mark("u_tests", "нет", SKILL["tests"])
        topics = self.data.get("topics", []) or []
        auto = False
        dated = [a for a in self.data.get("assessment", []) or [] if a.get("week")]
        mids = [a for a in dated if any(w in str(a.get("name", "")).lower()
                                        for w in ("рубеж", "модул", "аралык", "midterm"))] or dated
        given_mid = self.extra.get("midterm_questions") or {}
        prev = 0
        for k, a in enumerate(mids, 1):
            n += 1
            doc.add_heading(f"{self.sec}.{n}. {t('u_midterm', lang)} {k} ({t('week', lang).lower()} {a['week']})", level=2)
            qs = given_mid.get(k) or given_mid.get(str(k))
            if not qs:
                qs = self._auto_questions([tp for tp in topics if prev < int(tp.get("week", 0) or 0) <= int(a["week"])])
                auto = True
            prev = int(a["week"])
            for i, q in enumerate(qs, 1):
                para(doc, f"{i}. {q}")
        n += 1
        doc.add_heading(f"{self.sec}.{n}. {t('u_exam', lang)}", level=2)
        qs = self.extra.get("exam_questions")
        if not qs:
            qs, auto = self._auto_questions(topics), True
        for i, q in enumerate(qs, 1):
            para(doc, f"{i}. {q}")
        if auto:
            self.mark("u_assessment", "черновик: вопросы по шаблону из тем и РО", SKILL["questions"])
        else:
            self.mark("u_assessment", "готово", "umkd.yaml")

    def glossary(self):
        self.heading("u_glossary")
        terms: dict[str, str] = {}
        for it in self.extra.get("glossary", []) or []:
            terms[str(it.get("term", "")).strip()] = str(it.get("definition", "")).strip()
        source = "umkd.yaml" if terms else ""
        for md in self.lectures:
            _, body = split_frontmatter(md.read_text(encoding="utf-8"))
            for runs in section_items(parse_md(body), ("понят", "термин", "түшүнүк", "concept", "glossary")):
                text = runs_text(runs)
                m = re.match(r"\s*(.+?)\s+[—–-]\s+(.+)", text)
                if m and m.group(1).strip() not in terms:
                    terms[m.group(1).strip()] = m.group(2).strip()
                    source = source or "лекции"
        if not terms:
            para(self.doc, "[Глоссарий пуст: нет раздела «Ключевые понятия» в лекциях]", italic=True)
            self.mark("u_glossary", "нет", SKILL["glossary"])
            return
        table(self.doc, [t("term", self.lang), t("definition", self.lang)],
                   [[k, v] for k, v in sorted(terms.items(), key=lambda kv: kv[0].lower())], widths_cm=[4.5, 12])
        self.mark("u_glossary", f"готово ({len(terms)} терминов)", source)

    def guidelines(self):
        self.heading("u_guidelines")
        text = self.extra.get("guidelines")
        if text:
            self.md(text, shift=1)
            self.mark("u_guidelines", "готово", "umkd.yaml")
            return
        d, doc = self.data, self.doc
        hours = d.get("hours", {}) or {}
        para(doc, f"Дисциплина «{course_title(d, self.lang)}» рассчитана на {d.get('credits', '')} кредитов "
                  f"({sum(int(v or 0) for v in hours.values())} ч), из них самостоятельная работа — {hours.get('srs', 0)} ч.")
        para(doc, "Как работать с курсом:", bold=True, indent=False)
        steps = [
            "До лекции прочитайте вопросы для подготовки (перевёрнутый класс) и материал из списка литературы.",
            "На лекции отвечайте на вопросы, участвуйте в разборе кейсов; конспект сверяйте с конспектом лекций УМКД.",
            "Практические и лабораторные задания выполняйте по плану из раздела «Планы практических и лабораторных занятий».",
            "Задания СРС сдавайте в срок; каждое оценивается по критериям, указанным в задании.",
            "Перед рубежным контролем и экзаменом пройдите вопросы для самопроверки и глоссарий.",
        ]
        for i, s in enumerate(steps, 1):
            para(doc, f"{i}. {s}")
        assess = d.get("assessment", []) or []
        if assess:
            para(doc, "Итоговая оценка складывается так: " +
                 "; ".join(f"{a.get('name')} — {a.get('weight')}%" for a in assess) + ".")
        for v in (d.get("policies", {}) or {}).values():
            para(doc, f"– {v}")
        oh = (d.get("teacher") or {}).get("office_hours")
        if oh:
            para(doc, f"Консультации преподавателя: {oh}.")
        self.mark("u_guidelines", "черновик: типовой текст из паспорта", SKILL["guidelines"])

    def literature(self):
        self.heading("literature")
        lang, n = self.lang, 0
        groups = {"main": t("main_lit", lang), "extra": t("extra_lit", lang), "online": t("online_lit", lang)}
        bad = 0
        for g, title in groups.items():
            items = [it for it in passport_literature(self.data) if it["group"] == g]
            if not items:
                continue
            para(self.doc, f"{title}:", bold=True, indent=False)
            for it in items:
                n += 1
                if it["verified"]:
                    para(self.doc, f"{n}. {it['text']}")
                else:
                    bad += 1
                    add_runs(para(self.doc), [{"text": f"{n}. {it['text']} {t('unverified', lang)}"}])
        if bad:
            self.warns.append(f"литература: {bad} источников без проверки помечены {t('unverified', lang)}")
        self.mark("literature", "готово" if n and not bad else ("нет" if not n else "есть непроверенные"),
                  "" if n and not bad else "/academic-core:new-course (литература в паспорте)")

    def approval(self):
        self.heading("u_approval")
        d, doc = self.data, self.doc
        prof, teacher = d.get("_profile", {}), d.get("teacher", {}) or {}
        program = (d.get("course", {}) or {}).get("program", "")
        para(doc, f"Учебно-методический комплекс дисциплины «{course_title(d, self.lang)}» составлен в соответствии "
                  f"с ГОС ВПО Кыргызской Республики и учебным планом по направлению {program}.")
        para(doc)
        rows = [
            ("Разработчик", ", ".join(x for x in (teacher.get("position"), teacher.get("degree")) if x), teacher.get("name", "")),
            (f"Рассмотрен на заседании кафедры «{prof.get('department', '')}»", "протокол № ____ от «___» __________ 20__ г.", ""),
            ("Заведующий кафедрой", "", ""),
            ("Одобрен учебно-методической комиссией факультета", "протокол № ____ от «___» __________ 20__ г.", ""),
            ("Председатель УМК факультета", "", ""),
        ]
        for role, detail, name in rows:
            para(doc, role, bold=True, indent=False)
            if detail:
                para(doc, detail, indent=False)
            para(doc, f"_______________________  {name or '_______________________'}", indent=False)
            para(doc)
        self.mark("u_approval", "готово (заполнить подписи)", "")

    def build(self) -> Path:
        d = self.data
        title_page(self.doc, d, self.lang, t("umkd", self.lang), approve=True)
        add_toc(self.doc, self.lang)
        for part in (self.syllabus, self.program, self.lectures_part, self.practice, self.srs_part,
                     self.assessment, self.glossary, self.guidelines, self.literature, self.approval):
            part()
        path = self.out / "umkd.docx"
        if self.composer is not None:
            self.composer.save(str(path))
        else:
            self.doc.save(path)
        return path


def main() -> None:
    ap = argparse.ArgumentParser(description="Сборка УМКД дисциплины в docx")
    ap.add_argument("course", help="путь к course.yaml")
    a = ap.parse_args()
    try:
        data = load_course(a.course)
        if not data.get("topics"):
            raise DataError("в course.yaml нет тем (topics) — заполните паспорт курса")
        u = Umkd(data)
        path = u.build()
    except DataError as e:
        fail(str(e))
    print(path)
    print("\nРазделы УМКД:")
    for name, state, how in u.status:
        print(f"  {name}: {state}" + (f" → {how}" if how and not state.startswith("готово") else ""))
    for w in u.warns:
        print(f"Предупреждение: {w}")


if __name__ == "__main__":
    main()
