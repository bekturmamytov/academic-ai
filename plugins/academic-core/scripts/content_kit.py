"""Общий код скриптов учебных материалов: лекции, СРС, УМКД, пособия, отчёты.

Markdown разбирается в простое дерево блоков (parse_md), из него собирается
docx (render_docx) или JSON для Typst (пособие). Без pandoc.
Импорт: sys.path.insert(0, str(Path(__file__).parent)); from content_kit import ...
"""
from __future__ import annotations

import difflib
import re
from pathlib import Path

import yaml

from common import DataError, find_up, read_yaml
from i18n import LABELS, t

LABELS.update({
    "toc": {"ru": "Содержание", "ky": "Мазмуну", "en": "Contents"},
    "toc_hint": {
        "ru": "Оглавление обновится при открытии в Word (если нет — правой кнопкой по этому тексту → «Обновить поле»).",
        "ky": "Мазмуну Word'до ачканда жаңырат (болбосо — оң баскыч → «Талааны жаңыртуу»).",
        "en": "The table of contents updates when opened in Word (otherwise right-click → Update Field).",
    },
    "lecture": {"ru": "Лекция", "ky": "Лекция", "en": "Lecture"},
    "lectures_book": {"ru": "Конспект лекций", "ky": "Лекциялардын конспекти", "en": "Lecture notes"},
    "points_range": {"ru": "Баллы", "ky": "Упай", "en": "Score"},
    "grade": {"ru": "Оценка", "ky": "Баа", "en": "Grade"},
})

MARKER_RE = re.compile(
    r"\[(?:НЕ ПРОВЕРЕНО|нужен источник|нужны ваши данные|ТЕКШЕРИЛГЕН ЭМЕС|булак керек|"
    r"UNVERIFIED|source needed)[^\]]*\]",
    re.IGNORECASE,
)
INLINE_RE = re.compile(
    r"\*\*(?P<b>.+?)\*\*|__(?P<b2>.+?)__|(?<![\w*])\*(?P<i>[^*\s][^*]*?)\*(?!\*)"
    r"|`(?P<code>[^`]+)`|\[(?P<lt>[^\]]+)\]\((?P<lu>[^)\s]+)\)"
)


# ---------- Markdown → блоки ----------

def split_frontmatter(text: str) -> tuple[dict, str]:
    text = text.lstrip("\ufeff")
    if text.startswith("---"):
        m = re.match(r"---\s*\n(.*?)\n---\s*\n?(.*)", text, re.S)
        if m:
            try:
                return yaml.safe_load(m.group(1)) or {}, m.group(2)
            except yaml.YAMLError as e:
                raise DataError(f"ошибка в шапке (frontmatter) markdown: {e}")
    return {}, text


def parse_inline(text: str) -> list[dict]:
    runs, pos = [], 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            runs.append({"text": text[pos:m.start()]})
        if m.group("b") or m.group("b2"):
            runs.append({"text": m.group("b") or m.group("b2"), "b": True})
        elif m.group("i"):
            runs.append({"text": m.group("i"), "i": True})
        elif m.group("code"):
            runs.append({"text": m.group("code"), "code": True})
        else:
            runs.append({"text": f'{m.group("lt")} ({m.group("lu")})'})
        pos = m.end()
    if pos < len(text):
        runs.append({"text": text[pos:]})
    return runs


def runs_text(runs: list[dict]) -> str:
    return "".join(r["text"] for r in runs)


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse_md(text: str) -> list[dict]:
    """Блоки: h, p, ul, ol, table, quote, img, pagebreak."""
    blocks: list[dict] = []
    para: list[str] = []
    lines = text.splitlines()

    def flush():
        if para:
            blocks.append({"t": "p", "runs": parse_inline(" ".join(s.strip() for s in para))})
            para.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            flush()
        elif m := re.match(r"(#{1,6})\s+(.*)", s):
            flush()
            blocks.append({"t": "h", "level": len(m.group(1)), "runs": parse_inline(m.group(2).strip())})
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.fullmatch(r"\|?[\s:\-|]+\|?", lines[i].strip()):
                    rows.append([parse_inline(c) for c in _cells(lines[i])])
                i += 1
            if rows:
                blocks.append({"t": "table", "header": rows[0], "rows": rows[1:]})
            continue
        elif m := re.match(r"(\s*)[-*+]\s+(.*)", line):
            flush()
            if blocks and blocks[-1]["t"] == "ul":
                items = blocks[-1]["items"]
            else:
                items = []
                blocks.append({"t": "ul", "items": items})
            items.append({"level": min(len(m.group(1).expandtabs(4)) // 2, 2), "runs": parse_inline(m.group(2))})
        elif m := re.match(r"(\s*)(\d+)[.)]\s+(.*)", line):
            flush()
            if blocks and blocks[-1]["t"] == "ol":
                items = blocks[-1]["items"]
            else:
                items = []
                blocks.append({"t": "ol", "items": items})
            items.append({"level": min(len(m.group(1).expandtabs(4)) // 2, 2), "n": m.group(2),
                          "runs": parse_inline(m.group(3))})
        elif s.startswith(">"):
            flush()
            q = s.lstrip(">").strip()
            if blocks and blocks[-1]["t"] == "quote" and q:
                blocks[-1]["runs"] += [{"text": " "}] + parse_inline(q)
            elif q:
                blocks.append({"t": "quote", "runs": parse_inline(q)})
        elif m := re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", s):
            flush()
            blocks.append({"t": "img", "alt": m.group(1), "path": m.group(2)})
        elif s in ("\\pagebreak", "<!-- pagebreak -->"):
            flush()
            blocks.append({"t": "pagebreak"})
        elif re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", s):
            flush()
        else:
            # строка-продолжение пункта списка
            if not para and blocks and blocks[-1]["t"] in ("ul", "ol") and line.startswith((" ", "\t")):
                blocks[-1]["items"][-1]["runs"] += [{"text": " "}] + parse_inline(s)
            else:
                para.append(s)
        i += 1
    flush()
    return blocks


def section_items(blocks: list[dict], heading_words: tuple[str, ...]) -> list[list[dict]]:
    """Пункты списков (runs) под заголовком, в котором есть одно из слов."""
    out, inside, lvl = [], False, 0
    for b in blocks:
        if b["t"] == "h":
            title = runs_text(b["runs"]).lower()
            if inside and b["level"] <= lvl:
                inside = False
            if any(w in title for w in heading_words):
                inside, lvl = True, b["level"]
            continue
        if inside and b["t"] in ("ul", "ol"):
            out += [it["runs"] for it in b["items"]]
        elif inside and b["t"] == "p":
            out.append(b["runs"])
    return out


# ---------- блоки → docx ----------

def add_runs(p, runs: list[dict], *, size=None, bold=None) -> int:
    """Добавляет runs в абзац; пометки [НЕ ПРОВЕРЕНО] и т. п. — жёлтым. Возвращает число пометок."""
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Pt

    marks = 0
    for r in runs:
        parts = MARKER_RE.split(r["text"])
        found = MARKER_RE.findall(r["text"])
        for k, chunk in enumerate(parts):
            for txt, is_mark in ((chunk, False), (found[k] if k < len(found) else "", True)):
                if not txt:
                    continue
                run = p.add_run(txt)
                run.bold = True if r.get("b") or bold else None
                run.italic = True if r.get("i") else None
                if r.get("code"):
                    run.font.name = "Courier New"
                if size:
                    run.font.size = Pt(size)
                if is_mark:
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    marks += 1
    return marks


def page_break(doc):
    from docx.enum.text import WD_BREAK

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _list_para(doc, runs, marker: str, level: int) -> int:
    from docx.shared import Cm

    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(1.25 + 0.75 * level + 0.6)
    pf.first_line_indent = Cm(-0.6)
    p.add_run(marker + "\t")
    pf.tab_stops.add_tab_stop(pf.left_indent)
    return add_runs(p, runs)


def render_docx(doc, blocks: list[dict], base_dir: Path | None = None, *, shift: int = 0,
                new_page: bool = False) -> int:
    """Пишет блоки в документ. shift сдвигает уровни заголовков (# → Heading 1+shift);
    new_page — начать с новой страницы (разрыв перед первым абзацем, без пустых страниц)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    from docx.text.paragraph import Paragraph

    from common import table

    start = len(doc.element.body)
    marks = 0
    for b in blocks:
        kind = b["t"]
        if kind == "h":
            lvl = b["level"] + shift
            if lvl <= 3:
                h = doc.add_heading(level=lvl)
                marks += add_runs(h, [{"text": runs_text(b["runs"])}])
            else:
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.keep_with_next = True
                marks += add_runs(p, b["runs"], bold=True)
        elif kind == "p":
            marks += add_runs(doc.add_paragraph(), b["runs"])
        elif kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.25)
            p.paragraph_format.first_line_indent = Cm(0)
            marks += add_runs(p, [dict(r, i=True) for r in b["runs"]])
        elif kind in ("ul", "ol"):
            for it in b["items"]:
                marker = f'{it["n"]}.' if kind == "ol" else ("–" if it["level"] == 0 else "•")
                marks += _list_para(doc, it["runs"], marker, it["level"])
        elif kind == "table":
            width = len(b["header"])
            rows = [[runs_text(c) for c in (r + [[]] * width)[:width]] for r in b["rows"]]
            table(doc, [runs_text(c) for c in b["header"]], rows)
            marks += sum(len(MARKER_RE.findall(" ".join(r))) for r in rows)
        elif kind == "img":
            path = (base_dir / b["path"]) if base_dir else Path(b["path"])
            if path.exists():
                doc.add_picture(str(path), width=Cm(14))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if b["alt"]:
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_runs(cap, [{"text": b["alt"]}], size=12)
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(f"[Рисунок: {b['alt'] or b['path']} — файл не найден]")
                r.italic = True
                r.font.size = Pt(12)
        elif kind == "pagebreak":
            page_break(doc)
    if new_page:
        first = next((el for el in list(doc.element.body)[start - 1:] if el.tag.endswith("}p")), None)
        if first is not None:
            Paragraph(first, doc._body).paragraph_format.page_break_before = True
    return marks


def add_toc(doc, lang: str, levels: str = "1-2", *, new_page: bool = False):
    """Заголовок «Содержание» + поле TOC; Word обновит его при открытии."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    from common import para

    para(doc, t("toc", lang).upper(), bold=True, align="center").paragraph_format.page_break_before = new_page
    p = doc.add_paragraph()
    run = p.add_run()
    for kind, text in (("begin", None), (None, f'TOC \\o "{levels}" \\h \\z \\u'), ("separate", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    hint = p.add_run(t("toc_hint", lang))
    hint.italic = True
    end = p.add_run()
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), "end")
    end._r.append(el)
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)


# ---------- шкала оценок КР ----------

KR_SCALE = [  # (от, до, буква, оценка ru, ky, en)
    (87, 100, "A", "отлично", "эң жакшы", "excellent"),
    (80, 86, "B", "хорошо", "жакшы", "good"),
    (74, 79, "C", "хорошо", "жакшы", "good"),
    (68, 73, "D", "удовлетворительно", "канааттандырарлык", "satisfactory"),
    (61, 67, "E", "удовлетворительно", "канааттандырарлык", "satisfactory"),
    (0, 60, "F", "неудовлетворительно", "канааттандырарлык эмес", "fail"),
]


def kr_grade(score: float) -> tuple[str, int]:
    """Буква и индекс строки KR_SCALE для балла 0–100."""
    for i, row in enumerate(KR_SCALE):
        if score >= row[0]:
            return row[2], i
    return "F", len(KR_SCALE) - 1


# ---------- литература из паспорта ----------

def _norm(s: str) -> str:
    return re.sub(r"[\W_]+", " ", s.lower()).strip()


def passport_literature(data: dict) -> list[dict]:
    lit = data.get("literature", {}) or {}
    items = []
    for group in ("main", "extra", "online"):
        for it in lit.get(group, []) or []:
            if isinstance(it, str):
                it = {"text": it, "verified": False}
            items.append({"text": it.get("text", ""), "verified": bool(it.get("verified")), "group": group})
    return items


def literature_status(text: str, data: dict) -> str:
    """'ok' — проверенный источник паспорта; 'unverified' — в паспорте без проверки; 'absent' — нет в паспорте."""
    n = _norm(MARKER_RE.sub("", text))
    if not n:
        return "absent"
    best, status = 0.0, "absent"
    for it in passport_literature(data):
        pn = _norm(it["text"])
        if not pn:
            continue
        score = 1.0 if (n in pn or pn in n or n[:40] == pn[:40]) else difflib.SequenceMatcher(None, n, pn).ratio()
        if score > best:
            best, status = score, ("ok" if it["verified"] else "unverified")
    return status if best >= 0.8 else "absent"


def find_course(start: Path) -> Path | None:
    start = start if start.is_dir() else start.parent
    return find_up(start, "course.yaml")


def load_yaml_doc(path: str | Path) -> dict:
    p = Path(path).expanduser().resolve()
    if not p.exists():
        raise DataError(f"не найден файл {p}")
    return read_yaml(p)


def load_context(start: Path, course: str | Path | None = None) -> dict:
    """Паспорт курса, если он есть; иначе только профиль и бренд вуза (для пособий и отчётов)."""
    from common import load_course

    start = start if start.is_dir() else start.parent
    path = Path(course).expanduser().resolve() if course else find_course(start)
    if path:
        return load_course(path)
    prof = find_up(start, "profile.yaml")
    uni = find_up(start, "university")
    brand = uni / "brand.yaml" if uni else None
    profile = read_yaml(prof) if prof else {}
    return {
        "_dir": start, "_profile": profile, "_brand_dir": uni,
        "_brand": read_yaml(brand) if brand and brand.exists() else {},
        "teacher": {k: profile.get(k, "") for k in ("name", "degree", "position", "email")},
        "language": profile.get("language", "ru"),
    }


def mark_runs(runs: list[dict]) -> list[dict]:
    """Разбивает runs так, чтобы пометки [нужен источник] стали отдельными runs с mark=True."""
    out = []
    for r in runs:
        pos = 0
        for m in MARKER_RE.finditer(r["text"]):
            if m.start() > pos:
                out.append(dict(r, text=r["text"][pos:m.start()]))
            out.append(dict(r, text=m.group(0), mark=True))
            pos = m.end()
        if pos < len(r["text"]):
            out.append(dict(r, text=r["text"][pos:]))
    return out
