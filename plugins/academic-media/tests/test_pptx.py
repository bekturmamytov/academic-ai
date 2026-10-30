from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from common import load_slides
from slides_from_pptx import extract


def make_pptx(path: Path) -> Path:
    prs = Presentation()
    prs.core_properties.subject = "Базы данных"
    prs.core_properties.author = "Асанова А. Т."

    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Нормализация"
    s.placeholders[1].text = "Лекция 5"
    s.notes_slide.notes_text_frame.text = "Здравствуйте, коллеги."

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Аномалии"
    body = s.placeholders[1].text_frame
    body.text = "Аномалия вставки"
    body.add_paragraph().text = "Аномалия удаления"
    s.notes_slide.notes_text_frame.text = "Есть две аномалии."

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Нормальные формы"
    tbl = s.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for r, row in enumerate([["Форма", "Требование"], ["1НФ", "Атомарность"], ["2НФ", "Нет частичных"]]):
        for c, val in enumerate(row):
            tbl.cell(r, c).text = val

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Вопросы?"

    prs.save(str(path))
    return path


def test_extract_types_notes_table(tmp_path):
    data = extract(make_pptx(tmp_path / "lec.pptx"))
    assert data["course"] == "Базы данных"
    assert data["teacher"] == "Асанова А. Т."
    kinds = [s["type"] for s in data["slides"]]
    assert kinds == ["title", "bullets", "table", "closing"]
    assert data["slides"][0]["subtitle"] == "Лекция 5"
    assert data["slides"][0]["notes"] == "Здравствуйте, коллеги."
    assert data["slides"][1]["bullets"] == ["Аномалия вставки", "Аномалия удаления"]
    assert data["slides"][2]["table"] == {"header": ["Форма", "Требование"],
                                          "rows": [["1НФ", "Атомарность"], ["2НФ", "Нет частичных"]]}
    assert data["slides"][3]["notes"] == ""


def test_load_slides_accepts_pptx(tmp_path):
    deck = load_slides(make_pptx(tmp_path / "lec.pptx"))
    assert deck["title"] == "Нормализация"
    assert deck["accent"] == "#1F4E79"
    assert len(deck["slides"]) == 4
